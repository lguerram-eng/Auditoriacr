# -*- coding: utf-8 -*-
"""Construye el papel de trabajo de auditoría de Obligaciones Financieras (cuenta 21 / detalle 62)
de GUAICARAMO S.A.S. con corte 31-jul-2026.

Fuentes (carpeta UPLOADS):
  - Obl_Financieros.xlsx  -> Sheet1: movimiento auxiliar cuenta 21 (ene-jul 2026); Hoja1: balance de prueba (29-sep-2026)
  - 62._Obligaciones_Financieras.xlsx -> detalle de obligaciones del cliente a jul-26
  - Balance_2026_GUUU.xlsx -> balance de prueba (24-ago-2026)
  - ilovepdf_merged_1.pdf -> extractos / certificados bancarios (datos transcritos en CERTS)
"""
import datetime as dt
import re
import sys

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

UP = sys.argv[1] if len(sys.argv) > 1 else "/root/.claude/uploads/9b3ddf83-4eba-59aa-a915-ef35201f01de"
OUT = sys.argv[2] if len(sys.argv) > 2 else "Cruce_Obligaciones_Financieras_62_Jul2026.xlsx"
F_MOV = f"{UP}/fef0ced2-Obl_Financieros.xlsx"
F_DET = f"{UP}/2b0f39e2-62._Obligaciones_Financieras.xlsx"
F_BAL = f"{UP}/a3876b2a-Balance_2026_GUUU.xlsx"

CORTE = dt.date(2026, 7, 31)

# ----------------------------------------------------------------------------- estilos
FN = "Arial"
f_base = Font(name=FN, size=9)
f_in = Font(name=FN, size=9, color="0000FF")
f_link = Font(name=FN, size=9, color="008000")
f_bold = Font(name=FN, size=9, bold=True)
f_hdr = Font(name=FN, size=9, bold=True, color="FFFFFF")
f_title = Font(name=FN, size=13, bold=True, color="1F3864")
f_sub = Font(name=FN, size=9, italic=True, color="595959")
f_sec = Font(name=FN, size=10, bold=True, color="1F3864")
fill_hdr = PatternFill("solid", fgColor="1F3864")
fill_sec = PatternFill("solid", fgColor="D9E1F2")
fill_tot = PatternFill("solid", fgColor="F2F2F2")
fill_key = PatternFill("solid", fgColor="FFFF00")
fill_bad = PatternFill("solid", fgColor="FCE4D6")
thin = Side(style="thin", color="BFBFBF")
box = Border(top=thin, bottom=thin, left=thin, right=thin)
NUM = '#,##0.00;(#,##0.00);"-"'
NUM0 = '#,##0;(#,##0);"-"'
PCT = '0.00%;(0.00%);"-"'
PB = '#,##0.0" pb";(#,##0.0" pb");"-"'
DATE = "yyyy-mm-dd"

wb = openpyxl.Workbook()
wb.remove(wb.active)


def ws_new(name, title, subtitle, widths):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws["A1"] = title
    ws["A1"].font = f_title
    ws["A2"] = subtitle
    ws["A2"].font = f_sub
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws


def put(ws, r, c, v, font=None, fmt=None, fill=None, bold=False, wrap=False, border=True, align=None):
    cell = ws.cell(row=r, column=c, value=v)
    if font is None:
        font = f_base
        if isinstance(v, str) and v.startswith("="):
            font = f_link if "!" in v else f_base
    if bold:
        font = Font(name=font.name, size=font.size, bold=True, color=font.color)
    cell.font = font
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    if border:
        cell.border = box
    if wrap or align:
        cell.alignment = Alignment(wrap_text=wrap, vertical="top", horizontal=align)
    return cell


def header(ws, r, labels, c0=1, height=30):
    for i, lab in enumerate(labels):
        cell = put(ws, r, c0 + i, lab, font=f_hdr, fill=fill_hdr, wrap=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[r].height = height


def section(ws, r, text, ncols):
    for c in range(1, ncols + 1):
        ws.cell(row=r, column=c).fill = fill_sec
    ws.cell(row=r, column=1, value=text).font = f_sec


def note(ws, r, c, text):
    ws.cell(row=r, column=c).comment = Comment(text, "Auditoría")


# ============================================================================ 1. MAYOR (movimiento)
src = openpyxl.load_workbook(F_MOV, data_only=True)
mov_rows = []
gran_total = None
for i, r in enumerate(src["Sheet1"].iter_rows(values_only=True), start=1):
    if i == 1:
        gran_total = (r[5], r[6], r[7])
        continue
    if r[1] is None:
        continue
    mov_rows.append((i,) + tuple(r))

ID_RULES = [  # (regex sobre descripción, ID)
    (r"1159383796|1155297469|115297469|CRED 7469", "BOG-1159383796"),
    (r"830624|CRED 0624", "DAV-830624"),
    (r"607691|CRED 7691", "DAV-607691"),
    (r"569032|CRED 9032", "DAV-569032"),
    (r"748906|CRED 8906|CREDI 8906", "DAV-8906"),
    (r"1019441|3000101441", "DAV-L1019441"),
    (r"1019448", "DAV-L1019448"),
    (r"1021275|1275-3", "DAV-L1021275"),
    (r"453436746", "BOG-453436746"),
    (r"856670629|CRED 0629", "BOG-856670629"),
    (r"858177259|CRED 7259", "BOG-858177259"),
    (r"556449224|LEAS 9224", "BOG-L556449224"),
    (r"556449180|LEAS 9180", "BOG-L556449180"),
    (r"556449117|LEAS 9117", "BOG-L556449117"),
    (r"556449064|LEAS 9064", "BOG-L556449064"),
    (r"556449288|LEAS 9288", "BOG-L556449288"),
    (r"557285005|55728505|LEAS 5005", "BOG-L557285005"),
    (r"631309940|63130940|CRED 9940", "POP-631309940"),
    (r"0141|CRED 4476", "BBVA-0141"),
    (r"1260105867|12660105867|CRED 5867", "BCL-1260105867"),
    (r"362095|LEAS 2095", "BCL-L362095"),
    (r"336817", "BCL-L336817"),
    (r"331330", "BCL-L331330"),
    (r"386735", "BCL-L386735"),
]


def assign_id(acct, terc, desc):
    if acct == "21050301":
        return "TC"
    if "VALLEY" in (terc or ""):
        return "N/A-COMISION"
    if "CRUCE DE DOCUMENTOS" in desc:
        return "N/A-CRUCE"
    if "DAVIVIENDA" in (terc or "") and "CREDI 8906" in desc:
        return "DAV-8906"
    for pat, ident in ID_RULES:
        if re.search(pat, desc):
            return ident
    if acct == "21300102" and "BANCOLOMBIA" in (terc or ""):
        return "BCL-L362095"
    return "SIN-ID"


CLASE = {
    "21050100": "Capital crédito CP",
    "21050201": "Capital crédito LP",
    "21050301": "Tarjeta de crédito",
    "21202040": "Capital leasing CP",
    "21202140": "Capital leasing LP",
    "21300101": "Interés OF por pagar",
    "21300102": "Interés leasing por pagar",
}

ws = ws_new("Mayor_Mov", "Movimiento auxiliar cuenta 21 – Obligaciones financieras (ene–jul 2026)",
            "Fuente: Obl_Financieros.xlsx / Sheet1 (Siesa). Columnas N–P asignadas por auditoría (ID de obligación, clase y tipo de movimiento).",
            [7, 10, 24, 30, 16, 16, 16, 11, 11, 13, 22, 60, 13, 17, 21, 20])
MH = 4
header(ws, MH, ["Fila origen", "Cuenta", "Nombre cuenta", "Tercero", "Débito", "Crédito", "Neto (D-C)", "Fecha",
                "NIT", "Documento", "Tipo documento", "Descripción", "Usuario", "ID obligación (auditoría)",
                "Clase", "Tipo movimiento"])
r = MH + 1
for row in mov_rows:
    (i, _a, acct, nom, _cc, terc, deb, cre, _net, fecha, nit, _t, doc, tdoc, desc, _x, user, *_rest) = row
    acct = str(acct).strip()
    desc = (desc or "").strip()
    ident = assign_id(acct, terc, desc)
    if ident.startswith("N/A") or ident == "TC":
        tipo = "Reclasificación / no obligación" if ident != "TC" else "Tarjeta de crédito"
    elif acct.startswith("2130"):
        tipo = "Causación" if (cre or 0) > 0 else ("Ajuste redondeo" if (deb or 0) < 1000 else "Reversión por pago")
    else:
        if (deb or 0) > 0 and (cre or 0) > 0:
            tipo = "Renovación / prórroga"
        elif (cre or 0) > 0:
            tipo = "Desembolso / nueva obligación"
        else:
            tipo = "Pago de capital"
    vals = [i, acct, nom, terc, deb or 0, cre or 0, f"=E{r}-F{r}", fecha, nit, (doc or "").strip(), tdoc, desc, user,
            ident, CLASE.get(acct, "Otra"), tipo]
    for c, v in enumerate(vals, start=1):
        fnt = f_in if c in (1, 2, 3, 4, 5, 6, 8, 9, 10, 11, 12, 13) else (f_bold if c == 14 else f_base)
        fmt = NUM if c in (5, 6, 7) else (DATE if c == 8 else None)
        put(ws, r, c, v, font=fnt, fmt=fmt)
    r += 1
M_FIRST, M_LAST = MH + 1, r - 1
put(ws, r, 4, "TOTAL MOVIMIENTO", bold=True, fill=fill_tot)
for c in (5, 6, 7):
    L = get_column_letter(c)
    put(ws, r, c, f"=SUM({L}{M_FIRST}:{L}{M_LAST})", fmt=NUM, bold=True, fill=fill_tot)
M_TOT = r
ws.freeze_panes = ws.cell(row=MH + 1, column=5)
ws.auto_filter.ref = f"A{MH}:P{M_LAST}"


def MR(col):
    return f"Mayor_Mov!${col}${M_FIRST}:${col}${M_LAST}"


R_CTA, R_DEB, R_CRE, R_FEC, R_DOC, R_ID = MR("B"), MR("E"), MR("F"), MR("H"), MR("J"), MR("N")

# ============================================================================ 2. BALANCES
def load_bal(f, off):
    wsb = openpyxl.load_workbook(f, data_only=True)["Hoja1"]
    d = {}
    for rr in wsb.iter_rows(values_only=True):
        if off:
            a, n, v = rr[1], rr[2], (rr[3], rr[4], rr[5], rr[6])
        else:
            a, n, v = rr[0], rr[1], (rr[7], rr[9], rr[11], rr[13])
        a = str(a or "").strip()
        if a and a[0].isdigit():
            d[a] = (str(n).strip(), [x if isinstance(x, (int, float)) else 0 for x in v])
    return d


BA = load_bal(F_MOV, 0)  # balance 29-sep-2026 (archivo Obl_Financieros / Hoja1)
BB = load_bal(F_BAL, 1)  # balance 24-ago-2026 (Balance_2026_GUUU)
n_diff_all = sum(1 for k in BA if k in BB and any(abs(x - y) > 0.005 for x, y in zip(BA[k][1], BB[k][1])))
BAL_ACCTS = ["1", "2", "3", "4", "5", "6", "7", "21", "2105", "210501", "21050100", "210502", "21050201", "21050301",
             "2120", "212020", "21202040", "212021", "21202140", "2130", "213001", "21300101", "21300102",
             "5305", "530520", "53052001", "53052002", "53052010", "53052020", "53052030", "53052040"]

ws = ws_new("Balance_21", "Balances de prueba – cuentas 21, 5305 y totales de clase",
            "Fuentes: Obl_Financieros.xlsx/Hoja1 (emitido 29-sep-2026) y Balance_2026_GUUU.xlsx (emitido 24-ago-2026). Libro NIIF, cifras en pesos; saldos crédito con signo negativo.",
            [11, 42, 18, 18, 18, 18, 18, 18, 18, 18, 14, 14, 14, 14])
BH = 5
put(ws, 4, 3, "Balance 29-sep-2026 (Obl_Financieros/Hoja1)", bold=True, border=False)
put(ws, 4, 7, "Balance 24-ago-2026 (Balance_2026_GUUU)", bold=True, border=False)
put(ws, 4, 11, "Diferencias (29-sep menos 24-ago)", bold=True, border=False)
header(ws, BH, ["Cuenta", "Descripción", "Saldo inicial 2026/01", "Débitos", "Créditos", "Saldo final 2026/07",
                "Saldo inicial 2026/01", "Débitos", "Créditos", "Saldo final 2026/07", "Dif. SI", "Dif. Déb.",
                "Dif. Créd.", "Dif. SF"])
BAL_ROW = {}
r = BH + 1
for a in BAL_ACCTS:
    put(ws, r, 1, a, font=f_in)
    put(ws, r, 2, BA[a][0], font=f_in)
    for j in range(4):
        put(ws, r, 3 + j, BA[a][1][j], font=f_in, fmt=NUM)
        put(ws, r, 7 + j, BB[a][1][j], font=f_in, fmt=NUM)
        put(ws, r, 11 + j, f"={get_column_letter(3 + j)}{r}-{get_column_letter(7 + j)}{r}", fmt=NUM)
    BAL_ROW[a] = r
    r += 1
B_FIRST, B_LAST = BH + 1, r - 1
r += 1
put(ws, r, 1, "Comparación integral de las dos versiones del balance (todas las cuentas):", bold=True, border=False)
r += 1
put(ws, r, 1, "Cuentas comparadas", border=True)
put(ws, r, 3, len(BA), font=f_in, fmt=NUM0)
r += 1
put(ws, r, 1, "Cuentas con alguna diferencia (SI, Déb., Créd. o SF)")
put(ws, r, 3, n_diff_all, font=f_in, fmt=NUM0)
note(ws, r, 3, "Comparación cuenta a cuenta de las 1.268 cuentas de ambos archivos realizada por auditoría (script). Resultado: 0 diferencias.")
BAL_NDIFF = f"Balance_21!$C${r}"


def BAL(acct, col):  # col: 'SI','D','C','SF'
    c = {"SI": "C", "D": "D", "C": "E", "SF": "F"}[col]
    return f"Balance_21!${c}${BAL_ROW[acct]}"


# ============================================================================ 3. CRUCE MOVIMIENTO VS BALANCE
ws = ws_new("Cruce_Mov_Balance", "Cruce movimiento auxiliar (cuenta 21) vs balance de prueba – Integridad",
            "SI balance + Débitos movimiento − Créditos movimiento = SF calculado; se compara con SF del balance. Signo: crédito negativo.",
            [11, 34, 18, 18, 18, 18, 18, 16, 18, 14, 18, 14, 12])
header(ws, 4, ["Cuenta", "Descripción", "Saldo inicial balance", "Débitos movimiento", "Créditos movimiento",
               "SF calculado", "SF balance", "Diferencia SF", "Débitos balance", "Dif. débitos", "Créditos balance",
               "Dif. créditos", "Estado"])
LEAF = ["21050100", "21050201", "21050301", "21202040", "21202140", "21300101", "21300102"]
r = 5
for a in LEAF:
    put(ws, r, 1, a, font=f_in)
    put(ws, r, 2, f"=INDEX(Balance_21!$B${B_FIRST}:$B${B_LAST},MATCH(A{r},Balance_21!$A${B_FIRST}:$A${B_LAST},0))")
    put(ws, r, 3, "=" + BAL(a, "SI"), fmt=NUM)
    put(ws, r, 4, f"=SUMIFS({R_DEB},{R_CTA},A{r})", fmt=NUM)
    put(ws, r, 5, f"=SUMIFS({R_CRE},{R_CTA},A{r})", fmt=NUM)
    put(ws, r, 6, f"=C{r}+D{r}-E{r}", fmt=NUM)
    put(ws, r, 7, "=" + BAL(a, "SF"), fmt=NUM)
    put(ws, r, 8, f"=F{r}-G{r}", fmt=NUM)
    put(ws, r, 9, "=" + BAL(a, "D"), fmt=NUM)
    put(ws, r, 10, f"=D{r}-I{r}", fmt=NUM)
    put(ws, r, 11, "=" + BAL(a, "C"), fmt=NUM)
    put(ws, r, 12, f"=E{r}-K{r}", fmt=NUM)
    put(ws, r, 13, f'=IF(MAX(ABS(H{r}),ABS(J{r}),ABS(L{r}))<=Resumen!$C$6,"OK","DIFERENCIA")')
    r += 1
put(ws, r, 2, "TOTAL CUENTA 21", bold=True, fill=fill_tot)
for c in range(3, 13):
    L = get_column_letter(c)
    put(ws, r, c, f"=SUM({L}5:{L}{r - 1})", fmt=NUM, bold=True, fill=fill_tot)
put(ws, r, 13, f'=IF(MAX(ABS(H{r}),ABS(J{r}),ABS(L{r}))<=Resumen!$C$6,"OK","DIFERENCIA")', bold=True, fill=fill_tot)
CMB_TOT = r
r += 1
put(ws, r, 2, "Total cuenta 21 según balance (mayor)")
put(ws, r, 3, "=" + BAL("21", "SI"), fmt=NUM)
put(ws, r, 9, "=" + BAL("21", "D"), fmt=NUM)
put(ws, r, 11, "=" + BAL("21", "C"), fmt=NUM)
put(ws, r, 7, "=" + BAL("21", "SF"), fmt=NUM)
r += 1
put(ws, r, 2, "Diferencia suma de auxiliares vs cuenta 21")
for c in (3, 7, 9, 11):
    L = get_column_letter(c)
    put(ws, r, c, f"={L}{CMB_TOT}-{L}{r - 1}", fmt=NUM)
r += 2

section(ws, r, "B. Integridad del archivo de movimiento ('Gran total' del reporte vs suma de líneas)", 13)
r += 1
header(ws, r, ["", "Concepto", "Gran total reporte", "Suma líneas", "Diferencia"])
r += 1
for lab, val, col in (("Débitos", gran_total[0], "E"), ("Créditos", gran_total[1], "F"), ("Neto", gran_total[2], "G")):
    put(ws, r, 2, lab)
    put(ws, r, 3, val, font=f_in, fmt=NUM)
    put(ws, r, 4, f"=Mayor_Mov!{col}{M_TOT}", fmt=NUM)
    put(ws, r, 5, f"=C{r}-D{r}", fmt=NUM)
    r += 1
put(ws, r, 2, "Nº de líneas del movimiento")
put(ws, r, 4, f"=COUNTA({R_CTA})", fmt=NUM0)
r += 2

section(ws, r, "C. Consistencia jerárquica del balance y relación pasivo–gasto", 13)
r += 1
header(ws, r, ["", "Prueba", "Valor A", "Valor B", "Diferencia", "", "Comentario"])
r += 1
tests = [
    ("SF 21 = 2105 + 2120 + 2130", "=" + BAL("21", "SF"), f"={BAL('2105','SF')}+{BAL('2120','SF')}+{BAL('2130','SF')}", ""),
    ("SF 2105 = 210501 + 210502", "=" + BAL("2105", "SF"), f"={BAL('210501','SF')}+{BAL('210502','SF')}", ""),
    ("SF 210502 = 21050201 + 21050301 (tarjeta de crédito agrupada en LP)", "=" + BAL("210502", "SF"),
     f"={BAL('21050201','SF')}+{BAL('21050301','SF')}", "Tarjeta de crédito (pasivo corriente) agrupada dentro de 'Obligaciones financieras a largo plazo' – ver hallazgo de clasificación."),
    ("SF 2120 = 21202040 + 21202140", "=" + BAL("2120", "SF"), f"={BAL('21202040','SF')}+{BAL('21202140','SF')}", ""),
    ("SF 2130 = 21300101 + 21300102", "=" + BAL("2130", "SF"), f"={BAL('21300101','SF')}+{BAL('21300102','SF')}", ""),
    ("Ecuación: Activo + Pasivo + Patrimonio + Resultado (4+5+6+7) = 0", f"={BAL('1','SF')}+{BAL('2','SF')}+{BAL('3','SF')}",
     f"=-({BAL('4','SF')}+{BAL('5','SF')}+{BAL('6','SF')}+{BAL('7','SF')})", "Resultado del periodo ene–jul incluido vía clases 4 a 7."),
    ("Intereses OF por pagar (21300101) = gasto causado neto (53052030)", "=-" + BAL("21300101", "SF"), "=" + BAL("53052030", "SF"),
     "El cliente reversa la causación al pagar y lleva el pago a 53052020; el saldo de 53052030 debe igualar el pasivo."),
    ("Intereses leasing por pagar (21300102) = gasto causado neto (53052040)", "=-" + BAL("21300102", "SF"), "=" + BAL("53052040", "SF"), ""),
    ("SI 21300101 = créditos 53052020 (reclasificación apertura)", "=-" + BAL("21300101", "SI"), "=" + BAL("53052020", "C"),
     "Asiento de apertura: Db 53052030 / Cr 53052020 por el saldo inicial causado."),
    ("SI 21300102 = créditos 53052010 (reclasificación apertura)", "=-" + BAL("21300102", "SI"), "=" + BAL("53052010", "C"), ""),
]
for lab, a_, b_, com in tests:
    put(ws, r, 2, lab, wrap=True)
    put(ws, r, 3, a_, fmt=NUM)
    put(ws, r, 4, b_, fmt=NUM)
    put(ws, r, 5, f"=C{r}-D{r}", fmt=NUM)
    put(ws, r, 7, com, wrap=True)
    ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=13)
    ws.row_dimensions[r].height = 26
    r += 1
CMB_ECUACION = r - 5  # fila de la ecuación patrimonial
r += 1
section(ws, r, "D. Comparación de versiones del balance", 13)
r += 1
put(ws, r, 2, "Cuentas con diferencias entre balance 24-ago y 29-sep")
put(ws, r, 3, "=" + BAL_NDIFF, fmt=NUM0)
put(ws, r, 4, f'=IF(C{r}=0,"Sin cambios posteriores al cierre de jul-26","Revisar ajustes posteriores")')
ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=8)
ws.freeze_panes = "C5"

# ============================================================================ 4. DETALLE 62
d_ws = openpyxl.load_workbook(F_DET, data_only=True).active
drows = [rr for rr in d_ws.iter_rows(values_only=True)]
rates = {drows[i][1]: drows[i][2] for i in range(1, 7)}
DET_ID = {
    "LEASING 1019441": "DAV-L1019441", "LEASING 1019448": "DAV-L1019448", "LEASING 10212753": "DAV-L1021275",
    "457300569032": "DAV-569032", "457300830624": "DAV-830624", "457300607691": "DAV-607691",
    "453436746": "BOG-453436746", "856670629": "BOG-856670629", "858177259": "BOG-858177259",
    "1159383796": "BOG-1159383796", "LEASING 556449224": "BOG-L556449224", "LEASING 556449180": "BOG-L556449180",
    "LEASING 556449117": "BOG-L556449117", "LEASING 556449064": "BOG-L556449064", "LEASING 556449288": "BOG-L556449288",
    "LEASING 557285005": "BOG-L557285005", "631309940-2": "POP-631309940", "252059": "BBVA-0141",
    "1260105867": "BCL-1260105867", "LEASING 362095": "BCL-L362095", "LEASING 336817": "BCL-L336817",
    "LEASING 331330": "BCL-L331330",
}
det = []
subtot_cli = []
bank = None
for rr in drows[7:]:
    if rr[1] in ("BANCO DAVIVIENDA", "BANCO DE BOGOTA", "BANCO POPULAR", "BBVA", "BANCOLOMBIA S.A."):
        bank = rr[1]
        continue
    if rr[1] == "Total Obligaciones":
        total_cli = rr[3]
        continue
    key = str(rr[1]).strip() if rr[1] is not None else None
    key = key if key is None else key
    if key in DET_ID or (key and key.isdigit() and key.lstrip("0") in [k.lstrip("0") for k in DET_ID]):
        k2 = key if key in DET_ID else [k for k in DET_ID if k.lstrip("0") == key.lstrip("0")][0]
        det.append((bank, rr[1], DET_ID[k2], rr))
    elif rr[3] is not None:
        subtot_cli.append((bank, rr[3], rr[14]))

ws = ws_new("Detalle_62", "Detalle de obligaciones financieras del cliente (papel 62) a 31-jul-2026 – recálculo",
            "Fuente: 62._Obligaciones_Financieras.xlsx. Columnas A–Q: datos del cliente (azul). Columnas R en adelante: recálculo de auditoría.",
            [17, 19, 17, 9, 26, 17, 9, 17, 9, 10, 7, 11, 16, 16, 8, 11, 9, 11, 11, 10, 15, 11, 16, 16, 10])
put(ws, 4, 1, "Tasas de referencia 31-jul-26 (cliente)", bold=True, border=False)
RATE_CELL = {}
rr_ = 5
for k, v in rates.items():
    put(ws, rr_, 1, k, font=f_in)
    put(ws, rr_, 2, v, font=f_in, fmt="0.00000" if k not in ("TRM",) else NUM)
    RATE_CELL[k] = f"Detalle_62!$B${rr_}"
    rr_ += 1
DH = 13
header(ws, DH, ["Banco", "Nº obligación (cliente)", "ID auditoría", "Tipo", "Descripción", "Saldo deuda jul-26",
                "% deuda (cliente)", "Tipo tasa", "Puntos adic.", "Tasa EA cliente", "Veces año", "Tipo pago",
                "Corto plazo (cliente)", "Largo plazo (cliente)", "Cuotas rest.", "Vencimiento", "Tasa pond. cliente",
                "Índice aplicado", "Tasa EA recalculada", "Dif. EA (pb)", "CP + LP − Saldo", "Estado CP/LP",
                "CP corregido", "LP corregido", "% deuda recalc."], height=40)
r = DH + 1
D_FIRST = r
for bank, num, ident, rr in det:
    tipo = "Leasing" if "-L" in ident else "Crédito"
    vals = [bank, num, ident, tipo, rr[2], rr[3], rr[4], rr[5], rr[6], rr[7], rr[8], rr[9], rr[10] or 0,
            rr[11] or 0, rr[12], rr[13], rr[14]]
    for c, v in enumerate(vals, start=1):
        fmt = {6: NUM, 7: PCT, 9: PCT, 10: PCT, 13: NUM, 14: NUM, 16: DATE, 17: PCT}.get(c)
        put(ws, r, c, v, font=f_bold if c == 3 else f_in, fmt=fmt)
    put(ws, r, 18, f'=IF(LEFT(H{r},11)="IBR 6 MESES",{RATE_CELL["IBR 6 MESES"]},IF(LEFT(H{r},11)="IBR 3 MESES",'
                   f'{RATE_CELL["IBR 3 MESES"]},IF(H{r}="DTF",{RATE_CELL["DTF"]},{RATE_CELL["IBR"]})))', fmt="0.000%")
    put(ws, r, 19, f"=(1+(R{r}+I{r})/K{r})^K{r}-1", fmt="0.000%")
    put(ws, r, 20, f"=(J{r}-S{r})*10000", fmt=PB)
    put(ws, r, 21, f"=M{r}+N{r}-F{r}", fmt=NUM)
    put(ws, r, 22, f'=IF(ABS(U{r})>Resumen!$C$6,"ERROR",IF(OR(M{r}<0,N{r}<0),"ERROR","OK"))')
    put(ws, r, 23, f"=MIN(MAX(M{r},0),F{r})", fmt=NUM)
    put(ws, r, 24, f"=F{r}-W{r}", fmt=NUM)
    put(ws, r, 25, f"=F{r}/SUMIFS($F${D_FIRST}:$F${D_FIRST + len(det) - 1},$A${D_FIRST}:$A${D_FIRST + len(det) - 1},A{r})", fmt=PCT)
    r += 1
D_LAST = r - 1
put(ws, r, 5, "TOTAL OBLIGACIONES (recalculado)", bold=True, fill=fill_tot)
for c in (6, 13, 14, 21, 23, 24):
    L = get_column_letter(c)
    put(ws, r, c, f"=SUM({L}{D_FIRST}:{L}{D_LAST})", fmt=NUM, bold=True, fill=fill_tot)
D_TOT = r
r += 1
put(ws, r, 5, "Total obligaciones según cliente")
put(ws, r, 6, total_cli, font=f_in, fmt=NUM)
r += 1
put(ws, r, 5, "Diferencia (recalculado − cliente)")
put(ws, r, 6, f"=F{D_TOT}-F{r - 1}", fmt=NUM)
D_TOTDIF = r
r += 2


def DR(col):
    return f"Detalle_62!${col}${D_FIRST}:${col}${D_LAST}"


section(ws, r, "Subtotales por banco – recálculo vs cliente", 25)
r += 1
header(ws, r, ["Banco", "", "", "", "", "Subtotal recalculado", "", "", "", "", "", "", "Subtotal cliente",
               "Diferencia", "", "", "Tasa pond. recalc.", "Tasa pond. cliente", "Dif. (pb)"])
r += 1
for bank_, sub_cli, tp_cli in subtot_cli:
    put(ws, r, 1, bank_, font=f_in)
    put(ws, r, 6, f"=SUMIFS($F${D_FIRST}:$F${D_LAST},$A${D_FIRST}:$A${D_LAST},A{r})", fmt=NUM)
    put(ws, r, 13, sub_cli, font=f_in, fmt=NUM)
    put(ws, r, 14, f"=F{r}-M{r}", fmt=NUM)
    put(ws, r, 17, f"=SUMPRODUCT(($A${D_FIRST}:$A${D_LAST}=A{r})*$F${D_FIRST}:$F${D_LAST}*$J${D_FIRST}:$J${D_LAST})/F{r}", fmt="0.000%")
    put(ws, r, 18, tp_cli, font=f_in, fmt="0.000%")
    put(ws, r, 19, f"=(R{r}-Q{r})*10000", fmt=PB)
    r += 1
r += 1
put(ws, r, 1, "Tasa EA promedio ponderada total (cliente)", bold=True)
put(ws, r, 6, f"=SUMPRODUCT({DR('F')},{DR('J')})/F{D_TOT}", fmt="0.000%")
D_TPOND = f"Detalle_62!$F${r}"
r += 1
put(ws, r, 1, "Tasa EA promedio ponderada total (recalculada)", bold=True)
put(ws, r, 6, f"=SUMPRODUCT({DR('F')},{DR('S')})/F{D_TOT}", fmt="0.000%")
r += 1
put(ws, r, 1, "Nota: la tasa EA recalculada usa la convención del cliente (1 + (índice + puntos)/n)^n − 1 con n = 'Veces año'. "
              "Diferencias indican que el cliente usó una periodicidad distinta a la declarada.", font=f_sub, border=False)
ws.freeze_panes = ws.cell(row=DH + 1, column=4)

# ============================================================================ 5. CERTIFICADOS
# (pág, banco, obligación, ID, tipo extracto, fecha corte, capital s/cert, abono capital posterior hasta 31-jul,
#  tasa EA cert, vencimiento cert, observación)
CERTS = [
    (1, "Davivienda", "7000457300830624", "DAV-830624", "Extracto crédito corporativo", dt.date(2026, 3, 21),
     "=7499998630.82-127.57", 0, 0.1551, dt.date(2044, 3, 21),
     "Saldo anterior 21-sep-25 $7.499.998.630,82 − abono capital 21-mar-26 $127,57. Próximo pago 21-sep-26 sólo intereses $565.738.188,80 (IBR semestral + 3,00)."),
    (2, "Banco de Bogotá", "00856670629", "BOG-856670629", "Extracto Sustitución Finagro", dt.date(2026, 5, 27),
     1639499992.59, 182166667.00, 0.1494, None,
     "Cuota 12 (21-jun-26): capital $182.166.667 + intereses $59.390.887,23. Valor aprobado $2.186 MM, plazo 60 meses."),
    (4, "Bancolombia", "01260105867", "BCL-1260105867", "Detalle del crédito (Sucursal Virtual)", dt.date(2026, 8, 11),
     3333333334.00, 0, 0.1417, dt.date(2028, 4, 22),
     "Capital vigente al 11-ago-26; sin abonos entre 31-jul y 11-ago. Intereses corrientes causados $24.607.407; próximo pago 22-oct-26. IBR + 1,5."),
    (6, "Banco de Bogotá", "00453436746", "BOG-453436746", "Extracto Ordinaria Comercial", dt.date(2026, 6, 22),
     2999999977.39, 62500000.00, 0.1696, None,
     "Cuota 100 (17-jul-26): capital $62.500.000 + intereses $38.894.999,71. Aprobado $7.500 MM, plazo 146 meses."),
    (8, "Banco de Bogotá", "00858177259", "BOG-858177259", "Extracto Sustitución Finagro", dt.date(2026, 7, 31),
     1124999996.41, 0, 0.1665, None,
     "Corte 31-jul-26 (fecha de auditoría). Próxima cuota 25-ago-26: capital $125.000.000 + intereses $45.166.249,86."),
    (10, "Banco de Bogotá", "01155297469 → prorrogado 1159383796", "BOG-1159383796", "Extracto Sustitución Finagro",
     dt.date(2026, 6, 17), 10000000000.00, 0, 0.1239, None,
     "Cuota 2 (14-jul-26) por $10.299.617.737,55 incluye capital $10.000 MM. Contabilidad registra prórroga (NB-1016, 29-jul) al crédito 1159383796 – sin certificado del nuevo crédito."),
    (12, "Davivienda", "7000457300607691", "DAV-607691", "Extracto crédito corporativo", dt.date(2026, 1, 28),
     3959998435.90, "=272300000-272299382.05", 0.1258, dt.date(2041, 7, 28),
     "Pago 28-jul-26 de $272.300.000 vs intereses facturados $272.299.382,05: excedente $617,95 se aplica a capital."),
    (13, "Davivienda", "7000457300569032", "DAV-569032", "Extracto crédito corporativo", dt.date(2026, 1, 28),
     2999995943.80, "=206288000-206287213.20", 0.1258, dt.date(2041, 7, 28),
     "Pago 28-jul-26 de $206.288.000 vs intereses facturados $206.287.213,20: excedente $786,80 se aplica a capital."),
    ("p.14 + LEASING_1275-3_DAVIVIENDA.pdf", "Davivienda Leasing", "000030001021275-3", "DAV-L1021275", "Factura leasing", dt.date(2026, 7, 6),
     None, 0, 0.1441, None,
     "Canon 006 (23-jul-26) sólo intereses $64.041.851; 16 cuotas pactadas. El extracto NO informa saldo de capital → ver saldo implícito en Recalculo_Intereses."),
    (16, "Banco de Bogotá Leasing", "00557285005", "BOG-L557285005", "Estado de cuenta leasing", dt.date(2026, 7, 12),
     1047113003.63, 62866117.12, 0.1604, dt.date(2029, 7, 22),
     "Saldo capital tras pago 22-abr-26; canon 22-jul-26 capital $62.866.117,12 / costo financiero $39.578.631,88. Opción de compra $16.220.758."),
    (17, "Banco de Bogotá Leasing", "00556449064", "BOG-L556449064", "Estado de cuenta leasing", dt.date(2026, 7, 17),
     280274638.64, 26602533.70, 0.1425, dt.date(2028, 7, 27),
     "Saldo tras pago 27-abr-26; canon 27-jul-26 capital $26.602.533,70 / costo financiero $9.497.274,30. DTF."),
    (18, "Banco de Bogotá Leasing", "00556449117", "BOG-L556449117", "Estado de cuenta leasing", dt.date(2026, 7, 17),
     280274638.73, 26602533.71, 0.1425, dt.date(2028, 7, 27),
     "Saldo tras pago 27-abr-26; canon 27-jul-26 capital $26.602.533,71 / costo financiero $9.497.274,29. DTF."),
    (19, "Banco de Bogotá Leasing", "00556449180", "BOG-L556449180", "Estado de cuenta leasing", dt.date(2026, 7, 13),
     417650348.81, 21538319.01, 0.1422, dt.date(2030, 1, 23),
     "Saldo tras pago 23-abr-26; canon 23-jul-26 capital $21.538.319,01 / costo financiero $14.118.857,99. DTF."),
    (20, "Banco de Bogotá Leasing", "00556449224", "BOG-L556449224", "Estado de cuenta leasing", dt.date(2026, 7, 13),
     417650348.67, 21538319.02, 0.1422, dt.date(2030, 1, 23),
     "Saldo tras pago 23-abr-26; canon 23-jul-26 capital $21.538.319,02 / costo financiero $14.118.857,98. DTF."),
    (21, "Banco de Bogotá Leasing", "00556449288", "BOG-L556449288", "Estado de cuenta leasing", dt.date(2026, 7, 10),
     417736780.24, 21542838.16, 0.1422, dt.date(2030, 1, 20),
     "Saldo tras pago 20-abr-26; canon 20-jul-26 capital $21.542.838,16 / costo financiero $14.121.779,84. DTF."),
    ("p.22 + LEASING_9441-2_DAVIVIENDA.pdf", "Davivienda Leasing", "000030001019441-2", "DAV-L1019441", "Factura leasing", dt.date(2026, 6, 25),
     None, 0, 0.1603, None,
     "Canon 033 (13-jul-26): capital $12.327.298, intereses $5.062.372, seguro $1.263.646. Último pago 11-jun-26 canon 032 $1.029.562 (sólo intereses). Sin saldo de capital."),
    ("p.24 + LEASING_9448-3_DAVIVIENDA.pdf", "Davivienda Leasing", "000030001019448-3", "DAV-L1019448", "Factura leasing", dt.date(2026, 6, 18),
     None, 0, 0.1593, None,
     "Canon 031 (6-jul-26): capital $8.867.342, intereses $4.177.893. Último pago 5-jun-26 canon 030: capital $8.901.235, intereses $4.402.982. Sin saldo de capital."),
]
ws = ws_new("Certificados", "Cruce detalle 62 vs certificados / extractos bancarios (PDF adjunto)",
            "Fuente: ilovepdf_merged_1.pdf (25 págs.). Las págs. 3, 5, 7, 9, 11, 15, 23 y 25 son anexos/instructivos sin cifras. Capital banco 31-jul = capital certificado − abonos a capital entre corte y 31-jul.",
            [5, 20, 20, 24, 16, 22, 11, 17, 16, 17, 17, 16, 10, 10, 10, 11, 11, 11, 70])
CH = 4
header(ws, CH, ["#", "Pág. PDF", "Banco", "Obligación (certificado)", "ID auditoría", "Tipo documento", "Fecha corte",
                "Capital s/certificado", "Abonos capital posteriores al corte (≤31-jul)", "Capital banco 31-jul-26",
                "Saldo detalle 62", "Diferencia detalle − banco", "Tasa EA certificado", "Tasa EA detalle",
                "Dif. (pb)", "Venc. certificado", "Venc. detalle", "¿Venc. coincide?", "Observaciones"], height=48)
r = CH + 1
C_FIRST = r
CERT_ROW = {}
for n, (pg, bk, ob, ident, td, fc, cap, abono, ea, venc, obs) in enumerate(CERTS, start=1):
    put(ws, r, 1, n)
    put(ws, r, 2, pg, font=f_in, wrap=True)
    put(ws, r, 3, bk, font=f_in)
    put(ws, r, 4, ob, font=f_in)
    put(ws, r, 5, ident, font=f_bold)
    put(ws, r, 6, td, font=f_in)
    put(ws, r, 7, fc, font=f_in, fmt=DATE)
    put(ws, r, 8, cap, font=f_in, fmt=NUM)
    put(ws, r, 9, abono, font=f_in, fmt=NUM)
    put(ws, r, 10, f'=IF(H{r}="","N/D",H{r}-I{r})', fmt=NUM)
    put(ws, r, 11, f"=SUMIFS({DR('F')},{DR('C')},E{r})", fmt=NUM)
    put(ws, r, 12, f'=IF(H{r}="","N/D",K{r}-J{r})', fmt=NUM)
    put(ws, r, 13, ea, font=f_in, fmt=PCT)
    put(ws, r, 14, f"=INDEX({DR('J')},MATCH(E{r},{DR('C')},0))", fmt=PCT)
    put(ws, r, 15, f"=(N{r}-M{r})*10000", fmt=PB)
    put(ws, r, 16, venc, font=f_in, fmt=DATE)
    put(ws, r, 17, f"=INDEX({DR('P')},MATCH(E{r},{DR('C')},0))", fmt=DATE)
    put(ws, r, 18, f'=IF(P{r}="","N/D",IF(P{r}=Q{r},"Sí","NO"))')
    put(ws, r, 19, obs, font=f_in, wrap=True)
    ws.row_dimensions[r].height = 36
    CERT_ROW[ident] = r
    r += 1
C_LAST = r - 1
put(ws, r, 6, "TOTAL obligaciones con saldo certificado", bold=True, fill=fill_tot)
for c in (10, 11):
    L = get_column_letter(c)
    put(ws, r, c, f'=SUMIFS({L}{C_FIRST}:{L}{C_LAST},$H${C_FIRST}:$H${C_LAST},"<>")', fmt=NUM, bold=True, fill=fill_tot)
put(ws, r, 12, f"=SUM(L{C_FIRST}:L{C_LAST})", fmt=NUM, bold=True, fill=fill_tot)
C_TOT = r
r += 1
put(ws, r, 6, "Diferencia detalle − banco: leasings Banco de Bogotá")
put(ws, r, 12, f'=SUMIFS(L{C_FIRST}:L{C_LAST},E{C_FIRST}:E{C_LAST},"BOG-L*")', fmt=NUM, fill=fill_key)
C_BOGL = f"Certificados!$L${r}"
r += 1
put(ws, r, 6, "Diferencia detalle − banco: créditos (redondeos)")
put(ws, r, 12, f'=L{C_TOT}-L{r - 1}', fmt=NUM)
C_CRED = f"Certificados!$L${r}"
r += 2

section(ws, r, "B. Cobertura de confirmación – obligaciones del detalle 62 SIN certificado/extracto en el PDF", 19)
r += 1
header(ws, r, ["", "", "Banco", "Obligación", "ID auditoría", "Tipo", "", "", "", "", "Saldo detalle 62"])
r += 1
B_FIRST = r
cert_ids = {c[3] for c in CERTS}
for bank, num, ident, rr in det:
    if ident in cert_ids:
        continue
    put(ws, r, 3, bank, font=f_in)
    put(ws, r, 4, str(num), font=f_in)
    put(ws, r, 5, ident, font=f_bold)
    put(ws, r, 6, "Leasing" if "-L" in ident else "Crédito")
    put(ws, r, 11, f"=SUMIFS({DR('F')},{DR('C')},E{r})", fmt=NUM)
    r += 1
put(ws, r, 3, "Bancolombia", font=f_in)
put(ws, r, 4, "Leasing 386735 (no está en detalle 62)", font=f_in)
put(ws, r, 5, "BCL-L386735", font=f_bold)
put(ws, r, 6, "Leasing")
put(ws, r, 11, f'=SUMIFS({R_CRE},{R_ID},E{r})-SUMIFS({R_DEB},{R_ID},E{r})', fmt=NUM)
note(ws, r, 11, "Valor según mayor (NB-00001014, 28-jul-26), no incluido en el detalle 62.")
r += 1
put(ws, r, 6, "Total sin certificado (detalle 62)", bold=True, fill=fill_tot)
put(ws, r, 11, f"=SUM(K{B_FIRST}:K{r - 2})", fmt=NUM, bold=True, fill=fill_tot)
C_SINCERT = r
r += 1
put(ws, r, 6, "% del saldo del detalle 62 cubierto con certificado/extracto")
put(ws, r, 11, f"=1-K{C_SINCERT}/Detalle_62!$F${D_TOT}", fmt=PCT, fill=fill_key)
C_COBERT = f"Certificados!$K${r}"
r += 2

section(ws, r, "C. Cruce de pagos/intereses según certificado vs registro contable (comprobantes del mayor)", 19)
r += 1
header(ws, r, ["", "", "ID auditoría", "Concepto", "Documento contable", "Fecha", "", "Valor s/certificado",
               "Valor s/contabilidad", "Diferencia", "", "", "", "", "", "", "", "", "Comentario"])
r += 1
P_FIRST = r
PAGOS = [  # (ID, concepto, doc, campo ('D' débito / 'C' crédito / valor fijo), valor certificado, comentario)
    ("DAV-830624", "Pago 21-mar-26 total aplicado (int. 445.207.872,43 + cap. 127,57)", "NI-00004268", "D", 445208000.00,
     "El cliente reversa el total como interés; los $127,57 abonados a capital no se registran."),
    ("DAV-607691", "Pago 28-ene-26 total aplicado (int. 243.932.730,53 + cap. 269,47)", "NI-00004061", "D", 243933000.00,
     "Capital $269,47 no registrado."),
    ("DAV-569032", "Pago 28-ene-26 total aplicado (int. 184.797.372,21 + cap. 627,79)", "NI-00004063", "D", 184798000.00,
     "Capital $627,79 no registrado."),
    ("DAV-607691", "Intereses 28-ene a 28-jul-26 facturados", "NI-00004743", "D", 272299382.05,
     "Contabilidad reversa $272.300.000 (pago redondeado); $617,95 corresponden a capital."),
    ("DAV-569032", "Intereses 28-ene a 28-jul-26 facturados", "NI-00004745", "D", 206287213.20,
     "Contabilidad reversa $206.288.000; $786,80 corresponden a capital."),
    ("BOG-856670629", "Cuota 11 – capital", "CL-00000383", "D", 182166667.53, "Diferencia de redondeo del banco."),
    ("BOG-856670629", "Cuota 11 – intereses", "NI-00004260", "D", 56089116.47, ""),
    ("BOG-856670629", "Cuota 12 – capital", "CL-00000410", "D", 182166667.00, ""),
    ("BOG-856670629", "Cuota 12 – intereses", "NI-00004662", "D", 59390887.23, ""),
    ("BOG-453436746", "Cuota 99 – capital", "CL-00000409", "D", 62500000.09, ""),
    ("BOG-453436746", "Cuota 100 – capital", "CL-00000422", "D", 62500000.00,
     "Intereses cuota 100 $38.894.999,71 coinciden con la descripción del CL-422 (gasto directo 53052020)."),
    ("BOG-858177259", "Cuota mayo – capital", "CL-00000407", "D", 125000000.56, ""),
    ("BOG-858177259", "Cuota mayo – intereses", "NI-00004526", "D", 46619930.44, ""),
    ("BOG-1159383796", "Crédito 1155297469 – intereses cuota 1", "NI-00004402", "D", 254610555.55, ""),
    ("BOG-1159383796", "Crédito 1155297469 – intereses cuota 2 (14-jul)", "NI-00004736", "D", 299617500.00, ""),
    ("BOG-L557285005", "Canon abr – capital", "CL-00000391", "D", 62648860.03, ""),
    ("BOG-L557285005", "Canon abr – costo financiero", "NI-00004396", "D", 36783056.97, ""),
    ("BOG-L557285005", "Canon jul – capital", "CL-00000429", "D", 62866117.12, ""),
    ("BOG-L557285005", "Canon jul – costo financiero", "NI-00004740", "D", 39578631.88, ""),
    ("BOG-L556449064", "Canon abr – capital", "CL-00000394", "D", 26079255.17, ""),
    ("BOG-L556449064", "Canon jul – capital", "CL-00000426", "D", 26602533.70, ""),
    ("BOG-L556449064", "Canon jul – costo financiero", "NI-00004737", "D", 9497274.30, ""),
    ("BOG-L556449117", "Canon abr – capital", "CL-00000389", "D", 26079255.16, ""),
    ("BOG-L556449117", "Canon jul – capital", "CL-00000424", "D", 26602533.71, ""),
    ("BOG-L556449117", "Canon jul – costo financiero", "NI-00004734", "D", 9497274.29, ""),
    ("BOG-L556449180", "Canon abr – capital", "CL-00000392", "D", 21315470.38, ""),
    ("BOG-L556449180", "Canon jul – capital", "CL-00000428", "D", 21538319.01, ""),
    ("BOG-L556449180", "Canon jul – costo financiero", "NI-00004739", "D", 14118857.99, ""),
    ("BOG-L556449224", "Canon abr – capital", "CL-00000393", "D", 21315470.39,
     "La descripción del CL-393 indica 'CAPITAL 34.950.056' (valor total del canon) – error de glosa."),
    ("BOG-L556449224", "Canon jul – capital", "CL-00000427", "D", 21538319.02, ""),
    ("BOG-L556449224", "Canon jul – costo financiero", "NI-00004738", "D", 14118857.98, ""),
    ("BOG-L556449288", "Canon abr – capital", "CL-00000390", "D", 21319941.88, ""),
    ("BOG-L556449288", "Canon jul – capital", "CL-00000430", "D", 21542838.16, ""),
    ("BOG-L556449288", "Canon jul – costo financiero", "NI-00004741", "D", 14121779.84, ""),
    ("DAV-L1021275", "Canon 005 (23-abr) – intereses", "NI-00004400", "D", 55427870.00, ""),
    ("DAV-L1021275", "Canon 006 (23-jul) – intereses", "NI-00004744", "D", 64041851.00, ""),
    ("DAV-L1019448", "Canon 030 (5-jun) – capital", "CL-00000415", "D", 8901235.00, ""),
    ("DAV-L1019448", "Canon 031 (6-jul) – capital", "CL-00000433", "D", 8867342.00, ""),
    ("DAV-L1019441", "Canon 032 (11-jun) – capital", "CL-00000416", "D", 0.00,
     "El banco reporta el canon 032 pagado el 11-jun por $1.029.562 sólo intereses (capital 0); contabilidad registró capital $11.973.677. Verificar aplicación del pago."),
    ("DAV-L1019441", "Canon 033 (13-jul) – capital", "CL-00000434", "D", 12327298.00, ""),
]
for ident, conc, doc, campo, vcert, com in PAGOS:
    put(ws, r, 3, ident, font=f_bold)
    put(ws, r, 4, conc, font=f_in, wrap=True)
    put(ws, r, 5, doc, font=f_in)
    put(ws, r, 6, f"=INDEX({R_FEC},MATCH(E{r},{R_DOC},0))", fmt=DATE)
    put(ws, r, 8, vcert, font=f_in, fmt=NUM)
    rng = R_DEB if campo == "D" else R_CRE
    put(ws, r, 9, f"=SUMIFS({rng},{R_DOC},E{r},{R_ID},C{r})", fmt=NUM)
    put(ws, r, 10, f"=I{r}-H{r}", fmt=NUM)
    put(ws, r, 19, com, font=f_in, wrap=True)
    ws.row_dimensions[r].height = 24
    r += 1
put(ws, r, 4, "TOTAL", bold=True, fill=fill_tot)
for c in (8, 9, 10):
    L = get_column_letter(c)
    put(ws, r, c, f"=SUM({L}{P_FIRST}:{L}{r - 1})", fmt=NUM, bold=True, fill=fill_tot)
C_PAGOS_TOT = r
r += 2
section(ws, r, "D. Documentos Davivienda Leasing recibidos por separado (LEASING_1275-3 / 9441-2 / 9448-3_DAVIVIENDA.pdf)", 19)
r += 1
header(ws, r, ["", "Archivo", "Contrato", "ID auditoría", "Fecha corte", "Canon / fecha", "Capital canon", "Intereses canon",
               "Seguro", "Tasa EA cobrada", "Último pago (fecha)", "Último pago (valor)", "¿Informa saldo capital?",
               "Coincide intereses con papel", "Coincide capital con papel", "Coincide tasa con papel", "", "", "Conclusión"], height=40)
r += 1
DOCS_ROWS = []
CERT_WS = ws
DOCS_DAV = [
    ("LEASING_1275-3_DAVIVIENDA.pdf", "000030001021275-3", "DAV-L1021275", dt.date(2026, 7, 6), "006 / 2026-07-23", 0, 64041851, 0, 0.1441,
     dt.date(2026, 4, 23), 55427870),
    ("LEASING_9441-2_DAVIVIENDA.pdf", "000030001019441-2", "DAV-L1019441", dt.date(2026, 6, 25), "033 / 2026-07-13", 12327298, 5062372, 1263646, 0.1603,
     dt.date(2026, 6, 11), 1029562),
    ("LEASING_9448-3_DAVIVIENDA.pdf", "000030001019448-3", "DAV-L1019448", dt.date(2026, 6, 18), "031 / 2026-07-06", 8867342, 4177893, 1263646, 0.1593,
     dt.date(2026, 6, 5), 14567863),
]
for arch, con, ident, fc, can, capc, intc, seg, ea, fup, vup in DOCS_DAV:
    DOCS_ROWS.append((r, ident))
    put(ws, r, 2, arch, font=f_in, wrap=True)
    put(ws, r, 3, con, font=f_in)
    put(ws, r, 4, ident, font=f_bold)
    put(ws, r, 5, fc, font=f_in, fmt=DATE)
    put(ws, r, 6, can, font=f_in)
    put(ws, r, 7, capc, font=f_in, fmt=NUM)
    put(ws, r, 8, intc, font=f_in, fmt=NUM)
    put(ws, r, 9, seg, font=f_in, fmt=NUM)
    put(ws, r, 10, ea, font=f_in, fmt=PCT)
    put(ws, r, 11, fup, font=f_in, fmt=DATE)
    put(ws, r, 12, vup, font=f_in, fmt=NUM)
    put(ws, r, 13, "NO", font=f_in, bold=True, fill=fill_bad)
    put(ws, r, 16, f'=IF(J{r}=M{CERT_ROW[ident]},"Sí","NO")')
    put(ws, r, 19, "Mismo documento que el PDF consolidado. Sólo factura del canon: el saldo de capital sigue sin certificar (PEND-1)."
        + (" Reconfirma canon 032 pagado sólo con intereses $1.029.562 (PEND-7)." if ident == "DAV-L1019441" else ""),
        font=f_in, wrap=True)
    ws.row_dimensions[r].height = 36
    r += 1
ws.freeze_panes = "F5"

# ============================================================================ 6. CRUCE DETALLE VS MAYOR
ws = ws_new("Cruce_Detalle_Mayor", "Cruce detalle 62 vs mayor (cuentas 2105 y 2120) – Integridad y exactitud",
            "Saldos del mayor tomados del balance con signo invertido (pasivo positivo). Movimientos de capital por obligación desde Mayor_Mov.",
            [18, 20, 10, 18, 18, 18, 18, 18, 60])
put(ws, 4, 1, "A. Totales", font=f_sec, border=False)
header(ws, 5, ["Concepto", "Cuentas mayor", "", "Saldo detalle 62", "Saldo mayor 31-jul", "Diferencia (mayor − detalle)", "Estado"])
put(ws, 6, 1, "Créditos (capital)")
put(ws, 6, 2, "21050100 + 21050201")
put(ws, 6, 4, f'=SUMIFS({DR("F")},{DR("D")},"Crédito")', fmt=NUM)
put(ws, 6, 5, f"=-({BAL('21050100','SF')}+{BAL('21050201','SF')})", fmt=NUM)
put(ws, 7, 1, "Leasing financiero (capital)")
put(ws, 7, 2, "21202040 + 21202140")
put(ws, 7, 4, f'=SUMIFS({DR("F")},{DR("D")},"Leasing")', fmt=NUM)
put(ws, 7, 5, f"=-({BAL('21202040','SF')}+{BAL('21202140','SF')})", fmt=NUM)
put(ws, 8, 1, "TOTAL", bold=True, fill=fill_tot)
put(ws, 8, 4, "=D6+D7", fmt=NUM, bold=True, fill=fill_tot)
put(ws, 8, 5, "=E6+E7", fmt=NUM, bold=True, fill=fill_tot)
for rr_ in (6, 7, 8):
    put(ws, rr_, 6, f"=E{rr_}-D{rr_}", fmt=NUM, fill=fill_key if rr_ == 7 else None)
    put(ws, rr_, 7, f'=IF(ABS(F{rr_})<=Resumen!$C$6,"OK","DIFERENCIA")')
put(ws, 9, 1, "Comprobación: total detalle = 'Total Obligaciones' del cliente")
put(ws, 9, 4, f"=Detalle_62!F{D_TOT}-D8", fmt=NUM)
put(ws, 10, 1, "Tarjetas de crédito (21050301) e intereses por pagar (2130) no forman parte del detalle 62.", font=f_sub, border=False)

put(ws, 12, 1, "B. Explicación de la diferencia en leasing (mayor − detalle)", font=f_sec, border=False)
header(ws, 13, ["Partida", "", "", "", "", "Valor", "", "", "Soporte"])
put(ws, 14, 1, "1. Leasing Bancolombia 386735 (camión VW Constellation) registrado en mayor el 28-jul-26 y omitido en el detalle 62", wrap=True)
ws.merge_cells("A14:E14")
put(ws, 14, 6, f'=SUMIFS({R_CRE},{R_ID},"BCL-L386735")-SUMIFS({R_DEB},{R_ID},"BCL-L386735")', fmt=NUM)
put(ws, 14, 9, "NB-00001014: Cr 21202040 $572.363.196 + Cr 21202140 $103.474.304.", wrap=True)
put(ws, 15, 1, "2. Leasings Banco de Bogotá: saldo según extractos mayor que el detalle 62 (detalle desactualizado)", wrap=True)
ws.merge_cells("A15:E15")
put(ws, 15, 6, f"=-{C_BOGL}", fmt=NUM)
put(ws, 15, 9, "Hoja Certificados, sección A (6 contratos).", wrap=True)
put(ws, 16, 1, "3. Diferencia residual no explicada – leasings Davivienda y Bancolombia sin extracto de saldo", wrap=True)
ws.merge_cells("A16:E16")
put(ws, 16, 6, "=F7-F14-F15", fmt=NUM, fill=fill_key)
put(ws, 16, 9, "Requiere certificados de saldo de los contratos 1019441, 1019448, 1021275 (Davivienda) y 362095, 336817, 331330 (Bancolombia).", wrap=True)
put(ws, 17, 1, "Total diferencia leasing", bold=True, fill=fill_tot)
put(ws, 17, 6, "=SUM(F14:F16)", fmt=NUM, bold=True, fill=fill_tot)
put(ws, 17, 7, '=IF(ABS(F17-F7)<0.01,"Cuadra","Revisar")')
for rr_ in (14, 15, 16):
    ws.row_dimensions[rr_].height = 30

put(ws, 19, 1, "C. Roll-forward de capital por obligación (detalle 62 + movimiento ene–jul) → saldo inicial implícito al 1-ene-26", font=f_sec, border=False)
header(ws, 20, ["ID auditoría", "Banco", "Tipo", "Saldo detalle 62 (31-jul)", "Pagos de capital (débitos)",
                "Desembolsos / renovaciones (créditos)", "Saldo inicial implícito 1-ene", "En detalle 62?", "Comentario"])
RF_IDS = [d[2] for d in det] + ["DAV-8906", "BCL-L386735", "N/A-COMISION", "N/A-CRUCE"]
RF_COM = {
    "DAV-8906": "Crédito Davivienda $10.000 MM cancelado el 28-ene-26 con el crédito Bogotá 1155297469 (refinanciación con otro banco: baja en cuentas NIIF 9.3.3.1).",
    "BCL-L386735": "Leasing nuevo 28-jul-26 no incluido en el detalle 62 → saldo inicial implícito negativo = omisión.",
    "N/A-COMISION": "Comisiones Valley Trading registradas y reclasificadas en 21202040 (neto 0). Partida ajena a obligaciones.",
    "N/A-CRUCE": "NI-4731 'cruce anticipo calderas' en 21202140 Davivienda (neto 0). Partida ajena a obligaciones.",
    "BOG-1159383796": "Incluye el crédito 1155297469 (desembolso 28-ene) y su prórroga al 1159383796 (29-jul).",
    "BBVA-0141": "Renovación 30-jun-26 (CE-35090) a la par. El detalle 62 lo identifica como 252059; el mayor como 9600290141/0141 (y 4476 en ene–mar).",
    "BCL-1260105867": "Mayor glosa '12660105867' (error de digitación).",
}
r = 21
RF_FIRST = r
CAP_ACCTS = ("21050100", "21050201", "21202040", "21202140")
for ident in RF_IDS:
    meta = next((d for d in det if d[2] == ident), None)
    bk = meta[0] if meta else {"DAV-8906": "BANCO DAVIVIENDA", "BCL-L386735": "BANCOLOMBIA S.A."}.get(ident, "N/A")
    tp = "Leasing" if ("-L" in ident or ident.startswith("N/A")) else "Crédito"
    put(ws, r, 1, ident, font=f_bold)
    put(ws, r, 2, bk, font=f_in)
    put(ws, r, 3, tp)
    put(ws, r, 4, f"=SUMIFS({DR('F')},{DR('C')},A{r})", fmt=NUM)
    put(ws, r, 5, "=" + "+".join(f'SUMIFS({R_DEB},{R_ID},$A{r},{R_CTA},"{a}")' for a in CAP_ACCTS), fmt=NUM)
    put(ws, r, 6, "=" + "+".join(f'SUMIFS({R_CRE},{R_ID},$A{r},{R_CTA},"{a}")' for a in CAP_ACCTS), fmt=NUM)
    put(ws, r, 7, f"=D{r}+E{r}-F{r}", fmt=NUM)
    put(ws, r, 8, f'=IF(COUNTIF({DR("C")},A{r})>0,"Sí","NO")')
    put(ws, r, 9, RF_COM.get(ident, ""), font=f_in, wrap=True)
    r += 1
RF_LAST = r - 1
for tp, accts_si in (("Crédito", ("21050100", "21050201")), ("Leasing", ("21202040", "21202140"))):
    put(ws, r, 1, f"Total {tp}", bold=True, fill=fill_tot)
    for c in (4, 5, 6, 7):
        L = get_column_letter(c)
        put(ws, r, c, f'=SUMIFS({L}{RF_FIRST}:{L}{RF_LAST},$C${RF_FIRST}:$C${RF_LAST},"{tp}")', fmt=NUM, bold=True, fill=fill_tot)
    r += 1
    put(ws, r, 1, f"Mayor {tp} (SI / Déb / Créd / —)")
    put(ws, r, 7, f"=-({BAL(accts_si[0],'SI')}+{BAL(accts_si[1],'SI')})", fmt=NUM)
    put(ws, r, 5, f"={BAL(accts_si[0],'D')}+{BAL(accts_si[1],'D')}", fmt=NUM)
    put(ws, r, 6, f"={BAL(accts_si[0],'C')}+{BAL(accts_si[1],'C')}", fmt=NUM)
    r += 1
    put(ws, r, 1, f"Diferencia {tp} (detalle implícito − mayor)")
    for c in (5, 6, 7):
        L = get_column_letter(c)
        put(ws, r, c, f"={L}{r - 2}-{L}{r - 1}", fmt=NUM, fill=fill_key if c == 7 else None)
    r += 2
ws.freeze_panes = "B21"

# ============================================================================ 7. RECÁLCULO DE INTERESES
ws = ws_new("Recalculo_Intereses", "Recálculo de intereses (NIIF 9 – costo amortizado, método del interés efectivo)",
            "Interés = Capital × ((1 + EA)^(días/base) − 1). Las fechas y tasas en azul provienen del certificado o del detalle 62 (ver columna 'Fuente').",
            [17, 30, 17, 9, 11, 11, 7, 16, 16, 16, 16, 13, 11, 13, 50])
put(ws, 4, 1, "Parte 1. Recálculo de intereses liquidados por el banco (prueba de exactitud de la liquidación bancaria)", font=f_sec, border=False)
header(ws, 5, ["ID auditoría", "Periodo / concepto", "Capital base", "Tasa EA", "Desde", "Hasta", "Días",
               "Recalc. base 365", "Recalc. base 360", "Interés banco", "Dif. mínima (recalc − banco)",
               "% dif.", "EA implícita (365)", "Conclusión", "Comentario"], height=40)
P1 = [
    ("DAV-830624", "Intereses 21-mar a 21-sep-26 (próximo pago)", "=Certificados!J{}".format(CERT_ROW["DAV-830624"]), 0.1551,
     dt.date(2026, 3, 21), dt.date(2026, 9, 21), 565738188.80, ""),
    ("DAV-607691", "Intereses 28-ene a 28-jul-26", "=Certificados!H{}".format(CERT_ROW["DAV-607691"]), 0.1258,
     dt.date(2026, 1, 28), dt.date(2026, 7, 28), 272299382.05,
     "La tasa 'cobrada' de 12,58% EA del extracto no reconcilia con el interés facturado; la tasa implícita (~14,4% EA) es coherente con IBR 10,69% + 3,40 pts. Confirmar con Davivienda."),
    ("DAV-569032", "Intereses 28-ene a 28-jul-26", "=Certificados!H{}".format(CERT_ROW["DAV-569032"]), 0.1258,
     dt.date(2026, 1, 28), dt.date(2026, 7, 28), 206287213.20, "Igual situación que 607691."),
    ("BOG-856670629", "Cuota 12: 21-mar a 21-jun-26", 1639499992.59, 0.1494, dt.date(2026, 3, 21), dt.date(2026, 6, 21), 59390887.23,
     "Finagro: el banco liquida en base 360."),
    ("BOG-858177259", "Próxima cuota: 25-may a 25-ago-26", "=Certificados!H{}".format(CERT_ROW["BOG-858177259"]), 0.1665,
     dt.date(2026, 5, 25), dt.date(2026, 8, 25), 45166249.86, "Finagro: base 360."),
    ("BOG-453436746", "Cuota 100: 17-jun a 17-jul-26", "=Certificados!H{}".format(CERT_ROW["BOG-453436746"]), 0.1696,
     dt.date(2026, 6, 17), dt.date(2026, 7, 17), 38894999.71, "Ordinaria comercial: base 365."),
    ("BOG-1159383796", "Crédito 1155297469 cuota 2: 14-abr a 14-jul-26", 10000000000, 0.1239, dt.date(2026, 4, 14),
     dt.date(2026, 7, 14), 299617500.00, "Finagro: base 360."),
    ("BCL-1260105867", "Intereses causados 22-jul a 11-ago-26", "=Certificados!H{}".format(CERT_ROW["BCL-1260105867"]), 0.1417,
     dt.date(2026, 7, 22), dt.date(2026, 8, 11), 24607407.00, ""),
    ("BOG-L557285005", "Canon jul: 22-abr a 22-jul-26", "=Certificados!H{}".format(CERT_ROW["BOG-L557285005"]), 0.1604,
     dt.date(2026, 4, 22), dt.date(2026, 7, 22), 39578631.88, ""),
    ("BOG-L556449064", "Canon jul: 27-abr a 27-jul-26", "=Certificados!H{}".format(CERT_ROW["BOG-L556449064"]), 0.1425,
     dt.date(2026, 4, 27), dt.date(2026, 7, 27), 9497274.30, ""),
    ("BOG-L556449117", "Canon jul: 27-abr a 27-jul-26", "=Certificados!H{}".format(CERT_ROW["BOG-L556449117"]), 0.1425,
     dt.date(2026, 4, 27), dt.date(2026, 7, 27), 9497274.29, ""),
    ("BOG-L556449180", "Canon jul: 23-abr a 23-jul-26", "=Certificados!H{}".format(CERT_ROW["BOG-L556449180"]), 0.1422,
     dt.date(2026, 4, 23), dt.date(2026, 7, 23), 14118857.99, ""),
    ("BOG-L556449224", "Canon jul: 23-abr a 23-jul-26", "=Certificados!H{}".format(CERT_ROW["BOG-L556449224"]), 0.1422,
     dt.date(2026, 4, 23), dt.date(2026, 7, 23), 14118857.98, ""),
    ("BOG-L556449288", "Canon jul: 20-abr a 20-jul-26", "=Certificados!H{}".format(CERT_ROW["BOG-L556449288"]), 0.1422,
     dt.date(2026, 4, 20), dt.date(2026, 7, 20), 14121779.84, ""),
]
r = 6
P1_FIRST = r
for ident, per, cap, ea, d0, d1, ib, com in P1:
    put(ws, r, 1, ident, font=f_bold)
    put(ws, r, 2, per, font=f_in, wrap=True)
    put(ws, r, 3, cap, font=f_link if str(cap).startswith("=") else f_in, fmt=NUM)
    put(ws, r, 4, ea, font=f_in, fmt=PCT)
    put(ws, r, 5, d0, font=f_in, fmt=DATE)
    put(ws, r, 6, d1, font=f_in, fmt=DATE)
    put(ws, r, 7, f"=F{r}-E{r}", fmt=NUM0)
    put(ws, r, 8, f"=C{r}*((1+D{r})^(G{r}/365)-1)", fmt=NUM)
    put(ws, r, 9, f"=C{r}*((1+D{r})^(G{r}/360)-1)", fmt=NUM)
    put(ws, r, 10, ib, font=f_in, fmt=NUM)
    put(ws, r, 11, f"=IF(ABS(H{r}-J{r})<=ABS(I{r}-J{r}),H{r}-J{r},I{r}-J{r})", fmt=NUM)
    put(ws, r, 12, f"=K{r}/J{r}", fmt=PCT)
    put(ws, r, 13, f"=(1+J{r}/C{r})^(365/G{r})-1", fmt=PCT)
    put(ws, r, 14, f'=IF(ABS(L{r})<=Resumen!$C$7,"Razonable","Diferencia")')
    put(ws, r, 15, com, font=f_in, wrap=True)
    ws.row_dimensions[r].height = 26
    r += 1
P1_LAST = r - 1
put(ws, r, 2, "TOTAL", bold=True, fill=fill_tot)
for c in (10, 11):
    L = get_column_letter(c)
    put(ws, r, c, f"=SUM({L}{P1_FIRST}:{L}{P1_LAST})", fmt=NUM, bold=True, fill=fill_tot)
put(ws, r, 12, f"=K{r}/J{r}", fmt=PCT, bold=True, fill=fill_tot)
P1_TOT = r
r += 2

put(ws, r, 1, "Parte 2. Causación de intereses al 31-jul-2026 – recálculo vs registrado en 2130 (prueba de corte y valuación)", font=f_sec, border=False)
r += 1
header(ws, r, ["ID auditoría", "Fuente capital / tasa", "Capital 31-jul", "Tasa EA", "Inicio causación (último pago)",
               "Fecha corte", "Días", "Base días", "Interés causado recalculado", "Causado s/mayor (2130)",
               "No registrado (recalc − mayor)", "Cuenta 2130", "Último pago s/mayor", "", "Comentario"], height=40)
r += 1
P2_FIRST = r
cap_cert = lambda i: f"=Certificados!J{CERT_ROW[i]}"
cap_det = lambda i: f"=SUMIFS({DR('F')},{DR('C')},\"{i}\")"
ea_det = lambda i: f"=INDEX({DR('J')},MATCH(\"{i}\",{DR('C')},0))"
P2 = [
    ("DAV-830624", "Certificado", cap_cert("DAV-830624"), 0.1551, dt.date(2026, 3, 21), 365,
     "Interés semestral vencido; el cliente causa cifras fijas mensuales hasta el día ~24 → subcausación."),
    ("DAV-607691", "Cert. / tasa detalle 62", cap_cert("DAV-607691"), ea_det("DAV-607691"), dt.date(2026, 7, 28), 365,
     "3 días del nuevo semestre. Tasa del nuevo periodo no certificada: se usa la del detalle 62."),
    ("DAV-569032", "Cert. / tasa detalle 62", cap_cert("DAV-569032"), ea_det("DAV-569032"), dt.date(2026, 7, 28), 365, "Ídem."),
    ("BBVA-0141", "Detalle 62 (sin certificado)", cap_det("BBVA-0141"), ea_det("BBVA-0141"), dt.date(2026, 6, 30), 365,
     "Renovado 30-jun-26 con intereses mensuales: no hay causación de julio en 2130. Verificar si el interés de julio se pagó y llevó a 53052020; si no, debe causarse."),
    ("BOG-453436746", "Certificado", cap_cert("BOG-453436746"), 0.1696, dt.date(2026, 7, 17), 365,
     "Crédito mensual: el cliente no causa entre cuotas (gasto al pago)."),
    ("BOG-856670629", "Certificado", cap_cert("BOG-856670629"), 0.1494, dt.date(2026, 6, 21), 360, "Finagro, base 360."),
    ("BOG-858177259", "Certificado", cap_cert("BOG-858177259"), 0.1665, dt.date(2026, 5, 25), 360, "Finagro, base 360."),
    ("BOG-1159383796", "Cert. 1155297469 / tasa detalle", cap_cert("BOG-1159383796"), ea_det("BOG-1159383796"),
     dt.date(2026, 7, 22), 360, "Inicio estimado 22-jul (vencimiento 22-ene-27, 6 meses, según detalle 62). Sin certificado del nuevo crédito; confirmar si hubo intereses 14 a 22-jul."),
    ("POP-631309940", "Detalle 62 (sin certificado)", cap_det("POP-631309940"), ea_det("POP-631309940"), dt.date(2026, 6, 17), 365,
     "Inicio estimado 17-jun (cuotas trimestrales al 17; vence 17-dic-26). Pago registrado 30-jun con intereses de mora $54.242."),
    ("BCL-1260105867", "Certificado", cap_cert("BCL-1260105867"), 0.1417, dt.date(2026, 7, 22), 365,
     "El banco certifica $24.607.407 causados del 22-jul al 11-ago (ver Parte 1)."),
    ("DAV-L1021275", "Detalle 62 / tasa extracto", cap_det("DAV-L1021275"), 0.1441, dt.date(2026, 7, 23), 365, ""),
    ("DAV-L1019441", "Detalle 62 / tasa extracto", cap_det("DAV-L1019441"), 0.1603, dt.date(2026, 7, 13), 365, ""),
    ("DAV-L1019448", "Detalle 62 / tasa extracto", cap_det("DAV-L1019448"), 0.1593, dt.date(2026, 7, 6), 365, ""),
    ("BOG-L557285005", "Certificado", cap_cert("BOG-L557285005"), 0.1604, dt.date(2026, 7, 22), 365, ""),
    ("BOG-L556449064", "Certificado", cap_cert("BOG-L556449064"), 0.1425, dt.date(2026, 7, 27), 365, ""),
    ("BOG-L556449117", "Certificado", cap_cert("BOG-L556449117"), 0.1425, dt.date(2026, 7, 27), 365, ""),
    ("BOG-L556449180", "Certificado", cap_cert("BOG-L556449180"), 0.1422, dt.date(2026, 7, 23), 365, ""),
    ("BOG-L556449224", "Certificado", cap_cert("BOG-L556449224"), 0.1422, dt.date(2026, 7, 23), 365, ""),
    ("BOG-L556449288", "Certificado", cap_cert("BOG-L556449288"), 0.1422, dt.date(2026, 7, 20), 365, ""),
    ("BCL-L362095", "Detalle 62 (sin certificado)", cap_det("BCL-L362095"), ea_det("BCL-L362095"), dt.date(2026, 6, 30), 365,
     "Inicio según periodo 'ABR-MAY-JUN' del CL-414. Confirmar fechas de canon con Bancolombia (vencimientos al día 14)."),
    ("BCL-L336817", "Detalle 62 (sin certificado)", cap_det("BCL-L336817"), ea_det("BCL-L336817"), dt.date(2026, 7, 19), 365,
     "Inicio = fecha del último canon registrado (CL-421)."),
    ("BCL-L331330", "Detalle 62 (sin certificado)", cap_det("BCL-L331330"), ea_det("BCL-L331330"), dt.date(2026, 7, 9), 365,
     "Inicio = fecha del último canon registrado (CL-420)."),
    ("BCL-L386735", "Mayor (no está en detalle 62)", f'=SUMIFS({R_CRE},{R_ID},"BCL-L386735")',
     f"=(1+({RATE_CELL['IBR 3 MESES']}+0.0135)/4)^4-1", dt.date(2026, 7, 28), 365, "IBR TV + 1,35 según glosa NB-1014."),
]
for ident, fuente, cap, ea, d0, base, com in P2:
    acct = "21300102" if "-L" in ident else "21300101"
    put(ws, r, 1, ident, font=f_bold)
    put(ws, r, 2, fuente, font=f_in)
    put(ws, r, 3, cap, fmt=NUM)
    put(ws, r, 4, ea, font=f_in if not str(ea).startswith("=") else None, fmt=PCT)
    put(ws, r, 5, d0, font=f_in, fmt=DATE)
    put(ws, r, 6, CORTE, font=f_in, fmt=DATE)
    put(ws, r, 7, f"=F{r}-E{r}", fmt=NUM0)
    put(ws, r, 8, base, font=f_in, fmt=NUM0)
    put(ws, r, 9, f"=C{r}*((1+D{r})^(G{r}/H{r})-1)", fmt=NUM)
    put(ws, r, 12, acct, font=f_in)
    put(ws, r, 13, f'=_xlfn.MAXIFS({R_FEC},{R_ID},A{r},{R_CTA},L{r},{R_DEB},">1000")', fmt='yyyy-mm-dd;;"sin reversión"')
    put(ws, r, 10, f'=SUMIFS({R_CRE},{R_ID},A{r},{R_CTA},L{r},{R_FEC},">"&M{r})-SUMIFS({R_DEB},{R_ID},A{r},{R_CTA},L{r},{R_FEC},">"&M{r})', fmt=NUM)
    put(ws, r, 11, f"=I{r}-J{r}", fmt=NUM, fill=fill_bad)
    put(ws, r, 15, com, font=f_in, wrap=True)
    ws.row_dimensions[r].height = 26
    r += 1
P2_LAST = r - 1
put(ws, r, 2, "TOTAL", bold=True, fill=fill_tot)
for c in (9, 10, 11):
    L = get_column_letter(c)
    put(ws, r, c, f"=SUM({L}{P2_FIRST}:{L}{P2_LAST})", fmt=NUM, bold=True, fill=fill_tot)
P2_TOT = r
r += 1
put(ws, r, 2, "  de los cuales créditos (21300101)")
for c in (9, 10, 11):
    L = get_column_letter(c)
    put(ws, r, c, f'=SUMIFS({L}{P2_FIRST}:{L}{P2_LAST},$L${P2_FIRST}:$L${P2_LAST},"21300101")', fmt=NUM)
P2_OF = r
r += 1
put(ws, r, 2, "  de los cuales leasing (21300102)")
for c in (9, 10, 11):
    L = get_column_letter(c)
    put(ws, r, c, f'=SUMIFS({L}{P2_FIRST}:{L}{P2_LAST},$L${P2_FIRST}:$L${P2_LAST},"21300102")', fmt=NUM)
P2_LS = r
r += 1
put(ws, r, 2, "Saldo 2130 según balance")
put(ws, r, 10, f"=-{BAL('2130','SF')}", fmt=NUM)
r += 1
put(ws, r, 2, "Diferencia causado por obligación vs balance 2130")
put(ws, r, 10, f"=J{P2_TOT}-J{r - 1}", fmt=NUM)
put(ws, r, 11, "Redondeos de reversiones (p.ej. 556449288 $0,16).", font=f_sub, border=False)
r += 2

put(ws, r, 1, "Parte 3. Saldo de capital implícito – leasings Davivienda sin saldo en el extracto", font=f_sec, border=False)
r += 1
header(ws, r, ["ID auditoría", "Canon", "Interés del canon", "Tasa EA", "Desde", "Hasta", "Días", "Capital implícito antes del canon",
               "Capital del canon", "Capital implícito después", "Saldo detalle 62", "Dif. detalle − implícito", "", "", "Comentario"], height=40)
r += 1
P3_FIRST = r
P3 = [
    ("DAV-L1021275", "Canon 006 (sólo intereses)", 64041851, 0.1441, dt.date(2026, 4, 23), dt.date(2026, 7, 23), 0),
    ("DAV-L1019441", "Canon 033", 5062372, 0.1603, dt.date(2026, 6, 11), dt.date(2026, 7, 13), 12327298),
    ("DAV-L1019448", "Canon 031", 4177893, 0.1593, dt.date(2026, 6, 5), dt.date(2026, 7, 6), 8867342),
]
P3_ROW = {}
for ident, can, it, ea, d0, d1, capc in P3:
    P3_ROW[ident] = r
    put(ws, r, 1, ident, font=f_bold)
    put(ws, r, 2, can, font=f_in)
    put(ws, r, 3, it, font=f_in, fmt=NUM)
    put(ws, r, 4, ea, font=f_in, fmt=PCT)
    put(ws, r, 5, d0, font=f_in, fmt=DATE)
    put(ws, r, 6, d1, font=f_in, fmt=DATE)
    put(ws, r, 7, f"=F{r}-E{r}", fmt=NUM0)
    put(ws, r, 8, f"=C{r}/((1+D{r})^(G{r}/365)-1)", fmt=NUM)
    put(ws, r, 9, capc, font=f_in, fmt=NUM)
    put(ws, r, 10, f"=H{r}-I{r}", fmt=NUM)
    put(ws, r, 11, f"=SUMIFS({DR('F')},{DR('C')},A{r})", fmt=NUM)
    put(ws, r, 12, f"=K{r}-J{r}", fmt=NUM)
    put(ws, r, 15, "Estimación indicativa (la tasa puede variar dentro del periodo). Solicitar certificado de saldo.", font=f_in, wrap=True)
    r += 1
put(ws, r, 2, "TOTAL", bold=True, fill=fill_tot)
for c in (10, 11, 12):
    L = get_column_letter(c)
    put(ws, r, c, f"=SUM({L}{P3_FIRST}:{L}{r - 1})", fmt=NUM, bold=True, fill=fill_tot)
P3_TOT = r
for rr_, ident in DOCS_ROWS:  # comparación diferida en Certificados sección D
    p3 = P3_ROW[ident]
    put(CERT_WS, rr_, 14, f'=IF(H{rr_}=Recalculo_Intereses!C{p3},"Sí","NO")')
    put(CERT_WS, rr_, 15, f'=IF(G{rr_}=Recalculo_Intereses!I{p3},"Sí","NO")')
ws.freeze_panes = "C6"

# ============================================================================ 8. GASTO DE INTERESES (prueba analítica)
ws = ws_new("Gasto_Intereses", "Prueba analítica sustantiva – gasto por intereses ene–jul 2026",
            "Expectativa = deuda promedio × ((1 + tasa promedio ponderada)^(días/365) − 1).",
            [58, 20, 60])
rows = [
    ("Deuda financiera 1-ene-26 (2105 sin TC + 2120)", f"=-({BAL('21050100','SI')}+{BAL('21050201','SI')}+{BAL('2120','SI')})", NUM, "Balance de prueba."),
    ("Deuda financiera 31-jul-26 (2105 sin TC + 2120)", f"=-({BAL('21050100','SF')}+{BAL('21050201','SF')}+{BAL('2120','SF')})", NUM, "Balance de prueba."),
    ("Deuda promedio simple", "=AVERAGE(B4:B5)", NUM, ""),
    ("Tasa EA promedio ponderada (detalle 62, tasas del cliente)", f"={D_TPOND}", PCT, "Detalle_62."),
    ("Días del periodo (1-ene a 31-jul-26)", "=DATE(2026,7,31)-DATE(2025,12,31)", NUM0, ""),
    ("Gasto de intereses esperado", "=B6*((1+B7)^(B8/365)-1)", NUM, ""),
    ("Intereses leasing pagados (53052010)", "=" + BAL("53052010", "SF"), NUM, ""),
    ("Intereses préstamos pagados (53052020)", "=" + BAL("53052020", "SF"), NUM, ""),
    ("Intereses préstamos causados (53052030)", "=" + BAL("53052030", "SF"), NUM, ""),
    ("Intereses leasing causados (53052040)", "=" + BAL("53052040", "SF"), NUM, ""),
    ("Gasto de intereses registrado", "=SUM(B10:B13)", NUM, "Excluye mora (53052002) y sobregiros (53052001)."),
    ("Causación no registrada al 31-jul (Recalculo_Intereses Parte 2)", f"=Recalculo_Intereses!K{P2_TOT}", NUM, ""),
    ("Gasto de intereses ajustado", "=B14+B15", NUM, ""),
    ("Diferencia ajustado − esperado", "=B16-B9", NUM, ""),
    ("% diferencia", "=B17/B9", PCT, ""),
    ("Umbral de expectativa (±)", 0.10, PCT, "Supuesto de auditoría: ±10% para una prueba analítica de precisión moderada (tasas variables IBR/DTF)."),
    ("Conclusión", '=IF(ABS(B18)<=B19,"Gasto razonable dentro del umbral","Fuera de umbral – investigar")', None, ""),
    ("Intereses de mora registrados (53052002)", "=" + BAL("53052002", "SF"), NUM, "Incluye $54.242 Banco Popular (CL-411)."),
]
header(ws, 3, ["Concepto", "Valor", "Fuente / comentario"])
for k, (lab, f, fmt, com) in enumerate(rows, start=4):
    put(ws, k, 1, lab, bold=lab.startswith(("Gasto de intereses", "Conclusión")))
    put(ws, k, 2, f, font=f_in if not str(f).startswith("=") else None, fmt=fmt, fill=fill_key if k in (19,) else None)
    put(ws, k, 3, com, font=f_sub)
GI_PCT, GI_CONC = "Gasto_Intereses!$B$18", "Gasto_Intereses!$B$20"

# ============================================================================ 9. CLASIFICACIÓN CP/LP
ws = ws_new("Clasif_CP_LP", "Clasificación corriente / no corriente (NIC 1 párr. 69–76) – detalle 62 vs mayor",
            "CP corregido = MIN(MAX(CP cliente,0); saldo). La porción corriente debe validarse contra las tablas de amortización.",
            [40, 18, 18, 18, 18, 18, 18, 50])
header(ws, 4, ["Concepto", "CP detalle (cliente)", "LP detalle (cliente)", "CP detalle corregido", "LP detalle corregido",
               "Cuenta CP mayor", "Cuenta LP mayor", "Comentario"])
for k, (tp, cp, lp) in enumerate((("Crédito", "21050100", "21050201"), ("Leasing", "21202040", "21202140")), start=5):
    put(ws, k, 1, f"{tp}s")
    for c, col in ((2, "M"), (3, "N"), (4, "W"), (5, "X")):
        put(ws, k, c, f'=SUMIFS({DR(col)},{DR("D")},"{tp}")', fmt=NUM)
    put(ws, k, 6, f"=-{BAL(cp,'SF')}", fmt=NUM)
    put(ws, k, 7, f"=-{BAL(lp,'SF')}", fmt=NUM)
put(ws, 7, 1, "TOTAL", bold=True, fill=fill_tot)
for c in range(2, 8):
    L = get_column_letter(c)
    put(ws, 7, c, f"={L}5+{L}6", fmt=NUM, bold=True, fill=fill_tot)
put(ws, 5, 8, "Las cuentas CP/LP del mayor se asignan por origen del crédito, no por vencimiento.", font=f_in, wrap=True)
put(ws, 6, 8, "El CP del mayor incluye $572 MM del leasing 386735 no incluido en el detalle.", font=f_in, wrap=True)
put(ws, 9, 1, "Reclasificación propuesta LP → CP (detalle corregido − mayor)", font=f_sec, border=False)
header(ws, 10, ["Concepto", "Reclasificación", "", "", "", "", "", "Comentario"])
put(ws, 11, 1, "Créditos: 21050201 → 21050100")
put(ws, 11, 2, "=D5-F5", fmt=NUM, fill=fill_key)
put(ws, 12, 1, "Leasing: 21202140 → 21202040 (sin 386735)")
put(ws, 12, 2, f'=D6-(F6-SUMIFS({R_CRE},{R_ID},"BCL-L386735",{R_CTA},"21202040"))', fmt=NUM, fill=fill_key)
put(ws, 12, 8, "La porción corriente del leasing 386735 ($572 MM registrada en CP) debe validarse con su tabla de amortización.", font=f_in, wrap=True)
put(ws, 13, 1, "Tarjetas de crédito: agrupadas en 210502 (LP) → presentar como corriente")
put(ws, 13, 2, f"=-{BAL('21050301','SF')}", fmt=NUM, fill=fill_key)
put(ws, 15, 1, "Errores de clasificación en el detalle 62", font=f_sec, border=False)
header(ws, 16, ["ID auditoría", "Saldo", "CP cliente", "LP cliente", "CP+LP − saldo", "", "", "Error"])
k = 17
for bank, num, ident, rr in det:
    cp, lp, sal = rr[10] or 0, rr[11] or 0, rr[3]
    if abs(cp + lp - sal) > 1 or cp < 0 or lp < 0 or cp > sal:
        rowd = D_FIRST + [d[2] for d in det].index(ident)
        put(ws, k, 1, ident, font=f_bold)
        put(ws, k, 2, f"=Detalle_62!F{rowd}", fmt=NUM)
        put(ws, k, 3, f"=Detalle_62!M{rowd}", fmt=NUM)
        put(ws, k, 4, f"=Detalle_62!N{rowd}", fmt=NUM)
        put(ws, k, 5, f"=C{k}+D{k}-B{k}", fmt=NUM)
        put(ws, k, 8, "CP mayor que el saldo y LP negativo." if cp > sal else "CP + LP no suma el saldo (LP posiblemente copiado de otra fila).",
            font=f_in, wrap=True)
        k += 1

# ============================================================================ 10. HALLAZGOS
ws = ws_new("Hallazgos", "Hallazgos, referencias NIIF y ajustes propuestos",
            "Montos vinculados a las hojas de trabajo. Ajustes sujetos a la materialidad del encargo (no informada).",
            [5, 16, 70, 18, 26, 55])
header(ws, 4, ["#", "Afirmación", "Hallazgo", "Monto (COP)", "Referencia normativa", "Recomendación / acción", "Estado"])
ws.column_dimensions["G"].width = 24
H_PEND = {2: "PEND-1, PEND-2", 3: "PEND-1, PEND-2", 5: "PEND-3, PEND-5, PEND-10", 6: "PEND-6", 8: "PEND-8",
          11: "PEND-7", 13: "PEND-1 a PEND-5", 15: "PEND-11"}
cdm = "Cruce_Detalle_Mayor"
H = [
    ("Integridad", "El leasing Bancolombia 386735 (camión VW Constellation, IBR TV + 1,35) se registró en el mayor el 28-jul-26 (NB-00001014) pero NO está en el detalle 62 del cliente.",
     f"={cdm}!F14", "NIIF 16 p.26; NIC 1 p.15", "Incluir el contrato en el detalle 62 con su tabla de amortización; solicitar contrato y certificado."),
    ("Exactitud / Integridad", "Diferencia de capital de leasing entre el mayor (2120) y el detalle 62. Se explica por el leasing 386735, por leasings de Bogotá con saldos desactualizados y por un residual sin explicar en leasings de Davivienda/Bancolombia.",
     f"={cdm}!F7", "NIIF 16 p.36; NIC 1 p.15", "Conciliar contrato a contrato contra certificados; actualizar el detalle 62 con saldos bancarios."),
    ("Exactitud", "Residual NO explicado de la diferencia de leasing (contratos sin extracto de saldo: Davivienda 1019441, 1019448, 1021275 y Bancolombia 362095, 336817, 331330).",
     f"={cdm}!F16", "NIA 505 / NIA 500", "Circularizar a Davivienda Leasing y Bancolombia; obtener tablas de amortización."),
    ("Exactitud", "Leasings Banco de Bogotá (pozos 1–5 y planta Jenbacher): el detalle 62 está por debajo del saldo certificado tras el canon de julio (p.ej. pozos 3 y 4 −$57 MM c/u).",
     f"=-{C_BOGL}", "NIIF 16 p.36", "Actualizar el detalle 62 con los saldos certificados."),
    ("Corte / Valuación", "Intereses causados al 31-jul-26 no registrados: el cliente causa cifras fijas hasta ~el día 24 y no causa en obligaciones de pago mensual (453436746, BBVA, leasings mensuales) ni en el periodo posterior al último pago.",
     f"=Recalculo_Intereses!K{P2_TOT}", "NIC 1 p.27-28 (base de acumulación); NIIF 9 p.5.4.1 y B5.4 (método del interés efectivo)", "Registrar AJE-1; implementar causación diaria con tabla de amortización por obligación."),
    ("Exactitud", "Davivienda 607691 y 569032: la tasa 'cobrada' informada (12,58% EA) no reconcilia con los intereses facturados (tasa implícita ~14,4% EA). El detalle 62 usa 16,02% EA.",
     f"=Recalculo_Intereses!K{P1_FIRST + 1}+Recalculo_Intereses!K{P1_FIRST + 2}", "NIIF 9 p.B5.4.5 (tasas variables); NIIF 7 p.39", "Solicitar a Davivienda la liquidación detallada (IBR aplicado y spread)."),
    ("Exactitud (detalle)", "Tasas EA del detalle 62 calculadas con periodicidad distinta a la declarada (BBVA, Bancolombia 1260105867 y 362095) y tasas del detalle distintas a las certificadas (p.ej. 453436746: 17,52% vs 16,96%; 856670629: 15,83% vs 14,94%).",
     None, "NIIF 7 p.33-40 (riesgo de tasa)", "Usar la tasa certificada por el banco y la periodicidad contractual."),
    ("Presentación", "Clasificación CP/LP: el detalle 62 tiene errores (Banco Popular CP $575 MM > saldo $287,5 MM y LP negativo; Pozo #2 CP + LP ≠ saldo) y las cuentas CP/LP del mayor no reflejan vencimientos a 12 meses.",
     "=Clasif_CP_LP!B11+Clasif_CP_LP!B12", "NIC 1 p.69-76", "Reclasificar la porción corriente (AJE-3) y corregir el detalle 62."),
    ("Presentación", "Las tarjetas de crédito por pagar (21050301) están agrupadas dentro de 'Obligaciones financieras a largo plazo' (210502).",
     "=Clasif_CP_LP!B13", "NIC 1 p.69", "Presentar como pasivo corriente."),
    ("Exactitud", "Abonos a capital aplicados por Davivienda por excedentes de pago (127,57 + 269,47 + 627,79 + 617,95 + 786,80) no registrados: el cliente lleva todo el pago a intereses.",
     "=127.57+269.47+627.79+617.95+786.8", "NIIF 9 p.3.3.1", "Inmaterial; registrar en el siguiente pago."),
    ("Ocurrencia / Exactitud", "Leasing Davivienda 1019441: el banco reporta el canon 032 (11-jun) pagado por $1.029.562, sólo intereses; la contabilidad registró capital $11.973.677 (CL-416).",
     11973677, "NIIF 16 p.36", "Obtener el estado de cuenta del contrato y verificar la aplicación del pago."),
    ("Control interno", "Partidas ajenas registradas en cuentas de obligaciones financieras: comisiones Valley Trading en 21202040 y cruce de anticipo de calderas (NI-4731) en 21202140 (netas cero).",
     "=4068134+31360500+8261283", "NIC 1 p.32 (no compensar)", "Restringir el uso de las cuentas 2120 a contratos de leasing con tercero y número de contrato."),
    ("Integridad / Confirmación", "Obligaciones sin certificado ni extracto en el PDF: BBVA ($4.700 MM), Banco Popular, Bogotá 1159383796 (nuevo crédito) y leasings Bancolombia.",
     f"=Certificados!K{C_SINCERT}", "NIA 505", "Enviar confirmaciones bancarias."),
    ("Exactitud (detalle)", "Vencimientos del detalle 62 distintos a los certificados: pozos 3 y 4 (abr-2028 vs 27-jul-2028), pozos 1 y 2 (20-ene vs 23-ene-2030), Bancolombia 1260105867 (21 vs 22-abr-2028).",
     None, "NIIF 7 p.39 (análisis de vencimientos)", "Corregir; afecta la clasificación CP/LP y la revelación de liquidez."),
    ("Valuación / Revelación", "Renovaciones y refinanciaciones (BBVA 30-jun; Bogotá 1155297469 → 1159383796 el 29-jul; Davivienda 8906 pagado con el crédito Bogotá 1155297469): se registraron a la par sin costos de transacción.",
     None, "NIIF 9 p.3.3.2 y B3.3.6 (prueba del 10%)", "Documentar la prueba del 10% (modificación sustancial) y los costos de transacción, si los hubo."),
    ("Control interno", "Causación no sistemática de intereses de leasing (p.ej. pozos 3 y 4: feb–mar $480.943 vs abr $8,6 MM) y glosas con valores errados (CL-393, NI-4397) o números de crédito errados (BBVA 252059 vs 0141; '12660105867').",
     None, "NIC 8 p.41; NIC 1 p.27", "Usar tablas de amortización por obligación; revisar las glosas."),
    ("Integridad", "Mayor vs balance: los débitos, créditos y saldos de las cuentas 21 del movimiento cuadran con los dos balances de prueba (24-ago y 29-sep), que son idénticos.",
     f"=Cruce_Mov_Balance!H{CMB_TOT}", "—", "Sin excepción."),
]
r = 5
H_ROW = {}
for n, (af, txt, monto, ref, rec) in enumerate(H, start=1):
    put(ws, r, 1, n)
    put(ws, r, 2, af, wrap=True)
    put(ws, r, 3, txt, wrap=True)
    put(ws, r, 4, monto if monto is not None else "Cualitativo", fmt=NUM, font=f_in if isinstance(monto, (int, float)) else None)
    put(ws, r, 5, ref, wrap=True)
    put(ws, r, 6, rec, wrap=True)
    put(ws, r, 7, f"Pendiente de soporte ({H_PEND[n]})" if n in H_PEND else "Concluido", wrap=True,
        fill=fill_bad if n in H_PEND else None, bold=True)
    ws.row_dimensions[r].height = 64
    H_ROW[n] = r
    r += 1
r += 1
section(ws, r, "Ajustes propuestos (AJE)", 6)
r += 1
header(ws, r, ["AJE", "Cuenta", "Descripción", "Débito", "Crédito", "Soporte"])
r += 1
AJE = [
    ("AJE-1", "53052030", "Intereses sobre préstamos causados", f"=Recalculo_Intereses!K{P2_OF}", None, "Causación de intereses de créditos al 31-jul-26"),
    ("AJE-1", "21300101", "Intereses financieros por pagar O.F.", None, f"=Recalculo_Intereses!K{P2_OF}", ""),
    ("AJE-1", "53052040", "Intereses leasing causados", f"=Recalculo_Intereses!K{P2_LS}", None, "Causación de intereses de leasing al 31-jul-26"),
    ("AJE-1", "21300102", "Intereses leasing por pagar", None, f"=Recalculo_Intereses!K{P2_LS}", ""),
    ("AJE-2", "21050201", "Obligación financiera LP", "=Clasif_CP_LP!B11", None, "Reclasificación de la porción corriente (presentación)"),
    ("AJE-2", "21050100", "Obligación financiera CP", None, "=Clasif_CP_LP!B11", ""),
    ("AJE-2", "21202140", "Leasing LP", "=Clasif_CP_LP!B12", None, ""),
    ("AJE-2", "21202040", "Leasing CP", None, "=Clasif_CP_LP!B12", ""),
]
A_FIRST = r
for aje, cta, des, db, cr, sop in AJE:
    put(ws, r, 1, aje)
    put(ws, r, 2, cta, font=f_in)
    put(ws, r, 3, des)
    put(ws, r, 4, db, fmt=NUM)
    put(ws, r, 5, cr, fmt=NUM)
    put(ws, r, 6, sop, wrap=True)
    r += 1
put(ws, r, 3, "Sumas iguales", bold=True, fill=fill_tot)
put(ws, r, 4, f"=SUM(D{A_FIRST}:D{r - 1})", fmt=NUM, bold=True, fill=fill_tot)
put(ws, r, 5, f"=SUM(E{A_FIRST}:E{r - 1})", fmt=NUM, bold=True, fill=fill_tot)
r += 1
put(ws, r, 3, "AJE-2 depende de validar la porción corriente del detalle 62 con tablas de amortización; el saldo del leasing (diferencia de $1.018 MM) se ajustará cuando se obtengan los certificados faltantes.",
    font=f_sub, border=False)

# ============================================================================ 10b. PENDIENTES
ws = ws_new("Pendientes", "Pendientes de soporte – el papel se emite sin esta información",
            "Las cifras afectadas se dejan calculadas con la información disponible y supuestos documentados; se actualizarán al recibir el soporte.",
            [9, 62, 22, 18, 16, 22, 14, 50])
header(ws, 4, ["Ref.", "Pendiente / información requerida", "Obligación (ID)", "Monto afectado (COP)", "Hallazgo relacionado",
               "Solicitar a", "Estado", "Tratamiento provisional en este papel"], height=36)
PEND = [
    ("PEND-1", "Certificados de saldo de capital al 31-jul-26 de los leasings Davivienda 1019441-2, 1019448-3 y 1021275-3 (los extractos sólo traen la factura del canon).",
     "DAV-L1019441 / L1019448 / L1021275", f"=Recalculo_Intereses!K{P3_TOT}", "2, 3, 13", "Davivienda Leasing",
     "Se usa el saldo del detalle 62; saldo implícito estimado en Recalculo_Intereses Parte 3. Recibidos LEASING_1275-3/9441-2/9448-3_DAVIVIENDA.pdf: son facturas de canon sin saldo de capital (Certificados sección D).",
     "Parcial – recibido sin saldo"),
    ("PEND-2", "Certificados de saldo de los leasings Bancolombia 362095, 336817, 331330 y contrato + tabla de amortización del nuevo leasing 386735.",
     "BCL-L362095 / L336817 / L331330 / L386735", f"={cdm}!F16", "1, 2, 3, 13", "Bancolombia",
     "Residual de leasing sin explicar queda abierto; 386735 se toma del mayor (NB-1014)."),
    ("PEND-3", "Confirmación BBVA del crédito renovado 30-jun-26 ($4.700 MM): número de obligación, tasa y si el interés de julio se pagó (y a qué cuenta se llevó).",
     "BBVA-0141", f'=SUMIFS(Recalculo_Intereses!K{P2_FIRST}:K{P2_LAST},Recalculo_Intereses!A{P2_FIRST}:A{P2_LAST},"BBVA-0141")',
     "5, 13", "BBVA / cliente", "Se incluye en AJE-1 la causación estimada de julio; se retira si se demuestra el pago en 53052020."),
    ("PEND-4", "Confirmación Banco Popular del crédito 631309940-2 (saldo, tasa, fechas de cuota).",
     "POP-631309940", f'=SUMIFS({DR("F")},{DR("C")},"POP-631309940")', "13", "Banco Popular",
     "Saldo del detalle 62; inicio de causación estimado 17-jun."),
    ("PEND-5", "Certificado del nuevo crédito Bogotá 1159383796 (prórroga del 1155297469): fecha de inicio, tasa y si hubo intereses entre el 14 y el 22-jul.",
     "BOG-1159383796", f'=SUMIFS(Recalculo_Intereses!K{P2_FIRST}:K{P2_LAST},Recalculo_Intereses!A{P2_FIRST}:A{P2_LAST},"BOG-1159383796")',
     "5, 13", "Banco de Bogotá", "Inicio estimado 22-jul con tasa del detalle 62 (IBR 3M + 0,8)."),
    ("PEND-6", "Liquidación detallada Davivienda créditos 607691 y 569032 (IBR aplicado y spread): la tasa declarada 12,58% EA no reconcilia con lo facturado.",
     "DAV-607691 / DAV-569032", f"=Recalculo_Intereses!K{P1_FIRST + 1}+Recalculo_Intereses!K{P1_FIRST + 2}", "6", "Davivienda",
     "Se acepta lo facturado (coincide con lo contabilizado); diferencia de tasa queda abierta."),
    ("PEND-7", "Estado de cuenta del leasing Davivienda 1019441: aplicación del pago del canon 032 (banco: sólo intereses $1.029.562; contabilidad: capital $11.973.677).",
     "DAV-L1019441", 11973677, "11", "Cliente", "Sin ajuste hasta verificar. LEASING_9441-2_DAVIVIENDA.pdf reconfirma el lado del banco; falta la explicación del cliente del registro CL-416.",
     "Parcial – banco reconfirma"),
    ("PEND-8", "Tablas de amortización por obligación para determinar la porción corriente a 12 meses (NIC 1.69).",
     "Todas", "=Clasif_CP_LP!B11+Clasif_CP_LP!B12", "8", "Cliente", "AJE-2 calculado con el CP del detalle 62 corregido; sujeto a validación."),
    ("PEND-9", "Materialidad del encargo (global, de ejecución y umbral de errores triviales).",
     "—", None, "Todos", "Socio / gerente del encargo", "Umbral de redondeo $1.000 y tolerancia de recálculo 2% (hoja Resumen, celdas amarillas)."),
    ("PEND-10", "Fechas reales de último pago/canon de Banco Popular y leasings Bancolombia (362095, 336817, 331330) para la causación al corte.",
     "POP / BCL-L*", f'=SUMIFS(Recalculo_Intereses!K{P2_FIRST}:K{P2_LAST},Recalculo_Intereses!A{P2_FIRST}:A{P2_LAST},"BCL-L*")+SUMIFS(Recalculo_Intereses!K{P2_FIRST}:K{P2_LAST},Recalculo_Intereses!A{P2_FIRST}:A{P2_LAST},"POP-*")',
     "5", "Cliente / bancos", "Fechas estimadas (anotadas en Recalculo_Intereses Parte 2)."),
    ("PEND-11", "Documentación de la prueba del 10% (NIIF 9 B3.3.6) y costos de transacción de las renovaciones BBVA y Bogotá y de la refinanciación Davivienda 8906.",
     "BBVA-0141 / BOG-1159383796 / DAV-8906", None, "15", "Cliente", "Se asume modificación no sustancial a la par, sin costos."),
]
r = 5
PN_FIRST = r
for ref, txt, ident, monto, hz, quien, trat, *est in PEND:
    est = est[0] if est else "Pendiente"
    put(ws, r, 1, ref, bold=True)
    put(ws, r, 2, txt, wrap=True)
    put(ws, r, 3, ident, wrap=True)
    put(ws, r, 4, monto if monto is not None else "Cualitativo", fmt=NUM, font=f_in if isinstance(monto, (int, float)) else None)
    put(ws, r, 5, hz)
    put(ws, r, 6, quien, wrap=True)
    put(ws, r, 7, est, bold=True, fill=fill_bad, wrap=True)
    put(ws, r, 8, trat, wrap=True)
    ws.row_dimensions[r].height = 52
    r += 1
PN_LAST = r - 1
r += 1
put(ws, r, 2, "Pendientes abiertos", bold=True, fill=fill_tot)
put(ws, r, 4, f'=COUNTIF(G{PN_FIRST}:G{PN_LAST},"<>Recibido")', fmt=NUM0, bold=True, fill=fill_tot)
PN_COUNT = f"Pendientes!$D${r}"
r += 1
put(ws, r, 2, "Instrucción: al recibir un soporte suficiente, cambiar el Estado a 'Recibido' (se descuenta del contador), actualizar los datos de entrada (azul) en la hoja correspondiente y recalcular.",
    font=f_sub, border=False)
ws.freeze_panes = "B5"

# ============================================================================ 11. RESUMEN
ws = wb.create_sheet("Resumen", 0)
ws.sheet_view.showGridLines = False
for i, w in enumerate([4, 62, 20, 20, 20, 16], start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws["A1"] = "GUAICARAMO S.A.S. – Auditoría de Obligaciones Financieras (cuenta 21 / papel 62) – corte 31-jul-2026"
ws["A1"].font = f_title
ws["A2"] = ("Cruce movimiento ↔ detalle 62 ↔ certificados bancarios ↔ balance de prueba; recálculo de intereses (NIIF 9 costo amortizado, NIIF 16, NIC 1). "
            "Cifras en pesos colombianos. Azul = dato de la fuente; negro = fórmula; verde = vínculo a otra hoja; amarillo = cifra clave.")
ws["A2"].font = f_sub
put(ws, 4, 2, "Parámetros", font=f_sec, border=False)
put(ws, 5, 2, "Fecha de corte")
put(ws, 5, 3, CORTE, font=f_in, fmt=DATE)
put(ws, 6, 2, "Umbral de diferencia trivial / redondeo (COP)")
put(ws, 6, 3, 1000, font=f_in, fmt=NUM, fill=fill_key)
note(ws, 6, 3, "Supuesto de auditoría: diferencias de hasta $1.000 se consideran redondeo. Ajustar según la materialidad del encargo.")
put(ws, 7, 2, "Tolerancia recálculo de intereses (%)")
put(ws, 7, 3, 0.02, font=f_in, fmt=PCT, fill=fill_key)
put(ws, 4, 5, "Estado del papel", font=f_sec, border=False)
put(ws, 5, 5, '="Emitido con " & ' + PN_COUNT + ' & " pendientes"', bold=True, fill=fill_bad)
ws.merge_cells("E5:F5")
put(ws, 6, 5, "Ver hoja 'Pendientes'", font=f_sub, border=False)
note(ws, 7, 3, "Supuesto de auditoría: ±2% por diferencias de convención (días/base, fecha de fijación del IBR/DTF).")

put(ws, 9, 2, "Resultados por prueba", font=f_sec, border=False)
header(ws, 10, ["", "Prueba", "Valor A", "Valor B", "Diferencia", "Estado"], c0=1)
TESTS = [
    ("1. Integridad: movimiento auxiliar vs balance – SF cuenta 21",
     f"=Cruce_Mov_Balance!F{CMB_TOT}", f"=Cruce_Mov_Balance!G{CMB_TOT}"),
    ("2. Integridad: balance 29-sep vs balance 24-ago – nº cuentas con diferencias", "=" + BAL_NDIFF, "=0"),
    ("3. Detalle 62: suma recalculada vs 'Total obligaciones' del cliente", f"=Detalle_62!F{D_TOT}", f"=Detalle_62!F{D_TOT + 1}"),
    ("4. Créditos: detalle 62 vs mayor (21050100 + 21050201)", f"={cdm}!D6", f"={cdm}!E6"),
    ("5. Leasing: detalle 62 vs mayor (21202040 + 21202140)", f"={cdm}!D7", f"={cdm}!E7"),
    ("6. Certificados: capital banco 31-jul vs detalle 62 (obligaciones con saldo certificado)",
     f'=SUMIFS(Certificados!J{C_FIRST}:J{C_LAST},Certificados!H{C_FIRST}:H{C_LAST},"<>")',
     f'=SUMIFS(Certificados!K{C_FIRST}:K{C_LAST},Certificados!H{C_FIRST}:H{C_LAST},"<>")'),
    ("7. Certificados: pagos/intereses certificados vs registrados (comprobantes)",
     f"=Certificados!H{C_PAGOS_TOT}", f"=Certificados!I{C_PAGOS_TOT}"),
    ("8. Recálculo de intereses liquidados por el banco (Parte 1)",
     f"=Recalculo_Intereses!J{P1_TOT}+Recalculo_Intereses!K{P1_TOT}", f"=Recalculo_Intereses!J{P1_TOT}"),
    ("9. Intereses causados al 31-jul: recalculado vs registrado en 2130",
     f"=Recalculo_Intereses!I{P2_TOT}", f"=Recalculo_Intereses!J{P2_TOT}"),
    ("10. Intereses por pagar: suma por obligación vs balance 2130",
     f"=Recalculo_Intereses!J{P2_TOT}", f"=-{BAL('2130','SF')}"),
]
r = 11
for lab, a_, b_ in TESTS:
    put(ws, r, 2, lab, wrap=True)
    put(ws, r, 3, a_, fmt=NUM)
    put(ws, r, 4, b_, fmt=NUM)
    put(ws, r, 5, f"=C{r}-D{r}", fmt=NUM)
    put(ws, r, 6, f'=IF(ABS(E{r})<=$C$6,"OK","DIFERENCIA")', bold=True)
    ws.row_dimensions[r].height = 26
    r += 1
r += 1
put(ws, r, 2, "Indicadores clave", font=f_sec, border=False)
r += 1
KPIS = [
    ("Saldo de capital según mayor 31-jul (2105 sin TC + 2120)", f"=Gasto_Intereses!B5", NUM),
    ("Saldo según detalle 62", f"=Detalle_62!F{D_TOT}", NUM),
    ("Leasing 386735 omitido en el detalle 62", f"={cdm}!F14", NUM),
    ("Leasings Bogotá: detalle por debajo de los certificados", f"={cdm}!F15", NUM),
    ("Diferencia de leasing sin explicar (requiere certificados)", f"={cdm}!F16", NUM),
    ("% del saldo del detalle 62 cubierto con certificados", f"={C_COBERT}", PCT),
    ("Intereses no causados al 31-jul (AJE-1)", f"=Recalculo_Intereses!K{P2_TOT}", NUM),
    ("Reclasificación LP → CP propuesta (créditos + leasing)", "=Clasif_CP_LP!B11+Clasif_CP_LP!B12", NUM),
    ("Gasto de intereses: desviación vs expectativa", f"={GI_PCT}", PCT),
    ("Conclusión prueba analítica", f"={GI_CONC}", None),
]
for lab, f, fmt in KPIS:
    put(ws, r, 2, lab)
    put(ws, r, 3, f, fmt=fmt, fill=fill_key)
    r += 1
r += 1
put(ws, r, 2, "Conclusión", font=f_sec, border=False)
r += 1
concl = [
    "INTEGRIDAD – Mayor vs balance: el movimiento auxiliar de la cuenta 21 cuadra exactamente (débitos, créditos y saldos) con los dos balances de prueba, que son idénticos entre sí. Sin excepción.",
    "INTEGRIDAD – Detalle 62 vs mayor: los créditos cuadran (diferencia de redondeo $0,67). El leasing NO cuadra: el mayor supera al detalle; la principal causa es el leasing Bancolombia 386735 omitido en el detalle.",
    "EXACTITUD – Certificados: los créditos cuadran con los extractos salvo redondeos de abonos a capital (< $5.000). Los leasings de Banco de Bogotá están subvalorados en el detalle 62. Los pagos registrados coinciden con los extractos, salvo el canon 032 del leasing 1019441.",
    "RECÁLCULO – La liquidación bancaria es razonable en base 365 (Bogotá Finagro en base 360), excepto Davivienda 607691/569032, cuya tasa declarada no reconcilia. Hay intereses causados no registrados al 31-jul (AJE-1) por la política de causación del cliente.",
    "PRESENTACIÓN – Se requiere reclasificar la porción corriente (NIC 1.69) y corregir los errores CP/LP del detalle 62. Ver la hoja 'Hallazgos'.",
    "PENDIENTES – El papel se emite sin certificados de saldo de leasings Davivienda/Bancolombia, sin confirmaciones de BBVA, Banco Popular y del crédito Bogotá 1159383796, sin tablas de amortización y sin materialidad del encargo. Las conclusiones afectadas quedan 'Pendiente de soporte' (ver hoja 'Pendientes').",
]
for t in concl:
    put(ws, r, 2, t, wrap=True)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
    ws.row_dimensions[r].height = 40
    r += 1
r += 1
put(ws, r, 2, "Hojas del libro", font=f_sec, border=False)
r += 1
for nm, ds in [("Mayor_Mov", "Movimiento auxiliar ene–jul con el ID de obligación asignado por auditoría"),
               ("Balance_21", "Cuentas 21/5305 de ambos balances de prueba y su comparación"),
               ("Cruce_Mov_Balance", "Integridad: movimiento vs balance, jerarquía de cuentas y relación pasivo–gasto"),
               ("Detalle_62", "Detalle del cliente con recálculo de tasas, subtotales y CP/LP"),
               ("Cruce_Detalle_Mayor", "Detalle 62 vs mayor y roll-forward de capital por obligación"),
               ("Certificados", "Detalle 62 vs extractos bancarios del PDF y pagos certificados vs contabilidad"),
               ("Recalculo_Intereses", "Recálculo de la liquidación bancaria, causación al corte y saldos implícitos"),
               ("Gasto_Intereses", "Prueba analítica del gasto por intereses"),
               ("Clasif_CP_LP", "Clasificación corriente/no corriente (NIC 1)"),
               ("Hallazgos", "Hallazgos, normas y ajustes propuestos"),
               ("Pendientes", "Información pendiente de soporte y tratamiento provisional")]:
    put(ws, r, 2, nm, bold=True)
    put(ws, r, 3, ds)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
    r += 1

for w in wb.worksheets:
    w.sheet_properties.tabColor = {"Resumen": "1F3864", "Hallazgos": "C00000"}.get(w.title, "8EA9DB")
    w.page_setup.orientation = "landscape"
    w.sheet_properties.pageSetUpPr.fitToPage = True
    w.page_setup.fitToWidth = 1
    w.page_setup.fitToHeight = 0
wb.save(OUT)
print("OK", OUT, "mov rows", len(mov_rows), "det rows", len(det))
