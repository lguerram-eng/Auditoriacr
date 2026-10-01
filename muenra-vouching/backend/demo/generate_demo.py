"""Genera datos FICTICIOS de demostración para Muenra Vouching.

Todas las empresas, NIT, números y valores son inventados. No se incluyen
documentos reales ni datos personales.

Uso:  python -m demo.generate_demo [directorio_salida]
Produce:
  <salida>/referencia_vouching_demo.xlsx   (6 hojas)
  <salida>/documentos/*.pdf|xml|png|tiff|jpg|docx
"""
from __future__ import annotations

import hashlib
import io
import random
import sys
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from PIL import Image, ImageDraw, ImageFilter, ImageFont

from app.services.normalization import compute_dv


def nit(base: str) -> str:
    """NIT ficticio con dígito de verificación válido, formato 900.123.456-7."""
    b = f"{int(base):,}".replace(",", ".")
    return f"{b}-{compute_dv(base)}"


CLIENT = ("Cliente Demostración S.A.S.", "900100200")


@dataclass(frozen=True)
class Company:
    name: str
    base: str

    @property
    def nit(self) -> str:
        return nit(self.base)


A = Company("Proveedor Andino S.A.S.", "900123456")
B = Company("Servicios Logísticos del Caribe Ltda.", "800234567")
C = Company("Consultores Boreal SAS", "901345678")
D = Company("Tecnología Cumbre S.A.", "860456789")
E = Company("Papelería La Estrella E.U.", "830567891")
F = Company("Distribuidora Nevado SAS", "900678912")
G = Company("Ingeniería Páramo S.A.S.", "901789123")
H = Company("Transportes Altiplano Ltda", "800891234")
I = Company("Inversiones Ficticias Ltda", "890999111")
J = Company("Montajes Sabana S.A.S.", "900912345")


def money(v: int) -> str:
    return f"$ {v:,.0f}".replace(",", ".")


# ---------------------------------------------------------------------------
# Renderizado
# ---------------------------------------------------------------------------


def _font(size: int, bold: bool = False):
    candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    ]
    for c in candidates:
        if Path(c).exists():
            return ImageFont.truetype(c, size)
    return ImageFont.load_default(size=size)


Lines = list  # [(texto, x, negrita)] o [(texto, x, negrita, x_valor, valor)]


def render_pdf(lines: Lines) -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter, invariant=1)
    c.setTitle("Documento ficticio de demostración")
    y = 740
    for item in lines:
        text, x, bold = item[0], item[1], item[2]
        c.setFont("Helvetica-Bold" if bold else "Helvetica", 11 if not bold else 13)
        if text:
            c.drawString(x, y, text)
        if len(item) > 3:
            c.drawString(item[3], y, item[4])
        y -= 20
        if y < 60:
            c.showPage()
            y = 740
    c.showPage()
    c.save()
    return buf.getvalue()


def render_image(lines: Lines, noise: bool = True, seed: int = 1) -> Image.Image:
    W, H = 1700, 2200
    img = Image.new("L", (W, H), 255)
    d = ImageDraw.Draw(img)
    y = 120
    for item in lines:
        text, x, bold = item[0], item[1], item[2]
        f = _font(40 if bold else 34, bold)
        if text:
            d.text((int(x * 2.6), y), text, fill=0, font=f)
        if len(item) > 3:
            d.text((int(item[3] * 2.6), y), item[4], fill=0, font=f)
        y += 62
    if noise:
        rnd = random.Random(seed)
        for _ in range(1500):
            img.putpixel((rnd.randrange(W), rnd.randrange(H)), rnd.randrange(150, 230))
        img = img.filter(ImageFilter.GaussianBlur(0.6))
    return img


def to_png(img: Image.Image) -> bytes:
    b = io.BytesIO()
    img.save(b, "PNG")
    return b.getvalue()


def to_tiff(img: Image.Image) -> bytes:
    b = io.BytesIO()
    img.save(b, "TIFF", compression="tiff_deflate")
    return b.getvalue()


def to_scanned_pdf(img: Image.Image) -> bytes:
    b = io.BytesIO()
    img.convert("RGB").save(b, "PDF", resolution=200)
    return b.getvalue()


# ---------------------------------------------------------------------------
# Contenido de documentos
# ---------------------------------------------------------------------------


def invoice(number: str, issuer: Company, issued: date, total: int, *, title="FACTURA DE VENTA", issuer_label=None,
            due: date | None = None, oc: str | None = None, contract: str | None = None, concept="Servicios profesionales",
            cufe: str | None = None, iva_rate: float = 0.19) -> Lines:
    sub = round(total / (1 + iva_rate))
    iva = total - sub
    lines: Lines = [
        (issuer_label or issuer.name.upper(), 50, True),
        (f"NIT: {issuer.nit}", 50, False),
        ("Régimen común - Responsable de IVA", 50, False),
        ("", 50, False),
        (f"{title} No. {number}", 50, True),
        (f"Fecha de emisión: {issued.strftime('%d/%m/%Y')}", 50, False),
    ]
    if due:
        lines.append((f"Fecha de vencimiento: {due.strftime('%d/%m/%Y')}", 50, False))
    lines += [
        (f"Cliente: {CLIENT[0]}", 50, False),
        (f"NIT cliente: {nit(CLIENT[1])}", 50, False),
        (f"Concepto: {concept}", 50, False),
    ]
    if oc:
        lines.append((f"Orden de compra No. {oc}", 50, False))
    if contract:
        lines.append((f"Contrato No. {contract}", 50, False))
    lines += [
        ("Forma de pago: Crédito 30 días", 50, False),
        ("", 50, False),
        ("Subtotal", 50, False, 380, money(sub)),
        (f"IVA {int(iva_rate * 100)}%", 50, False, 380, money(iva)),
        ("TOTAL A PAGAR", 50, True, 380, money(total)),
        ("Moneda: COP", 50, False),
    ]
    if cufe:
        lines += [("CUFE:", 50, False), (cufe[:48], 50, False), (cufe[48:], 50, False)]
    lines.append(("Documento ficticio generado para demostración de Muenra Vouching", 50, False))
    return lines


def cufe_for(number: str) -> str:
    return hashlib.sha384(f"cufe-demo-{number}".encode()).hexdigest()


def ubl_invoice(number: str, issuer: Company, issued: date, total: int, cufe: str, attached: bool = False, concept="Servicio de transporte de carga") -> bytes:
    sub = round(total / 1.19)
    iva = total - sub
    inv = f"""<?xml version="1.0" encoding="UTF-8"?>
<Invoice xmlns="urn:oasis:names:specification:ubl:schema:xsd:Invoice-2"
 xmlns:cac="urn:oasis:names:specification:ubl:schema:xsd:CommonAggregateComponents-2"
 xmlns:cbc="urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2">
  <cbc:UBLVersionID>UBL 2.1</cbc:UBLVersionID>
  <cbc:ID>{number}</cbc:ID>
  <cbc:UUID schemeName="CUFE-SHA384">{cufe}</cbc:UUID>
  <cbc:IssueDate>{issued.isoformat()}</cbc:IssueDate>
  <cbc:DueDate>{issued.replace(day=min(issued.day + 20, 28)).isoformat()}</cbc:DueDate>
  <cbc:InvoiceTypeCode>01</cbc:InvoiceTypeCode>
  <cbc:Note>{concept}</cbc:Note>
  <cbc:DocumentCurrencyCode>COP</cbc:DocumentCurrencyCode>
  <cac:OrderReference><cbc:ID>OC-5500</cbc:ID></cac:OrderReference>
  <cac:AccountingSupplierParty><cac:Party><cac:PartyTaxScheme>
    <cbc:RegistrationName>{issuer.name}</cbc:RegistrationName>
    <cbc:CompanyID schemeID="{compute_dv(issuer.base)}" schemeName="31">{issuer.base}</cbc:CompanyID>
  </cac:PartyTaxScheme></cac:Party></cac:AccountingSupplierParty>
  <cac:AccountingCustomerParty><cac:Party><cac:PartyTaxScheme>
    <cbc:RegistrationName>{CLIENT[0]}</cbc:RegistrationName>
    <cbc:CompanyID schemeID="{compute_dv(CLIENT[1])}" schemeName="31">{CLIENT[1]}</cbc:CompanyID>
  </cac:PartyTaxScheme></cac:Party></cac:AccountingCustomerParty>
  <cac:PaymentMeans><cbc:ID>2</cbc:ID></cac:PaymentMeans>
  <cac:TaxTotal><cbc:TaxAmount currencyID="COP">{iva}.00</cbc:TaxAmount>
    <cac:TaxSubtotal><cbc:TaxableAmount currencyID="COP">{sub}.00</cbc:TaxableAmount><cbc:TaxAmount currencyID="COP">{iva}.00</cbc:TaxAmount>
      <cac:TaxCategory><cbc:Percent>19.00</cbc:Percent><cac:TaxScheme><cbc:ID>01</cbc:ID><cbc:Name>IVA</cbc:Name></cac:TaxScheme></cac:TaxCategory>
    </cac:TaxSubtotal></cac:TaxTotal>
  <cac:LegalMonetaryTotal>
    <cbc:LineExtensionAmount currencyID="COP">{sub}.00</cbc:LineExtensionAmount>
    <cbc:TaxExclusiveAmount currencyID="COP">{sub}.00</cbc:TaxExclusiveAmount>
    <cbc:TaxInclusiveAmount currencyID="COP">{total}.00</cbc:TaxInclusiveAmount>
    <cbc:AllowanceTotalAmount currencyID="COP">0.00</cbc:AllowanceTotalAmount>
    <cbc:PayableAmount currencyID="COP">{total}.00</cbc:PayableAmount>
  </cac:LegalMonetaryTotal>
  <cac:InvoiceLine><cbc:ID>1</cbc:ID><cbc:InvoicedQuantity unitCode="EA">2</cbc:InvoicedQuantity>
    <cbc:LineExtensionAmount currencyID="COP">{sub}.00</cbc:LineExtensionAmount>
    <cac:Item><cbc:Description>{concept}</cbc:Description></cac:Item>
    <cac:Price><cbc:PriceAmount currencyID="COP">{sub / 2:.2f}</cbc:PriceAmount></cac:Price></cac:InvoiceLine>
</Invoice>"""
    if not attached:
        return inv.encode("utf-8")
    body = inv.split("\n", 1)[1]
    return f"""<?xml version="1.0" encoding="UTF-8"?>
<AttachedDocument xmlns="urn:oasis:names:specification:ubl:schema:xsd:AttachedDocument-2"
 xmlns:cac="urn:oasis:names:specification:ubl:schema:xsd:CommonAggregateComponents-2"
 xmlns:cbc="urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2">
  <cbc:ID>AD-{number}</cbc:ID>
  <cac:Attachment><cac:ExternalReference><cbc:MimeCode>text/xml</cbc:MimeCode>
    <cbc:Description><![CDATA[{body}]]></cbc:Description>
  </cac:ExternalReference></cac:Attachment>
</AttachedDocument>""".encode("utf-8")


def contract_docx(number: str, contractor: Company, value: int) -> bytes:
    import docx

    d = docx.Document()
    d.add_heading(f"CONTRATO DE PRESTACIÓN DE SERVICIOS No. {number}", level=1)
    d.add_paragraph(f"Contratante: {CLIENT[0]} NIT {nit(CLIENT[1])}")
    d.add_paragraph(f"Contratista: {contractor.name} NIT {contractor.nit}")
    d.add_paragraph("Objeto: Montaje y puesta en marcha de estanterías industriales en la bodega principal.")
    d.add_paragraph(f"Valor del contrato: {money(value)}")
    d.add_paragraph("Vigencia: seis (6) meses contados a partir de la firma del acta de inicio.")
    d.add_paragraph("Fecha de suscripción: 15/01/2026")
    d.add_paragraph("Forma de pago: 100% contra entrega a satisfacción.")
    t = d.add_table(rows=3, cols=3)
    for i, row in enumerate([("Ítem", "Cantidad", "Valor"), ("Estanterías", "20", money(value - 500000)), ("Instalación", "1", money(500000))]):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    d.add_paragraph("Firmado por: Representante Legal Contratante (ficticio)")
    d.add_paragraph("Firmado por: Representante Legal Contratista (ficticio)")
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


# ---------------------------------------------------------------------------
# Escenario completo
# ---------------------------------------------------------------------------


REFERENCE_HEADER = ["ID_MUESTRA", "TIPO_DOCUMENTO", "NUMERO_DOCUMENTO", "FECHA", "TERCERO", "NIT", "VALOR_ESPERADO", "MONEDA",
                    "CONTRATO", "ORDEN_COMPRA", "CONCEPTO", "CENTRO_COSTO", "CUENTA_CONTABLE", "ARCHIVO_SOPORTE"]


def scenario() -> tuple[list[list], dict[str, bytes], list[list]]:
    """Devuelve (filas de referencia, documentos {nombre: bytes}, resultado esperado)."""
    d = date
    rows: list[list] = []
    docs: dict[str, bytes] = {}
    expected: list[list] = []

    def row(sid, tipo, num, fecha, tercero, company, valor, estado, contrato=None, oc=None, concepto="Servicios", archivo=None, nota=""):
        rows.append([sid, tipo, num, fecha, tercero, company.nit if company else None, valor, "COP", contrato, oc, concepto, "CC-100", "5135", archivo])
        expected.append([sid, estado, archivo, nota])

    row("M001", "FACTURA", "FV-1001", d(2026, 3, 2), "PROVEEDOR ANDINO SAS", A, 1_000_000, "COINCIDE", nota="Factura PDF con texto, coincidencia exacta")
    docs["M001_factura_FV-1001.pdf"] = render_pdf(invoice("FV-1001", A, d(2026, 3, 2), 1_000_000, due=d(2026, 4, 1)))

    row("M002", "FE", "FE-2001", d(2026, 3, 5), "Servicios Logisticos del Caribe", B, 2_380_000, "COINCIDE", nota="XML DIAN (AttachedDocument) + representación gráfica PDF")
    cufe = cufe_for("FE-2001")
    docs["M002_FE-2001.xml"] = ubl_invoice("FE-2001", B, d(2026, 3, 5), 2_380_000, cufe, attached=True)
    docs["M002_FE-2001_representacion.pdf"] = render_pdf(invoice("FE-2001", B, d(2026, 3, 5), 2_380_000, title="FACTURA ELECTRÓNICA DE VENTA", cufe=cufe, oc="OC-5500"))

    row("M003", "FACTURA", "C-3001", d(2026, 3, 10), "Consultores Boreal S.A.S.", C, 1_000_000, "COINCIDE CON TOLERANCIA", nota="Valor 1.090.000 vs 1.000.000 (9 %)")
    docs["M003_factura_C-3001.pdf"] = render_pdf(invoice("C-3001", C, d(2026, 3, 10), 1_090_000))

    row("M004", "FACTURA", "TC-4001", d(2026, 3, 12), "Tecnologia Cumbre SA", D, 1_000_000, "EXCEPCIÓN", nota="Valor 1.120.000 vs 1.000.000 (12 %)")
    docs["M004_factura_TC-4001.pdf"] = render_pdf(invoice("TC-4001", D, d(2026, 3, 12), 1_120_000))

    row("M005", "FACTURA", "PE-5001", d(2026, 3, 14), "Papeleria La Estrella", E, 0, "REVISIÓN MANUAL", nota="Valor esperado cero")
    docs["M005_factura_PE-5001.pdf"] = render_pdf(invoice("PE-5001", E, d(2026, 3, 14), 50_000, concept="Papelería"))

    row("M006", "FACTURA", "DN-6001", d(2026, 3, 16), "Distribuidora Nevado S.A.S", F, 750_000, "COINCIDE", nota="Factura escaneada PNG (OCR) con alias del emisor")
    docs["M006_factura_escaneada_DN-6001.png"] = to_png(render_image(invoice("DN-6001", F, d(2026, 3, 16), 750_000, issuer_label="NEVADO DISTRIBUCIONES"), seed=6))

    row("M007", "FACTURA", "FV-1002", d(2026, 3, 18), "Proveedor Andino", A, 2_500_000, "COINCIDE", nota="PDF escaneado (OCR), nombre 'PROVEEDOR ANDINO S A S'")
    docs["M007_factura_escaneada_FV-1002.pdf"] = to_scanned_pdf(render_image(invoice("FV-1002", A, d(2026, 3, 18), 2_500_000, issuer_label="PROVEEDOR ANDINO S A S"), seed=7))

    row("M008", "FACTURA", "IP-8001", d(2026, 3, 20), "Ingenieria Paramo SAS", G, 5_000_000, "COINCIDE", nota="Dos facturas suman el valor contabilizado")
    docs["M008_factura_IP-8001.pdf"] = render_pdf(invoice("IP-8001", G, d(2026, 3, 20), 3_000_000, concept="Diseño estructural fase 1"))
    docs["M008_factura_IP-8002.pdf"] = render_pdf(invoice("IP-8002", G, d(2026, 3, 21), 2_000_000, concept="Diseño estructural fase 2"))

    row("M009", "CUENTA_COBRO", "CC-901", d(2026, 3, 22), "Transportes Altiplano Ltda.", H, 1_800_000, "COINCIDE", nota="Un documento respalda dos partidas")
    row("M010", "CUENTA_COBRO", "CC-901", d(2026, 3, 22), "TRANSPORTES ALTIPLANO", H, 1_200_000, "COINCIDE", nota="Un documento respalda dos partidas")
    docs["M009_M010_cuenta_cobro_CC-901.pdf"] = render_pdf([
        ("CUENTA DE COBRO No. CC-901", 50, True), ("Fecha: 22/03/2026", 50, False),
        (f"{CLIENT[0].upper()}", 50, False), (f"NIT cliente: {nit(CLIENT[1])}", 50, False),
        ("DEBE A:", 50, True), (f"{H.name} NIT {H.nit}", 50, False),
        ("La suma de: TRES MILLONES DE PESOS ($ 3.000.000)", 50, False),
        ("Concepto: Fletes nacionales marzo 2026", 50, False),
        ("Valor total", 50, True, 380, money(3_000_000)),
        ("Cuenta de ahorros No. 123-456789-01 Banco Ficticio", 50, False),
        ("Firma: Representante Transportes Altiplano (ficticio)", 50, False),
    ])

    row("M011", "FACTURA", "AN-1101", d(2026, 3, 24), "Proveedor Andino SAS", A, 600_000, "EXCEPCIÓN", nota="El número coincide pero el NIT del documento es otro")
    docs["M011_factura_AN-1101.pdf"] = render_pdf(invoice("AN-1101", I, d(2026, 3, 24), 600_000))

    row("M012", "FACTURA", "FV-1201", d(2026, 3, 25), "Proveedor Andino SAS", A, 1_500_000, "SIN SOPORTE", nota="No se cargó soporte")

    row("M013", "FACTURA", "FV-1301", d(2026, 3, 26), "Proveedor Andino SAS", A, 450_000, "POSIBLE DUPLICADO", nota="El mismo archivo se cargó dos veces")
    pdf13 = render_pdf(invoice("FV-1301", A, d(2026, 3, 26), 450_000))
    docs["M013_factura_FV-1301.pdf"] = pdf13
    docs["M013_factura_FV-1301_copia.pdf"] = pdf13

    row("M014", "RC", "RC-1401", d(2026, 3, 27), "Distribuidora Nevado", F, 320_000, "COINCIDE", nota="Recibo de caja escaneado TIFF (OCR)")
    docs["M014_recibo_caja_RC-1401.tiff"] = to_tiff(render_image([
        (CLIENT[0].upper(), 50, True), (f"NIT {nit(CLIENT[1])}", 50, False), ("", 50, False),
        ("RECIBO DE CAJA No. RC-1401", 50, True), ("Fecha: 27/03/2026", 50, False),
        (f"Recibimos de: {F.name}", 50, False), (f"NIT: {F.nit}", 50, False),
        ("La suma de: $ 320.000", 50, False), ("Concepto: Abono factura de venta 7788", 50, False),
        ("Forma de pago: Transferencia", 50, False), ("Firma: Cajero (ficticio)", 50, False),
    ], seed=14))

    row("M015", "CE", "CE-1501", d(2026, 3, 28), "Montajes Sabana SAS", J, 4_200_000, "COINCIDE", contrato="CT-2026-015", oc="OC-7788",
        concepto="Montaje estanterías", nota="Egreso + contrato (DOCX) + orden de compra como complementarios")
    docs["M015_comprobante_egreso_CE-1501.pdf"] = render_pdf([
        (CLIENT[0].upper(), 50, True), (f"NIT {nit(CLIENT[1])}", 50, False),
        ("COMPROBANTE DE EGRESO No. CE-1501", 50, True), ("Fecha: 28/03/2026", 50, False),
        (f"Pagado a: {J.name}", 50, False), (f"NIT: {J.nit}", 50, False),
        ("Concepto: Pago factura MS-150 montaje estanterías", 50, False),
        ("Contrato No. CT-2026-015", 50, False), ("Orden de compra No. OC-7788", 50, False),
        ("Valor pagado", 50, True, 380, money(4_200_000)),
        ("Cuenta corriente No. 987-654321-00", 50, False), ("Elaborado por: Tesorería (ficticio)", 50, False),
    ])
    docs["M015_contrato_CT-2026-015.docx"] = contract_docx("CT-2026-015", J, 4_200_000)
    docs["M015_orden_compra_OC-7788.pdf"] = render_pdf([
        (CLIENT[0].upper(), 50, True), (f"NIT {nit(CLIENT[1])}", 50, False),
        ("ORDEN DE COMPRA No. OC-7788", 50, True), ("Fecha: 20/01/2026", 50, False),
        (f"Proveedor: {J.name}", 50, False), (f"NIT: {J.nit}", 50, False),
        ("Contrato No. CT-2026-015", 50, False), ("Descripción: Estanterías industriales y montaje", 50, False),
        ("Valor total", 50, True, 380, money(4_200_000)), ("Aprobado por: Gerencia (ficticio)", 50, False),
    ])

    row("M016", "FACTURA", "PE-1601", d(2026, 3, 29), "Papeleria La Estrella", E, 210_000, "DOCUMENTO ILEGIBLE", archivo="M016_soporte_ilegible.png", nota="Imagen sin texto legible")
    blank = Image.new("L", (1200, 1600), 250)
    dr = ImageDraw.Draw(blank)
    rnd = random.Random(16)
    for _ in range(30):
        x, y = rnd.randrange(1100), rnd.randrange(1500)
        dr.ellipse((x, y, x + rnd.randrange(5, 40), y + rnd.randrange(5, 40)), fill=rnd.randrange(120, 200))
    docs["M016_soporte_ilegible.png"] = to_png(blank.filter(ImageFilter.GaussianBlur(8)))

    row("M017", "FACTURA", "FV-1701", d(2026, 2, 1), "Proveedor Andino SAS", A, 880_000, "EXCEPCIÓN", nota="Fecha del documento 47 días después de la contabilizada")
    docs["M017_factura_FV-1701.jpg"] = _jpg(render_image(invoice("FV-1701", A, d(2026, 3, 20), 880_000), seed=17))

    docs["EXTRA_factura_no_referenciada_PE-9999.pdf"] = render_pdf(invoice("PE-9999", E, d(2026, 3, 30), 99_000, concept="Resmas de papel"))
    return rows, docs, expected


def _jpg(img: Image.Image) -> bytes:
    b = io.BytesIO()
    img.convert("RGB").save(b, "JPEG", quality=90)
    return b.getvalue()


def build_reference_workbook(rows: list[list], expected: list[list]) -> bytes:
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    head_font = Font(bold=True, color="FFFFFF")

    def sheet(title, header, data):
        ws = wb.create_sheet(title)
        ws.append(header)
        for c in ws[1]:
            c.fill, c.font = head_fill, head_font
        for r in data:
            ws.append(r)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(12, min(45, max(len(str(c.value or "")) for c in col) + 2))
        return ws

    wb.remove(wb.active)
    total = sum(r[6] for r in rows)
    sheet("CONFIGURACION", ["PARAMETRO", "VALOR", "DESCRIPCION"], [
        ["TOLERANCIA_VALOR", 0.10, "Diferencia porcentual máxima aceptada (10 %)"],
        ["POLITICA_VALOR_CERO", "REVISION_MANUAL", "REVISION_MANUAL o EXACTA"],
        ["UMBRAL_NOMBRE", 85, "Similitud mínima de nombres (%)"],
        ["TOLERANCIA_DIAS", 5, "Días de diferencia aceptados"],
        ["UMBRAL_RELACION", 0.60, "Puntaje mínimo para relacionar"],
        ["NIT_CLIENTE", nit(CLIENT[1]), "NIT de la entidad auditada"],
        ["NOMBRE_CLIENTE", CLIENT[0], "Entidad auditada (ficticia)"],
        ["PERIODO", "2026-03", "Periodo auditado"],
        ["CONTROL_TOTAL_REGISTROS", len(rows), "Control de integridad: cantidad de partidas"],
        ["CONTROL_TOTAL_VALOR", total, "Control de integridad: suma de VALOR_ESPERADO"],
    ])
    ws = sheet("REFERENCIA_VOUCHING", REFERENCE_HEADER, rows)
    for row_cells in ws.iter_rows(min_row=2):
        row_cells[3].number_format = "yyyy-mm-dd"
        row_cells[6].number_format = "#,##0"
    sheet("ENTIDADES_ALIAS", ["NIT", "NOMBRE_CANONICO", "ALIAS"], [
        [F.nit, F.name, "NEVADO DISTRIBUCIONES"],
        [B.nit, B.name, "SERLOCARIBE"],
        [A.nit, A.name, "ANDINO PROVEEDORES"],
    ])
    sheet("TIPOS_DOCUMENTO", ["CODIGO", "NOMBRE", "PALABRAS_CLAVE", "SOPORTA_VALOR"], [
        ["FACTURA_COMPRA", "Factura de compra", "factura de venta; factura", "SI"],
        ["FACTURA_ELECTRONICA", "Factura electrónica", "factura electrónica; cufe", "SI"],
        ["RECIBO_CAJA", "Recibo de caja", "recibo de caja; recibimos de", "SI"],
        ["COMPROBANTE_EGRESO", "Comprobante de egreso", "comprobante de egreso; pagado a", "SI"],
        ["CONTRATO", "Contrato", "contrato de prestación de servicios", "NO"],
        ["ORDEN_COMPRA", "Orden de compra", "orden de compra", "NO"],
        ["CUENTA_COBRO", "Cuenta de cobro", "cuenta de cobro; debe a", "SI"],
    ])
    sheet("CAMPOS_EXTRACCION", ["CAMPO", "DESCRIPCION", "OBLIGATORIO", "TIPO_DATO", "TIPOS_DOCUMENTO"], [
        ["numero_documento", "Número del documento", "SI", "TEXTO", "TODOS"],
        ["fecha_emision", "Fecha de emisión", "SI", "FECHA", "TODOS"],
        ["emisor_nit", "NIT del emisor", "SI", "NIT", "FACTURA"],
        ["valor_total", "Valor total", "SI", "VALOR", "TODOS"],
        ["cufe", "CUFE / CUDE", "NO", "TEXTO", "FACTURA_ELECTRONICA"],
        ["numero_contrato", "Número de contrato", "NO", "TEXTO", "CONTRATO"],
        ["orden_compra", "Orden de compra", "NO", "TEXTO", "FACTURA; ORDEN_COMPRA"],
    ])
    sheet("RESULTADO_ESPERADO", ["ID_MUESTRA", "ESTADO_ESPERADO", "ARCHIVO", "OBSERVACION"], expected)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate(out_dir: Path) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    docs_dir = out_dir / "documentos"
    docs_dir.mkdir(exist_ok=True)
    rows, docs, expected = scenario()
    (out_dir / "referencia_vouching_demo.xlsx").write_bytes(build_reference_workbook(rows, expected))
    for name, data in docs.items():
        (docs_dir / name).write_bytes(data)
    return out_dir


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parents[2] / "demo"
    generate(target)
    print(f"Datos ficticios generados en {target}")
