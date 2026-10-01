"""Exportación de resultados y evidencias a Excel (openpyxl)."""
from __future__ import annotations

import io
from collections import defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from .. import models
from ..config import APP_NAME, APP_VERSION
from ..models import Status
from . import dashboard
from .extraction.classifier import DOC_TYPES
from .settings_defaults import CRITERION_LABELS, PRIORITY_LABELS, merged_parameters

SHEETS = {
    "resumen": "Resumen ejecutivo",
    "resultado": "Resultado completo",
    "coincidencias": "Coincidencias",
    "excepciones": "Excepciones",
    "sin_soporte": "Partidas sin soporte",
    "no_referenciados": "Soportes no referenciados",
    "evidencias": "Evidencias",
    "historial": "Historial de revisiones",
    "parametros": "Parámetros utilizados",
    "registro": "Registro modelo OCR reglas",
}

BRAND = "1F3A5F"
ACCENT = "C8963E"
HEADER_FILL = PatternFill("solid", fgColor=BRAND)
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=16, color=BRAND)
THIN = Side(style="thin", color="D0D7DE")
STATUS_FILL = {
    Status.COINCIDE: "D4EDDA",
    Status.COINCIDE_TOLERANCIA: "E6F4EA",
    Status.EXCEPCION: "F8D7DA",
    Status.REVISION_MANUAL: "FFF3CD",
    Status.SIN_SOPORTE: "FDE2C8",
    Status.NO_REFERENCIADO: "E2E3E5",
    Status.POSIBLE_DUPLICADO: "F5C6CB",
    Status.ILEGIBLE: "D6D8DB",
    Status.ERROR_LECTURA: "D6D8DB",
}

RESULT_COLUMNS = [
    "ID de muestra", "Archivo", "Hash del archivo (SHA-256)", "Tipo documental", "Página de evidencia", "Número esperado",
    "Número extraído", "Fecha esperada", "Fecha extraída", "Tercero esperado", "Tercero extraído", "NIT esperado",
    "NIT extraído", "Valor esperado", "Valor extraído", "Diferencia absoluta", "Diferencia porcentual", "Tolerancia aplicada",
    "Similitud del nombre", "Coincidencia del NIT", "Coincidencia del DV", "Coincidencia del número", "Diferencia en días",
    "Confianza OCR", "Puntaje de relación", "Estado", "Estado automático", "Motivo de excepción", "Texto de evidencia",
    "Soportes complementarios", "Decisión del revisor", "Comentario del revisor", "Revisor", "Fecha de revisión",
]


def safe(v):
    """Evita inyección de fórmulas en Excel (CSV/Formula injection)."""
    if isinstance(v, str):
        v = ILLEGAL_CHARACTERS_RE.sub("", v)
        if v[:1] in ("=", "+", "-", "@", "\t", "\r"):
            return "'" + v
        return v[:32000]
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, datetime) and v.tzinfo:
        return v.astimezone(timezone.utc).replace(tzinfo=None)
    return v


def _write_table(ws, start_row: int, headers: list[str], rows: list[list], status_col: int | None = None, money_cols=(), pct_cols=()):
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=start_row, column=j, value=h)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = Border(bottom=THIN)
    for i, row in enumerate(rows, start=start_row + 1):
        for j, v in enumerate(row, start=1):
            c = ws.cell(row=i, column=j, value=safe(v))
            if j - 1 in money_cols:
                c.number_format = "#,##0.00"
            elif j - 1 in pct_cols:
                c.number_format = "0.00%"
            elif isinstance(v, (date, datetime)):
                c.number_format = "yyyy-mm-dd" if not isinstance(v, datetime) else "yyyy-mm-dd hh:mm"
        if status_col is not None:
            fill = STATUS_FILL.get(row[status_col])
            if fill:
                ws.cell(row=i, column=status_col + 1).fill = PatternFill("solid", fgColor=fill)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    if rows:
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{start_row + len(rows)}"
    for j, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(str(r[j - 1])) if j - 1 < len(r) and r[j - 1] is not None else 0 for r in rows[:300]])
        ws.column_dimensions[get_column_letter(j)].width = min(max(10, width + 2), 60)


def _brand(ws, title: str, project: models.Project) -> int:
    ws["A1"] = APP_NAME
    ws["A1"].font = TITLE_FONT
    ws["A2"] = title
    ws["A2"].font = Font(bold=True, size=12, color=ACCENT)
    ws["A3"] = safe(f"Proyecto: {project.code} — {project.name} | Cliente: {project.client_name} | Periodo: {project.period or '-'}")
    ws["A4"] = f"Generado: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')} | {APP_NAME} v{APP_VERSION}"
    ws["A4"].font = Font(italic=True, color="666666")
    return 6


def _result_rows(db: Session, project_id: int) -> list[tuple[models.VouchingResult, list]]:
    results = (
        db.query(models.VouchingResult)
        .join(models.ReferenceItem)
        .filter(models.VouchingResult.project_id == project_id)
        .order_by(models.ReferenceItem.row_number)
        .all()
    )
    links = db.query(models.MatchLink).filter_by(project_id=project_id).all()
    docs = {d.id: d for d in db.query(models.Document).filter_by(project_id=project_id)}
    users = {u.id: u for u in db.query(models.User)}
    by_ref = defaultdict(list)
    for ln in links:
        by_ref[ln.reference_item_id].append(ln)
    out = []
    for res in results:
        r = res.reference_item
        lns = by_ref.get(r.id, [])
        prim = [docs[ln.document_id] for ln in lns if ln.role == "PRINCIPAL" and ln.document_id in docs]
        rep = [docs[ln.document_id] for ln in lns if ln.role == "REPRESENTACION_GRAFICA" and ln.document_id in docs]
        comp = [docs[ln.document_id] for ln in lns if ln.role == "COMPLEMENTARIO" and ln.document_id in docs]
        files = prim + rep
        reviewer = users.get(res.reviewer_id)
        reasons = res.reasons or []
        out.append((res, [
            r.sample_id,
            "; ".join(d.filename for d in files),
            "; ".join(d.sha256 for d in files),
            "; ".join(sorted({DOC_TYPES.get(d.doc_type, d.doc_type or "") for d in prim})),
            res.evidence_page,
            r.doc_number,
            res.extracted_number,
            r.doc_date,
            res.extracted_date,
            r.third_party,
            res.extracted_party,
            r.nit,
            res.extracted_nit,
            r.expected_value,
            res.extracted_value,
            res.abs_difference,
            res.pct_difference,
            res.tolerance_applied,
            res.name_similarity,
            res.nit_match,
            res.dv_match,
            res.number_match,
            res.days_difference,
            res.ocr_confidence,
            res.score,
            res.status,
            res.auto_status,
            " | ".join(reasons) if res.status not in (Status.COINCIDE,) else "",
            res.evidence_text,
            "; ".join(f"{d.filename} ({DOC_TYPES.get(d.doc_type, d.doc_type)})" for d in comp),
            res.review_decision,
            res.review_comment,
            reviewer.email if reviewer else None,
            res.reviewed_at,
        ]))
    return out


def build_workbook(db: Session, project: models.Project, only: list[str] | None = None) -> bytes:
    keys = [k for k in SHEETS if not only or k in only]
    wb = Workbook()
    wb.remove(wb.active)
    rows = _result_rows(db, project.id)
    money, pct = (13, 14, 15), (16, 17)
    status_idx = RESULT_COLUMNS.index("Estado")

    if "resumen" in keys:
        ws = wb.create_sheet(SHEETS["resumen"][:31])
        r0 = _brand(ws, "Resumen ejecutivo de vouching", project)
        kpi = dashboard.kpis(db, project.id)
        labels = [
            ("Total de partidas", kpi["total_partidas"]), ("Total de documentos", kpi["total_documentos"]),
            ("Partidas con soporte", kpi["partidas_con_soporte"]), ("Partidas sin soporte", kpi["partidas_sin_soporte"]),
            ("Coincidencias completas", kpi["coincidencias"]), ("Coincidencias con tolerancia", kpi["coincidencias_tolerancia"]),
            ("Excepciones", kpi["excepciones"]), ("Pendientes de revisión", kpi["pendientes_revision"]),
            ("Documentos ilegibles", kpi["documentos_ilegibles"]), ("Valor total de la población", kpi["valor_poblacion"]),
            ("Valor soportado", kpi["valor_soportado"]), ("Valor con diferencias", kpi["valor_con_diferencias"]),
            ("Suma de diferencias absolutas", kpi["suma_diferencias_absolutas"]),
            ("Porcentaje de cobertura", kpi["porcentaje_cobertura"] / 100), ("Porcentaje de avance", kpi["porcentaje_avance"] / 100),
        ]
        _write_table(ws, r0, ["Indicador", "Valor"], [[a, b] for a, b in labels])
        for i, (label, _v) in enumerate(labels, start=r0 + 1):
            c = ws.cell(row=i, column=2)
            if "Valor" in label or "Suma" in label:
                c.number_format = "#,##0.00"
            elif "Porcentaje" in label:
                c.number_format = "0.00%"
        r1 = r0 + len(labels) + 3
        ws.cell(row=r1 - 1, column=1, value="Distribución por estado").font = Font(bold=True, color=BRAND)
        _write_table(ws, r1, ["Estado", "Partidas", "Valor esperado"], [[k, v["partidas"], v["valor"]] for k, v in kpi["por_estado"].items()], status_col=0, money_cols=(2,))
        run = db.query(models.MatchRun).filter_by(project_id=project.id).order_by(models.MatchRun.id.desc()).first()
        r2 = r1 + len(kpi["por_estado"]) + 3
        ws.cell(row=r2 - 1, column=1, value="Conciliación de conteos y valores").font = Font(bold=True, color=BRAND)
        rec = (run.reconciliation if run else {}) or {}
        _write_table(ws, r2, ["Control", "Resultado"], [[k, str(v)] for k, v in rec.items()])
        ws.cell(row=r2 + len(rec) + 2, column=1, value="Nota: una excepción representa una diferencia que debe investigarse; no constituye por sí misma un indicio de fraude.").font = Font(italic=True)
        ws.column_dimensions["A"].width = 42
        ws.column_dimensions["B"].width = 60

    def result_sheet(key, title, predicate):
        ws = wb.create_sheet(SHEETS[key][:31])
        r0 = _brand(ws, title, project)
        sel = [row for res, row in rows if predicate(res)]
        _write_table(ws, r0, RESULT_COLUMNS, sel, status_col=status_idx, money_cols=money, pct_cols=pct)

    if "resultado" in keys:
        result_sheet("resultado", "Resultado completo", lambda r: True)
    if "coincidencias" in keys:
        result_sheet("coincidencias", "Coincidencias", lambda r: r.status in (Status.COINCIDE, Status.COINCIDE_TOLERANCIA))
    if "excepciones" in keys:
        result_sheet("excepciones", "Excepciones y casos para revisión", lambda r: r.status in (
            Status.EXCEPCION, Status.REVISION_MANUAL, Status.POSIBLE_DUPLICADO, Status.ILEGIBLE, Status.ERROR_LECTURA))
    if "sin_soporte" in keys:
        result_sheet("sin_soporte", "Partidas sin soporte", lambda r: r.status == Status.SIN_SOPORTE)

    docs = db.query(models.Document).filter_by(project_id=project.id).order_by(models.Document.id).all()
    if "no_referenciados" in keys:
        ws = wb.create_sheet(SHEETS["no_referenciados"][:31])
        r0 = _brand(ws, "Soportes no referenciados y documentos con incidencias", project)
        sel = [d for d in docs if d.vouching_status in (Status.NO_REFERENCIADO, Status.POSIBLE_DUPLICADO, Status.ILEGIBLE, Status.ERROR_LECTURA)]
        data = []
        for d in sel:
            f = {x.field_name: x.effective_value for x in d.fields}
            data.append([d.filename, d.sha256, d.vouching_status, DOC_TYPES.get(d.doc_type, d.doc_type), f.get("numero_documento"),
                         f.get("fecha_emision"), f.get("emisor_nombre"), f.get("emisor_nit"), f.get("valor_total"), d.ocr_confidence, d.error])
        _write_table(ws, r0, ["Archivo", "Hash SHA-256", "Estado", "Tipo documental", "Número", "Fecha", "Emisor", "NIT emisor", "Valor total", "Confianza OCR", "Observación"], data, status_col=2)

    if "evidencias" in keys:
        ws = wb.create_sheet(SHEETS["evidencias"][:31])
        r0 = _brand(ws, "Evidencias por campo", project)
        users = {u.id: u.email for u in db.query(models.User)}
        data = []
        for d in docs:
            for f in sorted(d.fields, key=lambda x: x.field_name):
                data.append([d.filename, d.sha256, f.field_name, f.value, f.normalized_value, f.corrected_value, f.review_status,
                             f.page, ", ".join(f"{x:.4f}" for x in f.bbox) if f.bbox else None, f.evidence_text, f.method, f.confidence,
                             f.extracted_at, users.get(f.reviewed_by), f.reviewed_at, f.review_comment])
        _write_table(ws, r0, ["Archivo", "Hash SHA-256", "Campo", "Valor extraído (original)", "Valor normalizado", "Valor corregido",
                              "Estado de revisión", "Página", "Región (x0, y0, x1, y1)", "Texto de evidencia", "Método de extracción",
                              "Confianza", "Fecha de procesamiento", "Revisor", "Fecha de revisión", "Comentario"], data)

    if "historial" in keys:
        ws = wb.create_sheet(SHEETS["historial"][:31])
        r0 = _brand(ws, "Historial de revisiones", project)
        evs = db.query(models.ReviewEvent).filter_by(project_id=project.id).order_by(models.ReviewEvent.id).all()
        _write_table(ws, r0, ["Fecha", "Usuario", "Acción", "ID de muestra", "Documento", "Campo", "Valor anterior", "Valor nuevo", "Comentario"],
                     [[e.created_at, e.user_email, e.action, e.sample_id, e.document_id, e.field_name, e.old_value, e.new_value, e.comment] for e in evs])

    if "parametros" in keys:
        ws = wb.create_sheet(SHEETS["parametros"][:31])
        r0 = _brand(ws, "Parámetros utilizados", project)
        params = merged_parameters(project.settings)
        data = [[k, str(v)] for k, v in params.items() if k != "PESOS"]
        data += [[f"PESO {CRITERION_LABELS[k]}", v, PRIORITY_LABELS.get(k)] for k, v in params["PESOS"].items()]
        data += [["Fórmula de valor", "DIFERENCIA_PORCENTAJE = ABS(VALOR_EXTRAIDO - VALOR_ESPERADO) / ABS(VALOR_ESPERADO); coincide si <= TOLERANCIA_VALOR"],
                 ["Valor esperado cero", "No se divide entre cero: se exige coincidencia exacta; si difiere -> " + params["POLITICA_VALOR_CERO"]],
                 ["Procesamiento con IA autorizado", "SÍ" if project.allow_ai_processing else "NO"],
                 ["Uso de documentos para entrenamiento", "NO (bloqueado)"],
                 ["Retención (días)", project.retention_days]]
        _write_table(ws, r0, ["Parámetro", "Valor", "Prioridad"], data)

    if "registro" in keys:
        ws = wb.create_sheet(SHEETS["registro"][:31])
        r0 = _brand(ws, "Registro del modelo, OCR y reglas aplicadas", project)
        from .extraction.readers import ocr_available
        from .matching import ENGINE_VERSION
        from .extraction.pipeline import ENGINE_VERSION as EXTRACTION_VERSION
        try:
            import pytesseract
            tess = str(pytesseract.get_tesseract_version()) if ocr_available() else "no disponible"
        except Exception:  # pragma: no cover
            tess = "no disponible"
        from ..config import get_settings
        s = get_settings()
        meta = [
            ["Aplicativo", f"{APP_NAME} v{APP_VERSION}"], ["Motor de vouching", ENGINE_VERSION], ["Motor de extracción", EXTRACTION_VERSION],
            ["OCR", f"Tesseract {tess} (idiomas {s.ocr_languages}, {s.ocr_dpi} DPI)"], ["PDF con texto", "pdfplumber (coordenadas por palabra)"],
            ["XML", "Extracción determinística UBL 2.1 DIAN (lxml seguro, sin entidades externas)"],
            ["Clasificación", "Reglas por palabras clave ponderadas + hoja TIPOS_DOCUMENTO"],
            ["IA", f"{s.llm_model} (solo campos faltantes con cita literal verificada)" if s.llm_enabled and project.allow_ai_processing else "No utilizada"],
            ["Normalización de nombres", "Mayúsculas, sin tildes ni puntuación, sin sufijos jurídicos, alias, similitud difusa (rapidfuzz)"],
            ["Normalización de NIT", "Sin separadores; base y DV separados; DV verificado con algoritmo DIAN módulo 11"],
        ]
        _write_table(ws, r0, ["Componente", "Detalle"], meta)
        r1 = r0 + len(meta) + 3
        data = [[d.filename, d.sha256, d.file_type, d.processing_status, d.extraction_method, d.page_count, d.ocr_confidence,
                 DOC_TYPES.get(d.doc_type, d.doc_type), d.doc_type_confidence, d.classification_reason, d.av_status, d.processing_ms,
                 d.processed_at, d.error] for d in docs]
        _write_table(ws, r1, ["Archivo", "Hash SHA-256", "Formato", "Estado técnico", "Método", "Páginas", "Confianza OCR", "Tipo documental",
                              "Confianza clasificación", "Motivo de clasificación", "Antivirus", "Tiempo (ms)", "Procesado", "Observaciones"], data)

    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = BRAND
    wb.properties.creator = APP_NAME
    wb.properties.title = f"{APP_NAME} — {project.code}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
