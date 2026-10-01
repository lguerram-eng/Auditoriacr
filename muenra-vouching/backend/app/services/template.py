"""Plantilla simple de carga (una hoja CARGA, una fila por documento, cualquier tipo documental)."""
from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

SIMPLE_HEADER = [
    "ID", "TIPO_DOCUMENTO", "NUMERO_DOCUMENTO", "FECHA_DOCUMENTO", "NOMBRE_TERCERO", "NIT_IDENTIFICACION", "CONCEPTO",
    "SUBTOTAL", "IVA", "RETENCIONES", "VALOR_TOTAL", "MONEDA", "CONTRATO_OC", "ARCHIVO_ESPERADO", "TOLERANCIA_VALOR", "OBSERVACIONES",
]
WIDTHS = [12, 20, 22, 15, 30, 20, 34, 16, 14, 16, 18, 11, 20, 26, 17, 30]
DOC_TYPES = "FACTURA,FE,RC,CE,CONTRATO,OTROSI,ORDEN_COMPRA,CUENTA_COBRO,NOTA_CREDITO,NOTA_DEBITO,DOCUMENTO_SOPORTE,CERTIFICACION,EXTRACTO,SOPORTE_PAGO,OTRO"
EXAMPLES = [
    ["MV-001", "FACTURA", "FV-00125", date(2026, 9, 27), "Proveedor Ejemplo S.A.S.", "900123456", "Servicios profesionales",
     1000000, 190000, 0, 1190000, "COP", "OC-4500123", "factura_125.pdf", 0.1, "EJEMPLO"],
    ["MV-002", "RC", "RC-00315", date(2026, 9, 25), "Cliente Ejemplo Ltda.", "800987654", "Recaudo de cartera",
     850000, 0, 0, 850000, "COP", None, "rc_315.pdf", 0.1, "EJEMPLO"],
    ["MV-003", "CE", "CE-00902", date(2026, 9, 30), "Proveedor Ejemplo SAS", "901555444", "Pago de factura",
     500000, 95000, 0, 595000, "COP", "CT-2026-021", "ce_902.pdf", 0.1, "EJEMPLO"],
]
HEADER_ROW = 4
MAX_ROWS = 500


def build_simple_template(rows: list[list] | None = None, examples: bool = True) -> bytes:
    """Genera la plantilla simple. ``rows`` permite precargar filas (mismo orden de SIMPLE_HEADER)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "CARGA"
    n = len(SIMPLE_HEADER)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n)
    ws["A1"] = "MUENRA VOUCHING | Carga general de documentos"
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="17365D")
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28
    ws["A2"] = ("Una fila por documento. Sirve para FACTURA, RC, CE, CONTRATO, OC y OTRO. Obligatorias: ID, TIPO_DOCUMENTO, "
                "NUMERO_DOCUMENTO, FECHA_DOCUMENTO, NOMBRE_TERCERO, NIT_IDENTIFICACION y VALOR_TOTAL. "
                "Sustituya o elimine las filas amarillas de ejemplo.")
    ws["A2"].font = Font(bold=True)
    ws["A2"].fill = PatternFill("solid", fgColor="FFF2CC")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 32
    thin = Side(style="thin", color="BFBFBF")
    for j, (h, w) in enumerate(zip(SIMPLE_HEADER, WIDTHS), start=1):
        c = ws.cell(row=HEADER_ROW, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F75B5")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[HEADER_ROW].height = 38
    data = (EXAMPLES if examples else []) + (rows or [])
    yellow = PatternFill("solid", fgColor="FFF2CC")
    for i, r in enumerate(data, start=HEADER_ROW + 1):
        for j, v in enumerate(r, start=1):
            c = ws.cell(row=i, column=j, value=v)
            if examples and i < HEADER_ROW + 1 + len(EXAMPLES):
                c.fill = yellow
    last = HEADER_ROW + MAX_ROWS
    for row in ws.iter_rows(min_row=HEADER_ROW + 1, max_row=last):
        row[3].number_format = "yyyy-mm-dd"
        for k in (7, 8, 9, 10):
            row[k].number_format = "#,##0.00"
        row[14].number_format = "0%"
        row[5].number_format = "@"
    for formula, col in ((f'"{DOC_TYPES}"', "B"), ('"COP,USD,EUR,OTRA"', "L")):
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        dv.add(f"{col}{HEADER_ROW + 1}:{col}{last}")
        ws.add_data_validation(dv)
    dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True,
                        error="Use un valor entre 0 y 1 (0,10 = 10 %)", errorTitle="Tolerancia")
    dv.add(f"O{HEADER_ROW + 1}:O{last}")
    ws.add_data_validation(dv)
    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=2)
    wb.properties.creator = "Muenra Vouching"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
