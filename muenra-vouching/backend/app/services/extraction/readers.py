"""Lectura de PDF (texto y escaneado), imágenes, DOCX y OCR con Tesseract.

Prioridad (ver README):
1. XML  -> extracción determinística (xml_extractor.py)
2. PDF con texto -> extracción directa con coordenadas (pdfplumber)
3. PDF escaneado / imágenes -> OCR (Tesseract) con confianza por palabra
"""
from __future__ import annotations

import io
import logging
from functools import lru_cache
from statistics import mean

from PIL import Image, ImageOps, ImageSequence

from ...config import get_settings
from .base import Line, Method, Page

log = logging.getLogger(__name__)

Image.MAX_IMAGE_PIXELS = 120_000_000  # protección contra "decompression bombs"


# ---------------------------------------------------------------------------
# OCR
# ---------------------------------------------------------------------------


@lru_cache
def ocr_available() -> bool:
    try:
        import pytesseract

        pytesseract.get_tesseract_version()
        return True
    except Exception:  # pragma: no cover - depende del sistema
        return False


@lru_cache
def _ocr_lang() -> str:
    import pytesseract

    wanted = get_settings().ocr_languages.split("+")
    try:
        installed = set(pytesseract.get_languages(config=""))
    except Exception:  # pragma: no cover
        installed = {"eng"}
    langs = [lang for lang in wanted if lang in installed] or ["eng"]
    return "+".join(langs)


def _prepare(img: Image.Image) -> Image.Image:
    img = ImageOps.exif_transpose(img)
    if img.mode not in ("L", "RGB"):
        img = img.convert("RGB")
    gray = ImageOps.grayscale(img)
    # Escala las imágenes pequeñas para mejorar el OCR
    if gray.width < 1400:
        factor = 1400 / gray.width
        gray = gray.resize((int(gray.width * factor), int(gray.height * factor)), Image.LANCZOS)
    return ImageOps.autocontrast(gray)


def ocr_image(img: Image.Image, page_number: int = 1) -> Page:
    """Ejecuta OCR y devuelve líneas con región normalizada y confianza (0-100)."""
    if not ocr_available():
        raise RuntimeError("Tesseract OCR no está instalado en el servidor")
    import pytesseract

    prepared = _prepare(img)
    data = pytesseract.image_to_data(
        prepared, lang=_ocr_lang(), config="--oem 1 --psm 3", output_type=pytesseract.Output.DICT
    )
    W, H = prepared.size
    groups: dict[tuple, list[int]] = {}
    for i, word in enumerate(data["text"]):
        if not word or not word.strip():
            continue
        try:
            conf = float(data["conf"][i])
        except (TypeError, ValueError):
            conf = -1
        if conf < 0:
            continue
        key = (data["block_num"][i], data["par_num"][i], data["line_num"][i])
        groups.setdefault(key, []).append(i)
    lines: list[Line] = []
    for key in sorted(groups, key=lambda k: (min(data["top"][i] for i in groups[k]), k)):
        idx = groups[key]
        text = " ".join(data["text"][i].strip() for i in idx)
        x0 = min(data["left"][i] for i in idx)
        y0 = min(data["top"][i] for i in idx)
        x1 = max(data["left"][i] + data["width"][i] for i in idx)
        y1 = max(data["top"][i] + data["height"][i] for i in idx)
        conf = mean(float(data["conf"][i]) for i in idx)
        lines.append(Line(text=text, bbox=[round(x0 / W, 5), round(y0 / H, 5), round(x1 / W, 5), round(y1 / H, 5)], conf=round(conf, 2)))
    confs = [line.conf for line in lines if line.conf is not None]
    return Page(
        number=page_number,
        method=Method.OCR,
        lines=lines,
        width=float(img.width),
        height=float(img.height),
        ocr_confidence=round(mean(confs), 2) if confs else 0.0,
    )


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def _group_words(words: list[dict], width: float, height: float) -> list[Line]:
    """Agrupa palabras de pdfplumber en líneas usando la coordenada vertical."""
    words = sorted(words, key=lambda w: (round(w["top"]), w["x0"]))
    rows: list[list[dict]] = []
    for w in words:
        if rows and abs(rows[-1][0]["top"] - w["top"]) <= 3:
            rows[-1].append(w)
        else:
            rows.append([w])
    lines = []
    for row in rows:
        row.sort(key=lambda w: w["x0"])
        # Separa columnas muy distantes en líneas distintas para no mezclar etiquetas
        segment: list[dict] = [row[0]]
        segments = [segment]
        for w in row[1:]:
            if w["x0"] - segment[-1]["x1"] > 60:
                segment = [w]
                segments.append(segment)
            else:
                segment.append(w)
        for seg in segments:
            x0 = min(w["x0"] for w in seg)
            x1 = max(w["x1"] for w in seg)
            y0 = min(w["top"] for w in seg)
            y1 = max(w["bottom"] for w in seg)
            lines.append(
                Line(
                    text=" ".join(w["text"] for w in seg),
                    bbox=[round(x0 / width, 5), round(y0 / height, 5), round(x1 / width, 5), round(y1 / height, 5)],
                    conf=100.0,
                )
            )
    return lines


def render_pdf_page(data: bytes, page_index: int, dpi: int = 150) -> Image.Image:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(data)
    try:
        page = pdf[page_index]
        bitmap = page.render(scale=dpi / 72)
        return bitmap.to_pil().convert("RGB")
    finally:
        pdf.close()


def pdf_page_count(data: bytes) -> int:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(data)
    try:
        return len(pdf)
    finally:
        pdf.close()


def read_pdf(data: bytes, force_ocr: bool = False) -> tuple[list[Page], list[dict], list[str]]:
    import pdfplumber

    settings = get_settings()
    pages: list[Page] = []
    tables: list[dict] = []
    warnings: list[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for i, p in enumerate(pdf.pages):
            number = i + 1
            words = [] if force_ocr else p.extract_words(keep_blank_chars=False, use_text_flow=False)
            chars = sum(len(w["text"]) for w in words)
            if chars >= settings.ocr_min_text_chars:
                pages.append(Page(number=number, method=Method.PDF_TEXT, lines=_group_words(words, p.width, p.height), width=float(p.width), height=float(p.height)))
                try:
                    for t in p.extract_tables()[:10]:
                        clean = [[(c or "").strip() for c in row] for row in t if any(row)]
                        if len(clean) >= 2:
                            tables.append({"page": number, "header": clean[0], "rows": clean[1:], "method": Method.PDF_TEXT})
                except Exception as exc:  # pragma: no cover - tablas mal formadas
                    warnings.append(f"Página {number}: no fue posible leer tablas ({exc.__class__.__name__})")
            else:
                if not ocr_available():
                    warnings.append(f"Página {number}: sin texto y OCR no disponible")
                    pages.append(Page(number=number, method=Method.OCR, lines=[], width=float(p.width), height=float(p.height), ocr_confidence=0.0))
                    continue
                img = render_pdf_page(data, i, dpi=settings.ocr_dpi)
                ocr_page = ocr_image(img, number)
                ocr_page.width, ocr_page.height = float(p.width), float(p.height)
                pages.append(ocr_page)
    return pages, tables, warnings


# ---------------------------------------------------------------------------
# Imágenes
# ---------------------------------------------------------------------------


def image_frames(data: bytes) -> list[Image.Image]:
    img = Image.open(io.BytesIO(data))
    frames = [frame.copy() for frame in ImageSequence.Iterator(img)]
    return frames[:200]


def read_image(data: bytes) -> list[Page]:
    return [ocr_image(frame, i + 1) for i, frame in enumerate(image_frames(data))]


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def read_docx(data: bytes) -> tuple[list[Page], list[dict]]:
    import docx

    d = docx.Document(io.BytesIO(data))
    lines = [Line(text=p.text.strip(), conf=100.0) for p in d.paragraphs if p.text.strip()]
    tables = []
    for t in d.tables[:20]:
        rows = [[c.text.strip() for c in r.cells] for r in t.rows]
        if rows:
            tables.append({"page": 1, "header": rows[0], "rows": rows[1:], "method": Method.DOCX})
            for r in rows:
                lines.append(Line(text=" | ".join(r), conf=100.0))
    return [Page(number=1, method=Method.DOCX, lines=lines)], tables
