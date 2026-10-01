"""Prueba de integración completa con los datos ficticios de demostración.

Importa el Excel, carga todos los documentos (PDF, XML, PNG, TIFF, JPG, DOCX),
procesa la cola, ejecuta el motor y compara cada estado con RESULTADO_ESPERADO.
También verifica la conciliación (criterios de aceptación 12 y 13) y la exportación.
"""
from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from app import worker

from .conftest import requires_ocr


@pytest.fixture(scope="module")
def processed_project(client, admin_headers, demo_data):
    r = client.post("/api/projects", headers=admin_headers, json={
        "code": "DEMO-E2E", "name": "Vouching demostración", "client_name": "Cliente Demostración S.A.S.", "period": "2026-03"})
    assert r.status_code == 201, r.text
    pid = r.json()["id"]
    r = client.post(f"/api/projects/{pid}/imports", headers=admin_headers,
                    files={"file": ("referencia.xlsx", demo_data["xlsx"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    assert r.json()["importado"] is True
    files = [("files", (name, data, "application/octet-stream")) for name, data in demo_data["docs"].items()]
    r = client.post(f"/api/projects/{pid}/documents", headers=admin_headers, files=files)
    assert r.status_code == 200, r.text
    assert not r.json()["rechazados"], r.json()["rechazados"]
    worker.drain()
    return pid


@requires_ocr
def test_statuses_match_expected(client, admin_headers, processed_project, demo_data):
    results = client.get(f"/api/projects/{processed_project}/results", headers=admin_headers).json()
    got = {r["id_muestra"]: r for r in results}
    mismatches = []
    for sid, expected_status, _file, note in demo_data["expected"]:
        if got[sid]["estado"] != expected_status:
            mismatches.append((sid, expected_status, got[sid]["estado"], got[sid]["motivos"], note))
    assert not mismatches, "\n".join(map(str, mismatches))


@requires_ocr
def test_explanations_and_evidence(client, admin_headers, processed_project):
    results = client.get(f"/api/projects/{processed_project}/results", headers=admin_headers).json()
    m001 = next(r for r in results if r["id_muestra"] == "M001")
    detail = client.get(f"/api/results/{m001['id']}", headers=admin_headers).json()
    principal = detail["explicacion"]["soportes_principales"][0]
    keys = {c["clave"] for c in principal["criterios"]}
    assert {"nit", "numero", "valor", "fecha", "nombre", "contrato", "orden_compra", "concepto", "moneda", "tipo"} <= keys
    valor = next(c for c in principal["criterios"] if c["clave"] == "valor")
    assert valor["pagina"] == 1 and "1.000.000" in valor["evidencia"] and valor["esperado"] == "1.000.000"
    assert detail["valor_extraido"] == 1000000.0 and detail["diferencia_absoluta"] == 0.0
    # XML + representación gráfica relacionados con la misma partida sin duplicar valor
    m002 = next(r for r in results if r["id_muestra"] == "M002")
    roles = sorted(d["rol"] for d in m002["documentos"])
    assert roles == ["PRINCIPAL", "REPRESENTACION_GRAFICA"]
    # Soportes complementarios (contrato DOCX y OC) para el egreso
    m015 = next(r for r in results if r["id_muestra"] == "M015")
    assert sum(1 for d in m015["documentos"] if d["rol"] == "COMPLEMENTARIO") == 2
    # Un soporte principal de otra partida no se presenta como complementario
    m006 = next(r for r in results if r["id_muestra"] == "M006")
    assert [d["rol"] for d in m006["documentos"]] == ["PRINCIPAL"]


@requires_ocr
def test_reconciliation_counts_and_values(client, admin_headers, processed_project, demo_data):
    rec = client.get(f"/api/projects/{processed_project}/reconciliation", headers=admin_headers).json()
    assert rec["conteo_partidas_cuadra"] is True
    assert rec["valor_cuadra"] is True
    assert rec["conteo_documentos_cuadra"] is True
    assert rec["documentos_cargados"] == len(demo_data["docs"])
    assert rec["sin_duplicacion_de_valor"] is True
    assert rec["validacion_resultado_esperado"]["precision"] == 1.0
    docs = client.get(f"/api/projects/{processed_project}/documents", headers=admin_headers).json()
    by_name = {d["archivo"]: d for d in docs}
    assert by_name["EXTRA_factura_no_referenciada_PE-9999.pdf"]["estado"] == "SOPORTE NO REFERENCIADO"
    assert by_name["M016_soporte_ilegible.png"]["estado"] == "DOCUMENTO ILEGIBLE"


@requires_ocr
def test_dashboard(client, admin_headers, processed_project, demo_data):
    k = client.get(f"/api/projects/{processed_project}/dashboard", headers=admin_headers).json()
    assert k["total_partidas"] == len(demo_data["rows"])
    assert k["total_documentos"] == len(demo_data["docs"])
    assert k["valor_poblacion"] == sum(r[6] for r in demo_data["rows"])
    assert 0 < k["porcentaje_cobertura"] < 100


@requires_ocr
def test_human_review_keeps_original_value(client, admin_headers, processed_project):
    results = client.get(f"/api/projects/{processed_project}/results", headers=admin_headers).json()
    m004 = next(r for r in results if r["id_muestra"] == "M004")
    doc_id = next(d["id"] for d in m004["documentos"] if d["rol"] == "PRINCIPAL")
    doc = client.get(f"/api/documents/{doc_id}", headers=admin_headers).json()
    field = next(f for f in doc["campos"] if f["campo"] == "valor_total")
    r = client.post(f"/api/fields/{field['id']}/review", headers=admin_headers, json={"action": "CORREGIR", "corrected_value": "1.000.000"})
    assert r.status_code == 422  # comentario obligatorio
    r = client.post(f"/api/fields/{field['id']}/review", headers=admin_headers,
                    json={"action": "CORREGIR", "corrected_value": "1.000.000", "comment": "Prueba: valor según nota crédito"})
    assert r.status_code == 200, r.text
    assert r.json()["valor_original"] == "$ 1.120.000"
    worker.drain()
    m004 = next(r for r in client.get(f"/api/projects/{processed_project}/results", headers=admin_headers).json() if r["id_muestra"] == "M004")
    assert m004["estado"] == "COINCIDE"
    r = client.post(f"/api/results/{m004['id']}/review", headers=admin_headers, json={"decision": "APROBADO", "comment": "Revisado"})
    assert r.status_code == 200
    hist = client.get(f"/api/projects/{processed_project}/review-history", headers=admin_headers).json()
    assert {"CAMPO_CORREGIDO", "RESULTADO_APROBADO"} <= {h["accion"] for h in hist}
    # Revertir la corrección para no afectar otras pruebas
    client.post(f"/api/fields/{field['id']}/review", headers=admin_headers, json={"action": "ACEPTAR", "comment": "Se restituye"})
    worker.drain()


@requires_ocr
def test_export_contains_all_sheets_and_columns(client, admin_headers, processed_project, demo_data):
    r = client.get(f"/api/projects/{processed_project}/export", headers=admin_headers)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Resumen ejecutivo", "Resultado completo", "Coincidencias", "Excepciones", "Partidas sin soporte",
                             "Soportes no referenciados", "Evidencias", "Historial de revisiones", "Parámetros utilizados",
                             "Registro modelo OCR reglas"]
    ws = wb["Resultado completo"]
    assert ws["A1"].value == "Muenra Vouching"
    header = [c.value for c in ws[6]]
    for col in ("ID de muestra", "Archivo", "Hash del archivo (SHA-256)", "Tipo documental", "Página de evidencia", "Número extraído",
                "Fecha extraída", "Tercero extraído", "NIT extraído", "Valor extraído", "Valor esperado", "Diferencia absoluta",
                "Diferencia porcentual", "Similitud del nombre", "Coincidencia del NIT", "Coincidencia del número", "Diferencia en días",
                "Confianza OCR", "Estado", "Motivo de excepción", "Texto de evidencia", "Revisor", "Fecha de revisión"):
        assert col in header, col
    assert ws.max_row - 6 == len(demo_data["rows"])
    ev = wb["Evidencias"]
    assert ev.max_row > 100
