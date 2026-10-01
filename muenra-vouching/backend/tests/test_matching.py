"""Pruebas del motor de coincidencias con documentos sintéticos (sin OCR)."""
from __future__ import annotations

import itertools
from datetime import date
from decimal import Decimal

import pytest

from app import database, models
from app.models import Status
from app.services import exporter
from app.services.matching import reconcile, run_matching

_seq = itertools.count(1)


@pytest.fixture()
def db(client):
    s = database.SessionLocal()
    yield s
    s.close()


def make_project(db, refs: list[dict], settings: dict | None = None) -> models.Project:
    n = next(_seq)
    p = models.Project(code=f"MT-{n}", name="Motor", client_name="Cliente Demostración SAS", client_nit="900100200", settings=settings or {})
    db.add(p)
    db.flush()
    imp = models.ImportBatch(project_id=p.id, filename="x.xlsx", sha256="0" * 64, status="VALIDA", row_count=len(refs),
                             total_value=sum(Decimal(str(r.get("expected_value", 0))) for r in refs))
    db.add(imp)
    db.flush()
    for i, r in enumerate(refs, start=2):
        base = {"doc_type": "FACTURA", "currency": "COP", "doc_date": date(2026, 3, 1)}
        base.update(r)
        db.add(models.ReferenceItem(project_id=p.id, import_id=imp.id, row_number=i, **base))
    db.commit()
    return p


def make_doc(db, p, name: str, fields: dict, doc_type="FACTURA_COMPRA", file_type="pdf", sha=None, status="PROCESADO", conf=100.0, method="PDF_TEXTO"):
    d = models.Document(project_id=p.id, filename=name, stored_path="-", sha256=sha or f"{next(_seq):064d}", size_bytes=1,
                        file_type=file_type, mime="application/pdf", processing_status=status, doc_type=doc_type, full_text=" ".join(map(str, fields.values())))
    db.add(d)
    db.flush()
    for k, v in fields.items():
        db.add(models.ExtractedField(document_id=d.id, field_name=k, value=str(v), normalized_value=str(v), page=1,
                                     bbox=[0.1, 0.1, 0.5, 0.12], evidence_text=f"{k}: {v}", method=method, confidence=conf))
    db.commit()
    return d


def inv(number, nit="900123456", value=1000000, day=date(2026, 3, 1), name="Proveedor Andino S.A.S.", **extra):
    f = {"numero_documento": number, "emisor_nit": nit, "emisor_nombre": name, "valor_total": f"{value}.00",
         "fecha_emision": day.isoformat(), "moneda": "COP"}
    f.update(extra)
    return f


REF = dict(doc_number="FV-1", nit="900.123.456-8", third_party="PROVEEDOR ANDINO SAS", expected_value=Decimal("1000000"))


def results(db, p):
    run_matching(db, p.id)
    return {r.reference_item.sample_id: r for r in db.query(models.VouchingResult).filter_by(project_id=p.id)}


def test_exact_and_tolerance_and_exception(db):
    p = make_project(db, [{**REF, "sample_id": "A"}, {**REF, "sample_id": "B", "doc_number": "FV-2"}, {**REF, "sample_id": "C", "doc_number": "FV-3"}])
    make_doc(db, p, "a.pdf", inv("FV-1", value=1000000))
    make_doc(db, p, "b.pdf", inv("FV-2", value=1090000))
    make_doc(db, p, "c.pdf", inv("FV-3", value=1120000))
    r = results(db, p)
    assert r["A"].status == Status.COINCIDE
    assert r["B"].status == Status.COINCIDE_TOLERANCIA and r["B"].pct_difference == pytest.approx(0.09)
    assert r["C"].status == Status.EXCEPCION and r["C"].abs_difference == Decimal("120000")
    assert r["C"].tolerance_applied == 0.10


def test_tolerance_is_configurable(db):
    p = make_project(db, [{**REF, "sample_id": "A"}], settings={"TOLERANCIA_VALOR": 0.15})
    make_doc(db, p, "a.pdf", inv("FV-1", value=1120000))
    assert results(db, p)["A"].status == Status.COINCIDE_TOLERANCIA


def test_zero_expected_goes_to_manual_review(db):
    p = make_project(db, [{**REF, "sample_id": "Z", "expected_value": Decimal("0")}, {**REF, "sample_id": "Z2", "doc_number": "FV-2", "expected_value": Decimal("0")}])
    make_doc(db, p, "z.pdf", inv("FV-1", value=50000))
    make_doc(db, p, "z2.pdf", inv("FV-2", value=0))
    r = results(db, p)
    assert r["Z"].status == Status.REVISION_MANUAL and r["Z"].pct_difference is None
    assert r["Z2"].status == Status.COINCIDE


def test_name_alone_never_approves_when_nit_contradicts(db):
    p = make_project(db, [{**REF, "sample_id": "N", "doc_number": "XX-1"}])
    # Mismo nombre, mismo valor y fecha, pero otro NIT y otro número
    make_doc(db, p, "n.pdf", inv("YY-9", nit="800234567", name="Proveedor Andino SAS"))
    r = results(db, p)["N"]
    assert r.status == Status.SIN_SOPORTE
    # Con número coincidente se relaciona, pero como EXCEPCIÓN por el NIT
    p2 = make_project(db, [{**REF, "sample_id": "N2"}])
    make_doc(db, p2, "n2.pdf", inv("FV-1", nit="800234567"))
    r2 = results(db, p2)["N2"]
    assert r2.status == Status.EXCEPCION and any("NIT no coincide" in m for m in r2.reasons)


def test_explanation_has_separate_score_per_criterion(db):
    p = make_project(db, [{**REF, "sample_id": "E", "contract": "CT-9", "purchase_order": "OC-1", "concept": "Servicios"}])
    make_doc(db, p, "e.pdf", inv("FV-1", numero_contrato="CT-9", orden_compra="OC-1", concepto="Servicios profesionales"))
    r = results(db, p)["E"]
    crit = {c["clave"]: c for c in r.explanation["soportes_principales"][0]["criterios"]}
    assert crit["nit"]["puntaje"] == 1.0 and crit["numero"]["puntaje"] == 1.0 and crit["valor"]["puntaje"] == 1.0
    assert crit["contrato"]["puntaje"] == 1.0 and crit["orden_compra"]["puntaje"] == 1.0
    assert crit["nit"]["peso"] > crit["nombre"]["peso"] > crit["concepto"]["peso"]
    assert crit["valor"]["evidencia"] and crit["valor"]["pagina"] == 1
    assert r.status == Status.COINCIDE and r.nit_match == "COINCIDE"


def test_multiple_supports_sum(db):
    p = make_project(db, [{**REF, "sample_id": "S", "expected_value": Decimal("5000000")}])
    make_doc(db, p, "s1.pdf", inv("FV-1", value=3000000))
    make_doc(db, p, "s2.pdf", inv("FV-7", value=1500000))
    make_doc(db, p, "s3.pdf", inv("FV-8", value=500000))
    make_doc(db, p, "otro.pdf", inv("FV-99", value=777000))
    r = results(db, p)["S"]
    files = sorted(d["archivo"] for d in r.explanation["soportes_principales"])
    assert files == ["s1.pdf", "s2.pdf", "s3.pdf"]
    assert r.status == Status.COINCIDE and r.extracted_value == Decimal("5000000")
    docs = {d.filename: d for d in db.query(models.Document).filter_by(project_id=p.id)}
    assert docs["otro.pdf"].vouching_status == Status.NO_REFERENCIADO


def test_one_document_supports_several_items_without_duplicating_value(db):
    p = make_project(db, [
        {**REF, "sample_id": "P1", "doc_number": "CC-1", "expected_value": Decimal("1800000")},
        {**REF, "sample_id": "P2", "doc_number": "CC-1", "expected_value": Decimal("1200000")},
    ])
    make_doc(db, p, "cc.pdf", inv("CC-1", value=3000000), doc_type="CUENTA_COBRO")
    r = results(db, p)
    assert r["P1"].status == r["P2"].status == Status.COINCIDE
    links = db.query(models.MatchLink).filter_by(project_id=p.id, role="PRINCIPAL").all()
    assert sum(ln.allocated_value for ln in links) == Decimal("3000000")
    rec = reconcile(db, p.id)
    assert rec["sin_duplicacion_de_valor"] is True


def test_shared_document_exceeding_value_is_not_shared(db):
    p = make_project(db, [
        {**REF, "sample_id": "Q1", "doc_number": "CC-2", "expected_value": Decimal("1000000")},
        {**REF, "sample_id": "Q2", "doc_number": "CC-2", "expected_value": Decimal("1000000")},
    ])
    make_doc(db, p, "cc2.pdf", inv("CC-2", value=1000000))
    r = results(db, p)
    # Partidas idénticas en la población: ambas se marcan como posible duplicado y el documento no se reparte
    assert r["Q1"].status == r["Q2"].status == Status.POSIBLE_DUPLICADO
    links = db.query(models.MatchLink).filter_by(project_id=p.id, role="PRINCIPAL").all()
    assert len(links) == 1 and links[0].allocated_value == Decimal("1000000")


def test_duplicate_documents(db):
    p = make_project(db, [{**REF, "sample_id": "D"}])
    a = make_doc(db, p, "d.pdf", inv("FV-1"), sha="a" * 64)
    b = make_doc(db, p, "d_copia.pdf", inv("FV-1"), sha="a" * 64)
    b.duplicate_of_id = a.id
    db.commit()
    assert results(db, p)["D"].status == Status.POSIBLE_DUPLICADO
    p2 = make_project(db, [{**REF, "sample_id": "D2"}])
    make_doc(db, p2, "x1.pdf", inv("FV-1"))
    make_doc(db, p2, "x2.jpg", inv("FV-1"), file_type="jpg")  # mismo número y emisor, distinto archivo
    assert results(db, p2)["D2"].status == Status.POSIBLE_DUPLICADO


def test_xml_and_pdf_representation_are_grouped(db):
    p = make_project(db, [{**REF, "sample_id": "X"}])
    make_doc(db, p, "fe.xml", inv("FV-1", cufe="ab" * 48), doc_type="FACTURA_ELECTRONICA", file_type="xml", method="XML")
    make_doc(db, p, "fe.pdf", inv("FV-1", cufe="ab" * 48), doc_type="FACTURA_ELECTRONICA")
    r = results(db, p)["X"]
    assert r.status == Status.COINCIDE
    roles = {ln.role for ln in db.query(models.MatchLink).filter_by(project_id=p.id)}
    assert roles == {"PRINCIPAL", "REPRESENTACION_GRAFICA"}
    assert reconcile(db, p.id)["valor_asignado_a_soportes"] == "1000000.00"


def test_date_out_of_tolerance_and_currency(db):
    p = make_project(db, [{**REF, "sample_id": "F"}, {**REF, "sample_id": "M", "doc_number": "FV-5"}])
    make_doc(db, p, "f.pdf", inv("FV-1", day=date(2026, 5, 1)))
    make_doc(db, p, "m.pdf", {**inv("FV-5"), "moneda": "USD"})
    r = results(db, p)
    assert r["F"].status == Status.EXCEPCION and r["F"].days_difference == 61
    assert r["M"].status == Status.EXCEPCION


def test_low_ocr_confidence_requires_review(db):
    p = make_project(db, [{**REF, "sample_id": "O"}])
    make_doc(db, p, "o.png", inv("FV-1"), file_type="png", conf=40.0, method="OCR")
    assert results(db, p)["O"].status == Status.REVISION_MANUAL


def test_unreadable_and_error_documents(db):
    p = make_project(db, [{**REF, "sample_id": "U", "expected_file": "u.png"}, {**REF, "sample_id": "R", "doc_number": "FV-2", "expected_file": "r.pdf"}])
    make_doc(db, p, "u.png", {}, status="ILEGIBLE", doc_type=None)
    make_doc(db, p, "r.pdf", {}, status="ERROR", doc_type=None)
    r = results(db, p)
    assert r["U"].status == Status.ILEGIBLE and r["R"].status == Status.ERROR_LECTURA
    rec = reconcile(db, p.id)
    assert rec["documentos_por_estado"] == {Status.ILEGIBLE: 1, Status.ERROR_LECTURA: 1}


def test_manual_review_survives_rerun(db):
    p = make_project(db, [{**REF, "sample_id": "H"}])
    make_doc(db, p, "h.pdf", inv("FV-1", value=1500000))
    r = results(db, p)["H"]
    assert r.status == Status.EXCEPCION
    r.manual_status, r.review_decision = Status.COINCIDE_TOLERANCIA, "APROBADO"
    db.commit()
    r = results(db, p)["H"]
    assert r.auto_status == Status.EXCEPCION and r.status == Status.COINCIDE_TOLERANCIA


def test_field_correction_preserves_original(db):
    p = make_project(db, [{**REF, "sample_id": "K"}])
    d = make_doc(db, p, "k.pdf", inv("FV-1", value=1500000))
    f = next(x for x in d.fields if x.field_name == "valor_total")
    f.review_status, f.corrected_value = "CORREGIDO", "1000000.00"
    db.commit()
    assert results(db, p)["K"].status == Status.COINCIDE
    assert f.value == "1500000.00"


def test_export_sanitizes_formula_injection(db):
    p = make_project(db, [{**REF, "sample_id": "=HYPERLINK(\"http://x\")", "third_party": "+cmd|' /C calc'!A0"}])
    make_doc(db, p, "@evil.pdf", inv("FV-1"))
    run_matching(db, p.id)
    from io import BytesIO

    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(exporter.build_workbook(db, p)))
    ws = wb["Resultado completo"]
    values = [c.value for c in ws[7]]
    assert all(not (isinstance(v, str) and v[:1] in "=+@") for v in values)
    assert exporter.safe("=1+1") == "'=1+1" and exporter.safe("-2") == "'-2" and exporter.safe(5) == 5
