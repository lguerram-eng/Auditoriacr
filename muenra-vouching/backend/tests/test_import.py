"""Pruebas de importación del Excel de referencia e informe de integridad."""
import io

from openpyxl import Workbook

from app.services.excel_import import read_workbook, validate_workbook

HEADER = ["ID_MUESTRA", "TIPO_DOCUMENTO", "NUMERO_DOCUMENTO", "FECHA", "TERCERO", "NIT", "VALOR_ESPERADO"]


def workbook(rows, header=HEADER, config=None, extra_sheets=None) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    if config is not None:
        ws = wb.create_sheet("CONFIGURACION")
        ws.append(["PARAMETRO", "VALOR"])
        for r in config:
            ws.append(r)
    ws = wb.create_sheet("REFERENCIA_VOUCHING")
    ws.append(header)
    for r in rows:
        ws.append(r)
    for name, data in (extra_sheets or {}).items():
        s = wb.create_sheet(name)
        for r in data:
            s.append(r)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def validate(data, name="ref.xlsx"):
    return validate_workbook(read_workbook(name, data))


ROW = ["M1", "FACTURA", "FV-1", "2026-03-01", "Proveedor Andino SAS", "900.123.456-8", 1000]


def test_valid_workbook_and_sheet_names_case_insensitive():
    data = workbook([ROW, ["M2", "RC", "RC-2", "2026-03-02", "X SAS", "800234567", 2000]], config=[["TOLERANCIA_VALOR", "10%"], ["UMBRAL_NOMBRE", 0.9]])
    report, norm = validate(data)
    assert report["estado"] == "VALIDA", report
    assert report["resumen"]["partidas"] == 2 and report["resumen"]["valor_total"] == "3000.00"
    assert norm["params"]["TOLERANCIA_VALOR"] == 0.10 and norm["params"]["UMBRAL_NOMBRE"] == 90


def test_missing_required_columns_rejected():
    report, _ = validate(workbook([["M1", "FV-1"]], header=["ID_MUESTRA", "NUMERO_DOCUMENTO"]))
    assert report["estado"] == "RECHAZADA"
    assert any("VALOR_ESPERADO" in e for e in report["errores"])


def test_missing_reference_sheet_rejected():
    wb = Workbook()
    wb.active.title = "OTRA"
    b = io.BytesIO()
    wb.save(b)
    report, _ = validate(b.getvalue())
    assert report["estado"] == "RECHAZADA"


def test_duplicate_ids_rejected_and_business_duplicates_warned():
    report, _ = validate(workbook([ROW, ROW]))
    assert report["estado"] == "RECHAZADA"
    assert any("duplicados" in e for e in report["errores"])
    row2 = ["M2"] + ROW[1:]
    report, _ = validate(workbook([ROW, row2]))
    assert report["estado"] == "CON_ADVERTENCIAS"
    assert any("repetidas" in w for w in report["advertencias"])


def test_invalid_dates_empty_values_and_types():
    rows = [
        ["M1", "FACTURA", "FV-1", "31/02/2026", "X", "900.123.456-8", 100],
        ["M2", "FACTURA", None, "2026-03-01", None, "900.123.456-8", 100],
        ["M3", "FACTURA", "FV-3", "2026-03-01", "X", "900.123.456-1", 100],
    ]
    report, norm = validate(workbook(rows))
    issues = {(i["fila"], i["columna"]) for i in report["incidencias_por_fila"]}
    assert (2, "FECHA") in issues
    assert (3, "NUMERO_DOCUMENTO") in issues and (3, "TERCERO") in issues
    assert (4, "NIT") in issues  # DV incorrecto
    assert norm["references"][0]["doc_date"] is None
    report, _ = validate(workbook([["M1", "FACTURA", "FV-1", "2026-03-01", "X", "1", "mil pesos"]]))
    assert report["estado"] == "RECHAZADA"  # valor no numérico


def test_control_totals_reconciliation():
    rows = [ROW, ["M2"] + ROW[1:2] + ["FV-2"] + ROW[3:6] + [500]]
    ok, _ = validate(workbook(rows, config=[["CONTROL_TOTAL_REGISTROS", 2], ["CONTROL_TOTAL_VALOR", 1500]]))
    assert ok["conciliacion"]["registros_cuadran"] and ok["conciliacion"]["valor_cuadra"]
    bad, _ = validate(workbook(rows, config=[["CONTROL_TOTAL_REGISTROS", 3], ["CONTROL_TOTAL_VALOR", 1600]]))
    assert bad["estado"] == "RECHAZADA" and len(bad["errores"]) == 2


def test_auxiliary_sheets():
    data = workbook([ROW], extra_sheets={
        "ENTIDADES_ALIAS": [["NIT", "NOMBRE_CANONICO", "ALIAS"], ["900123456", "Proveedor Andino SAS", "ANDINO PROV"]],
        "TIPOS_DOCUMENTO": [["CODIGO", "NOMBRE", "PALABRAS_CLAVE"], ["RC", "Recibo", "recibo oficial; caja menor"]],
        "CAMPOS_EXTRACCION": [["CAMPO", "OBLIGATORIO"], ["valor_total", "SI"]],
        "RESULTADO_ESPERADO": [["ID_MUESTRA", "ESTADO_ESPERADO"], ["M1", "Coincide"], ["M9", "SIN SOPORTE"]],
    })
    report, norm = validate(data)
    assert norm["aliases"][0]["alias"] == "ANDINO PROV"
    assert norm["types"][0]["code"] == "RECIBO_CAJA" and norm["types"][0]["keywords"] == ["recibo oficial", "caja menor"]
    assert norm["fields"][0]["required"] is True
    assert norm["expected"][0]["expected_status"] == "COINCIDE"
    assert any("M9" in w for w in report["advertencias"])


def test_csv_import():
    csv = "ID_MUESTRA;TIPO_DOCUMENTO;NUMERO_DOCUMENTO;FECHA;TERCERO;NIT;VALOR_ESPERADO\nM1;FACTURA;FV-1;01/03/2026;X SAS;900123456-8;1.500.000\n"
    report, norm = validate(csv.encode("utf-8"), "ref.csv")
    assert report["estado"] in ("VALIDA", "CON_ADVERTENCIAS")
    assert str(norm["references"][0]["expected_value"]) == "1500000.00"


def test_import_endpoint_dry_run_and_rejection(client, admin_headers):
    pid = client.post("/api/projects", headers=admin_headers, json={"code": "IMP-1", "name": "Importación", "client_name": "Cliente"}).json()["id"]
    f = {"file": ("ref.xlsx", workbook([ROW, ROW]), "application/octet-stream")}
    r = client.post(f"/api/projects/{pid}/imports", headers=admin_headers, files=f)
    assert r.status_code == 422 and r.json()["informe"]["estado"] == "RECHAZADA"
    r = client.post(f"/api/projects/{pid}/imports?dry_run=true", headers=admin_headers, files={"file": ("ref.xlsx", workbook([ROW]), "application/octet-stream")})
    assert r.status_code == 200 and r.json()["importado"] is False
    assert client.get(f"/api/projects/{pid}/references", headers=admin_headers).json() == []
    r = client.post(f"/api/projects/{pid}/imports", headers=admin_headers, files={"file": ("ref.xlsx", workbook([ROW]), "application/octet-stream")})
    assert r.status_code == 200 and r.json()["importado"] is True
    assert len(client.get(f"/api/projects/{pid}/references", headers=admin_headers).json()) == 1
    r = client.post(f"/api/projects/{pid}/imports", headers=admin_headers, files={"file": ("ref.xlsx", b"%PDF-1.4 falso", "application/octet-stream")})
    assert r.status_code == 422
