"""Plantilla simple (hoja CARGA, una fila por documento) — importación y recorrido completo."""
from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app import worker
from app.services.excel_import import read_workbook, split_contract_or_order, validate_workbook
from app.services.template import HEADER_ROW, SIMPLE_HEADER, build_simple_template

from .conftest import requires_ocr


def validate(data: bytes):
    return validate_workbook(read_workbook("plantilla.xlsx", data))


def test_template_layout():
    wb = load_workbook(io.BytesIO(build_simple_template()))
    ws = wb["CARGA"]
    assert ws["A1"].value.startswith("MUENRA VOUCHING")
    assert [c.value for c in ws[HEADER_ROW]] == SIMPLE_HEADER
    assert ws.cell(HEADER_ROW + 1, 1).value == "MV-001"
    assert len(ws.data_validations.dataValidation) == 3


def test_simple_template_with_title_rows_imports():
    report, norm = validate(build_simple_template())
    assert report["estado"] == "CON_ADVERTENCIAS" and not report["errores"]
    assert any("CARGA" in i for i in report["informacion"])
    assert report["hojas_reconocidas"] == {"CARGA": True}
    assert report["columnas_mapeadas"]["VALOR_ESPERADO"] == "VALOR_TOTAL"
    assert report["columnas_mapeadas"]["NIT"] == "NIT_IDENTIFICACION"
    assert report["columnas_mapeadas"]["FECHA"] == "FECHA_DOCUMENTO"
    refs = {r["sample_id"]: r for r in norm["references"]}
    assert refs["MV-001"]["expected_value"] == Decimal("1190000.00") and refs["MV-001"]["doc_date"] == date(2026, 9, 27)
    assert refs["MV-001"]["purchase_order"] == "OC-4500123" and refs["MV-001"]["contract"] is None
    assert refs["MV-003"]["contract"] == "CT-2026-021" and refs["MV-003"]["purchase_order"] is None
    assert refs["MV-001"]["extra"]["IVA"] == "190000.00" and refs["MV-001"]["extra"]["TOLERANCIA_VALOR"] == 0.1
    assert refs["MV-002"]["expected_file"] == "rc_315.pdf"
    # Las filas de ejemplo se advierten, no se descartan
    assert sum(1 for i in report["incidencias_por_fila"] if "EJEMPLO" in i["detalle"]) == 3


def test_any_document_type_and_tolerance_formats():
    rows = [
        ["X-1", "OTRO", "ACTA-7", date(2026, 1, 5), "Tercero Uno SAS", "900123456", "Acta de entrega", None, None, None, 1000, "COP", None, None, 10, None],
        ["X-2", "CONTRATO", "CT-9", date(2026, 1, 6), "Tercero Dos", "800234567", "Arrendamiento", None, None, None, 2000, "COP", "CT-9", None, None, None],
        ["X-3", "CUENTA_COBRO", "CC-1", date(2026, 1, 7), "Persona Natural Ficticia", "1020304050", "Honorarios", None, None, None, 3000, "COP", None, None, 7, None],
    ]
    report, norm = validate(build_simple_template(rows, examples=False))
    assert report["estado"] == "VALIDA", report["incidencias_por_fila"]
    refs = {r["sample_id"]: r for r in norm["references"]}
    assert refs["X-1"]["extra"]["TOLERANCIA_VALOR"] == pytest.approx(0.10)  # 10 -> 10 %
    assert refs["X-3"]["extra"]["TOLERANCIA_VALOR"] == pytest.approx(0.07)
    assert "TOLERANCIA_VALOR" not in refs["X-2"]["extra"]


def test_split_contract_or_order():
    assert split_contract_or_order("OC-4500123") == (None, "OC-4500123")
    assert split_contract_or_order("Orden 77") == (None, "Orden 77")
    assert split_contract_or_order("CT-2026-021") == ("CT-2026-021", None)


def test_user_style_file_with_blank_row_before_header():
    """Mismo diseño de la plantilla entregada: título, instrucciones, fila vacía y encabezado en la fila 4."""
    report, norm = validate(build_simple_template())
    assert [r["row_number"] for r in norm["references"]] == [5, 6, 7]


def test_template_endpoint(client, admin_headers):
    r = client.get("/api/template/simple", headers=admin_headers)
    assert r.status_code == 200 and r.content[:2] == b"PK"


def _simple_rows_from_demo(rows):
    out = []
    for r in rows:
        sid, tipo, num, fecha, tercero, nit, valor, moneda, contrato, oc, concepto, _cc, _cta, archivo = r
        out.append([sid, tipo, num, fecha, tercero, nit, concepto, None, None, None, valor, moneda, contrato or oc, archivo, 0.10, None])
    return out


@requires_ocr
def test_end_to_end_with_simple_template(client, admin_headers, demo_data):
    r = client.post("/api/projects", headers=admin_headers, json={
        "code": "DEMO-SIMPLE", "name": "Plantilla simple", "client_name": "Cliente Demostración S.A.S.", "client_nit": "900.100.200-0"})
    pid = r.json()["id"]
    xlsx = build_simple_template(_simple_rows_from_demo(demo_data["rows"]), examples=False)
    r = client.post(f"/api/projects/{pid}/imports", headers=admin_headers, files={"file": ("carga.xlsx", xlsx, "application/octet-stream")})
    assert r.status_code == 200 and r.json()["importado"], r.text
    files = [("files", (n, d, "application/octet-stream")) for n, d in demo_data["docs"].items()]
    assert not client.post(f"/api/projects/{pid}/documents", headers=admin_headers, files=files).json()["rechazados"]
    worker.drain()
    got = {x["id_muestra"]: x["estado"] for x in client.get(f"/api/projects/{pid}/results", headers=admin_headers).json()}
    expected = {sid: st for sid, st, _f, _n in demo_data["expected"]}
    assert got == expected
    rec = client.get(f"/api/projects/{pid}/reconciliation", headers=admin_headers).json()
    assert rec["conteo_partidas_cuadra"] and rec["valor_cuadra"] and rec["conteo_documentos_cuadra"]


def test_row_tolerance_and_informative_components(client):
    from .test_matching import REF, inv, make_doc, make_project, results  # reutiliza utilidades
    from app import database

    db = database.SessionLocal()
    try:
        p = make_project(db, [
            {**REF, "sample_id": "T15", "extra": {"TOLERANCIA_VALOR": 0.15, "IVA": "190000.00"}},
            {**REF, "sample_id": "T10", "doc_number": "FV-2"},
        ])
        make_doc(db, p, "a.pdf", inv("FV-1", value=1120000, iva="150000.00"))
        make_doc(db, p, "b.pdf", inv("FV-2", value=1120000))
        r = results(db, p)
        assert r["T15"].status == "COINCIDE CON TOLERANCIA" and r["T15"].tolerance_applied == 0.15
        assert any("IVA (referencia)" in m for m in r["T15"].reasons)
        crit = {c["clave"] for c in r["T15"].explanation["soportes_principales"][0]["criterios"]}
        assert "iva" in crit
        assert r["T10"].status == "EXCEPCIÓN" and r["T10"].tolerance_applied == 0.10
    finally:
        db.close()
